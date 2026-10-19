"""Vigilance des TL validés : préparation des check-lists, lecture des retours, décisions sur les HIT.

Parcours
  1. Après validation des OV, les TL validés sans vigilance (ou à refaire) sont regroupés par
     **client + bénéficiaire** : un seul fichier si, pour un client, le bénéficiaire est le même sur tous
     ses TL ; sinon un fichier par bénéficiaire distinct (nom + IBAN + SWIFT).
  2. Chaque unité remplit une copie du modèle « CHECK LIST TRANSFERT LIBRE » (modèle avant vigilance,
     chargé par l'utilisateur dans les référentiels). Les cellules sont repérées par leur LIBELLÉ en colonne A
     (« Client: », « Banque du bénéficiaire »…) et non par une adresse fixe : une ligne ajoutée au modèle
     ne décale rien.
  3. Retours (modèle après vigilance, colonne « HIT / No HIT / KO ») : classés RAS / HIT / KO / INFO /
     MANQUANT / INCONNU. Un HIT bloque les TL de l'unité jusqu'à décision motivée (levée ou confirmation).
     Un KO « plateforme indisponible » = vigilance à refaire. Un résultat vide ou inconnu = à vérifier.
  4. Seuls les TL au statut RAS ou HIT levé passent au registre.

Journal append-only `04_transferts/vigilance/journal_vigilance.jsonl` (ENVOI, RETOUR, DECISION).
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import unicodedata
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

import pandas as pd

from .common import digits, is_blank, siege_racine

# ------------------------------------------------------------------ constantes
CIRCUIT = "TL"
TYPE_DOSSIER = "Création"
LABELS = {  # clé interne -> libellé de la colonne A du modèle (comparaison normalisée)
    "client": "Client:", "ref": "Réf Dossier:", "circuit": "Circuit:", "type": "Type de dossier:",
    "banque": "Banque du bénéficiaire", "pays_banque": "Pays de la banque du bénéficiaire",
    "beneficiaire": "Bénéficiaire", "pays_beneficiaire": "Pays du bénéficiaire",
    "lieu_travail": "Lieu de travail", "individus": "Divers individus sur les documents",
}
HEADER_LABEL = "Contrôle"            # ligne d'en-tête du tableau (A10 dans le modèle actuel)
RESULT_HEADER = "HIT / No HIT / KO"
STATUTS_OK = {"RAS", "HIT_LEVE"}     # TL autorisés à passer au registre
STATUTS_A_DECIDER = {"HIT", "A_VERIFIER"}
LIBELLE_STATUT = {
    None: "Vigilance à faire", "EN_ATTENTE": "Envoyée, retour attendu", "RAS": "RAS — aucune alerte",
    "HIT": "HIT — décision requise", "HIT_LEVE": "HIT levé (motivé)", "HIT_CONFIRME": "HIT confirmé — TL rejeté",
    "KO_A_REFAIRE": "Plateforme indisponible — à refaire", "A_VERIFIER": "Résultat à vérifier — décision requise",
}
SWIFT_PAYS = {"TR": "TURQUIE", "DZ": "ALGERIE", "FR": "FRANCE", "DE": "ALLEMAGNE", "IT": "ITALIE", "ES": "ESPAGNE",
              "CN": "CHINE", "IN": "INDE", "GB": "ROYAUME-UNI", "BE": "BELGIQUE", "CH": "SUISSE", "AE": "EMIRATS ARABES UNIS",
              "EG": "EGYPTE", "TN": "TUNISIE", "MA": "MAROC", "PK": "PAKISTAN", "PH": "PHILIPPINES", "US": "ETATS-UNIS"}
MAX_ZIP_FILES, MAX_ZIP_BYTES = 2000, 200_000_000


def text(v) -> str:
    return "" if is_blank(v) else str(v).strip()


def norm(v) -> str:
    s = "".join(c for c in unicodedata.normalize("NFKD", text(v)) if not unicodedata.combining(c))
    return " ".join(s.upper().replace("’", "'").split())


def safe(v: str) -> str:
    """Pas de formule Excel injectée par une valeur lue (=, +, -, @ en tête)."""
    v = text(v)
    return v.lstrip("=+-@ ").strip() if v[:1] in ("=", "+", "-", "@") else v


def safe_name(v: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]", "_", v)


def lines(v) -> list[str]:
    """Lignes d'une cellule ; les lignes vides internes sont gardées (sinon les rangs C/D se décalent)."""
    t = text(v)
    return t.replace("\r\n", "\n").replace("\r", "\n").split("\n") if t else []


def classify(raw) -> tuple[str, str]:
    n = norm(raw)
    if not n:
        return "MANQUANT", "Résultat à obtenir"
    if "RUBRIQUE NON RENSEIGNEE" in n or "RUBRIQUE NON RENSEINGEE" in n:
        return "INFO", "Rubrique vide, aucune action"
    if re.match(r"^KO\b", n) and "INDISPONIBLE" in n:
        return "KO", "Vigilance à refaire"
    if re.fullmatch(r"(OK\s*=+>\s*)?(RAS|NOT HIT|NO HIT)", n):
        return "RAS", "Aucune alerte"
    if re.fullmatch(r"(OK\s*=+>\s*)?HIT", n):
        return "HIT", "Alerte à examiner"
    return "INCONNU", "Résultat à vérifier"


def format_ref(sr) -> str:
    """Réf dossier = identifiant client (siège + racine), au format « SSSSS RRRRRR » du modèle."""
    d = digits(sr)
    return f"{d[:5]} {d[5:]}" if len(d) > 5 else d


def pays_banque(swift, pays_annexe) -> str:
    if text(pays_annexe):
        return norm(pays_annexe)
    s = norm(swift).replace(" ", "")
    return SWIFT_PAYS.get(s[4:6], s[4:6]) if len(s) >= 6 else ""


def same_person(a, b) -> bool:
    """Bénéficiaire = client, dans l'ordre nom/prénom ou prénom/nom (mêmes mots)."""
    wa, wb = norm(a).split(), norm(b).split()
    return bool(wa) and sorted(wa) == sorted(wb)


# ------------------------------------------------------------------ modèle : cellules repérées par libellé
def locate(ws) -> dict:
    """{clé: numéro de ligne} + 'first'/'last' : lignes du tableau des rubriques."""
    rows = {}
    want = {norm(lab): k for k, lab in LABELS.items()}
    header = None
    for r in range(1, ws.max_row + 1):
        a = norm(ws.cell(r, 1).value)
        if not a:
            continue
        if a in want and want[a] not in rows:
            rows[want[a]] = r
        if a == norm(HEADER_LABEL) and header is None:
            header = r
    missing = [LABELS[k] for k in LABELS if k not in rows]
    if missing or header is None:
        raise ValueError("Modèle de vigilance non reconnu : libellés absents en colonne A : "
                         + ", ".join(missing + ([] if header else [HEADER_LABEL])))
    last = header
    for r in range(header + 1, ws.max_row + 1):
        if text(ws.cell(r, 2).value):          # colonne Type (C / E / I / V) renseignée = rubrique
            last = r
    rows["first"], rows["last"], rows["header"] = header + 1, last, header
    return rows


def check_template(path: Path) -> list[str]:
    from openpyxl import load_workbook
    try:
        ws = load_workbook(path).worksheets[0]
        r = locate(ws)
    except Exception as exc:  # noqa: BLE001
        return [str(exc)]
    issues = []
    if norm(ws.cell(r["header"], 4).value) != norm(RESULT_HEADER):
        issues.append(f"En-tête de la colonne D : « {ws.cell(r['header'], 4).value} » (attendu « {RESULT_HEADER} »)")
    return issues


# ------------------------------------------------------------------ journal
def journal_path(P) -> Path:
    return P.tr_root / "vigilance" / "journal_vigilance.jsonl"


def read_journal(path: Path) -> list[dict]:
    if not path.exists():
        return []
    out = []
    for ln in path.read_text(encoding="utf-8").splitlines():
        try:
            out.append(json.loads(ln))
        except json.JSONDecodeError:
            continue
    return out


def append(path: Path, user: str, action: str, **kw) -> dict:
    ev = {"ts": datetime.now().isoformat(timespec="seconds"), "user": user, "action": action, **kw}
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(ev, ensure_ascii=False, default=str) + "\n")
    return ev


def statuses(journal: list[dict]) -> dict[str, dict]:
    """Statut courant de chaque TL (fichier OV) : dernier envoi, puis retour, puis décision."""
    unit_of, units, st = {}, {}, {}
    for ev in journal:
        a = ev.get("action")
        if a == "ENVOI":
            for u in ev.get("unites", []):
                units[u["uid"]] = dict(u, lot=ev.get("lot"))
                for f in u["fichiers"]:
                    unit_of[f] = u["uid"]
                    st[f] = {"statut": "EN_ATTENTE", "uid": u["uid"], "lot": ev.get("lot"), "ts": ev["ts"]}
        elif a == "RETOUR":
            u = units.get(ev.get("uid"))
            for f in (u or {}).get("fichiers", []):
                if unit_of.get(f) == ev["uid"]:
                    st[f] = {"statut": ev["statut"], "uid": ev["uid"], "lot": u.get("lot"), "ts": ev["ts"],
                             "detail": ev.get("detail"), "fichier_retour": ev.get("fichier_retour"),
                             "chemin": ev.get("chemin")}
        elif a == "DECISION":
            for f in ev.get("fichiers", []):
                if f in st and st[f]["statut"] in STATUTS_A_DECIDER:
                    st[f] = dict(st[f], statut="HIT_LEVE" if ev["decision"] == "LEVEE" else "HIT_CONFIRME",
                                 motif=ev.get("motif"), decide_par=ev.get("user"), decide_le=ev["ts"],
                                 fiche_annotee=ev.get("fiche_annotee"))
    return st


# ------------------------------------------------------------------ unités de vigilance
def beneficiary_key(d: dict) -> str:
    return "|".join([norm(d.get("beneficiaire")), norm(d.get("iban")).replace(" ", ""), norm(d.get("swift"))])


def build_units(dossiers: list[dict]) -> list[dict]:
    """dossiers : [{fichier, ref_client, client, beneficiaire, iban, swift, banque, pays_banque, individus[]}]
    -> unités (client + bénéficiaire), avec la liste des TL couverts."""
    groups: dict[tuple, list[dict]] = {}
    for d in dossiers:
        groups.setdefault((d["ref_client"], beneficiary_key(d)), []).append(d)
    units = []
    per_client = Counter(k[0] for k in groups)
    rank: Counter = Counter()
    for (ref, bkey), ds in sorted(groups.items()):
        d0 = ds[0]
        rank[ref] += 1
        indiv = []
        for d in ds:                                   # individus de tous les TL du groupe, sans doublon
            for x in d.get("individus") or []:
                if x and norm(x) not in {norm(y) for y in indiv}:
                    indiv.append(x)
        uid = hashlib.sha1(f"{ref}|{bkey}".encode()).hexdigest()[:8].upper()
        ben = "" if same_person(d0.get("beneficiaire"), d0.get("client")) else text(d0.get("beneficiaire"))
        units.append({
            "uid": uid, "ref": format_ref(ref), "ref_client": ref, "client": text(d0.get("client")),
            "beneficiaire": ben, "beneficiaire_lu": text(d0.get("beneficiaire")), "iban": text(d0.get("iban")),
            "swift": text(d0.get("swift")), "banque": text(d0.get("banque")), "pays_banque": text(d0.get("pays_banque")),
            "individus": indiv, "fichiers": [d["fichier"] for d in ds],
            "fichier": f"Vigilance_{safe_name(format_ref(ref))}"
                       + (f"_B{rank[ref]}" if per_client[ref] > 1 else "") + f"_{uid}.xlsx",
        })
    return units


def dossiers_from_ov(sy: pd.DataFrame, audit: pd.DataFrame, fichiers: list[str], parents=None,
                     individus_format: str = "entier") -> list[dict]:
    """Assemble les données de vigilance d'un TL depuis l'OV (synthèse + valeurs lues/corrigées).
    parents(ligne de synthèse) -> [père, mère] (valeurs retenues dans la matrice DOM du dossier rattaché)."""
    a = audit.set_index(audit["fichier"].astype(str)) if not audit.empty else pd.DataFrame()
    out = []
    for f in fichiers:
        r = sy[sy["FICHIER"] == f]
        if r.empty:
            continue
        r = r.iloc[0].to_dict()
        d = a.loc[f].to_dict() if f in a.index else {}
        if isinstance(d, pd.DataFrame):
            d = d.iloc[0].to_dict()
        ref = siege_racine(r.get("SIEGE_RACINE")) or text(r.get("SIEGE_RACINE"))
        indiv = []
        for p in (parents(r) if parents else []) or []:
            if text(p):
                indiv.append(fmt_individu(p, individus_format))
        sig = text(d.get("ANNEXE_I.nom_prenom_signataire"))
        for x in lines(sig):
            if x.strip():
                indiv.append(fmt_individu(x, individus_format))
        out.append({
            "fichier": f, "ref_client": ref, "client": text(r.get("NOM")),
            "beneficiaire": text(r.get("BENEFICIAIRE") or d.get("OV.beneficiaire_nom")),
            "iban": text(r.get("BENEFICIAIRE_COMPTE") or d.get("OV.beneficiaire_compte")),
            "swift": text(r.get("SWIFT") or d.get("OV.swift")),
            "banque": text(r.get("BANQUE") or d.get("OV.banque_nom") or d.get("ANNEXE_II.banque_etranger")),
            "pays_banque": pays_banque(r.get("SWIFT") or d.get("OV.swift"), d.get("ANNEXE_II.pays_destination")),
            "individus": indiv,
        })
    return out


def fmt_individu(v, mode: str = "entier") -> str:
    """'entier' : nom complet tel que lu ; 'diese' : premier mot # reste (convention prénom#nom du KYC)."""
    v = " ".join(text(v).split())
    if mode == "diese" and " " in v:
        a, b = v.split(" ", 1)
        return f"{a}#{b}"
    return v


# ------------------------------------------------------------------ préparation du lot
def prepare(units: list[dict], template: Path, dest: Path, lieu_travail: str) -> dict:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    if not units:
        raise ValueError("Aucun TL à soumettre à la vigilance")
    dest.mkdir(parents=True, exist_ok=True)
    for u in units:
        wb = load_workbook(template)
        ws = wb.worksheets[0]
        r = locate(ws)
        for row in range(r["first"], r["last"] + 1):
            if text(ws.cell(row, 2).value) and type(ws.cell(row, 4)).__name__ != "MergedCell":
                ws.cell(row, 4).value = None
        vals = {"client": u["client"], "ref": u["ref"], "circuit": CIRCUIT, "type": TYPE_DOSSIER,
                "banque": u["banque"], "pays_banque": u["pays_banque"], "beneficiaire": u["beneficiaire"],
                "lieu_travail": lieu_travail, "individus": "\n".join(u["individus"])}
        for k, v in vals.items():
            c = ws.cell(r[k], 3)
            c.value = safe(v) if k != "individus" else "\n".join(safe(x) for x in u["individus"])
            c.value = c.value or None
        ws.cell(r["individus"], 3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r["individus"]].height = max(30, 16 * max(1, len(u["individus"])))
        wb.save(dest / u["fichier"])
    lot = datetime.now().strftime("VIG_%Y%m%d_%H%M%S")
    manifest = {"lot": lot, "cree_le": datetime.now().isoformat(timespec="seconds"), "unites": units,
                "lieu_travail": lieu_travail}
    (dest / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for u in units:
            z.write(dest / u["fichier"], u["fichier"])
    (dest / f"{lot}_a_deposer.zip").write_bytes(buf.getvalue())
    recap = pd.DataFrame([{"Réf dossier": u["ref"], "Client": u["client"], "Bénéficiaire (si ≠ client)": u["beneficiaire"],
                           "Banque": u["banque"], "Pays banque": u["pays_banque"], "IBAN": u["iban"], "SWIFT": u["swift"],
                           "Divers individus": "\n".join(u["individus"]), "Nb TL": len(u["fichiers"]),
                           "TL couverts": "\n".join(u["fichiers"]), "Fichier": u["fichier"]} for u in units])
    recap.to_excel(dest / f"{lot}_recapitulatif.xlsx", index=False)
    manifest["zip"] = str(dest / f"{lot}_a_deposer.zip")
    manifest["recap"] = str(dest / f"{lot}_recapitulatif.xlsx")
    return manifest


# ------------------------------------------------------------------ lecture des retours
def iter_return_files(uploads: list[tuple[str, bytes]]):
    """(nom, octets) des .xlsx reçus, directement ou dans un ZIP (lu en mémoire, sans extraction sur disque)."""
    for name, data in uploads:
        if name.lower().endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                infos = [i for i in z.infolist() if not i.is_dir()]
                if len(infos) > MAX_ZIP_FILES or sum(i.file_size for i in infos) > MAX_ZIP_BYTES:
                    raise ValueError(f"{name} : ZIP trop volumineux")
                for i in infos:
                    yield Path(i.filename).name, z.read(i)
        else:
            yield Path(name).name, data


def read_return(data: bytes) -> dict:
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data), data_only=True).worksheets[0]
    r = locate(ws)
    rubriques = []
    for row in range(r["first"], r["last"] + 1):
        lab = text(ws.cell(row, 1).value)
        if not lab or not text(ws.cell(row, 2).value):   # titres de section (« VIGILANCE ») ignorés
            continue
        rubriques.append({"ligne": row, "rubrique": lab, "valeur": text(ws.cell(row, 3).value),
                          "resultat": text(ws.cell(row, 4).value), "commentaire": text(ws.cell(row, 5).value)})
    return {"client": text(ws.cell(r["client"], 3).value), "ref": text(ws.cell(r["ref"], 3).value),
            "beneficiaire": text(ws.cell(r["beneficiaire"], 3).value), "banque": text(ws.cell(r["banque"], 3).value),
            "individus": lines(ws.cell(r["individus"], 3).value), "rubriques": rubriques,
            "individus_row": r["individus"]}


def match_unit(name: str, ret: dict, units: dict[str, dict]) -> tuple[dict | None, str]:
    for uid, u in units.items():                    # 1. identifiant TL-One dans le nom du fichier
        if uid in name.upper():
            return u, "nom de fichier"
    cand = [u for u in units.values() if norm(u["ref"]).replace(" ", "") == norm(ret["ref"]).replace(" ", "")]
    if len(cand) > 1:                               # 2. même client : bénéficiaire puis banque
        cand = [u for u in cand if norm(u["beneficiaire"]) == norm(ret["beneficiaire"])] or cand
    if len(cand) > 1:
        cand = [u for u in cand if norm(u["banque"]) == norm(ret["banque"])] or cand
    if len(cand) == 1:
        return cand[0], "réf dossier"
    return None, "ambigu" if cand else "inconnu"


def analyse(u: dict, ret: dict) -> tuple[str, list[list], list[str]]:
    """Statut de l'unité + détail par rubrique / individu + contrôles de structure."""
    detail, ctl = [], []
    if norm(ret["client"]) != norm(u["client"]):
        ctl.append(f"Client différent : « {ret['client']} » (envoyé « {u['client']} »)")
    for rb in ret["rubriques"]:
        names = lines(rb["valeur"]) if rb["ligne"] == ret["individus_row"] else [rb["valeur"]]
        results = lines(rb["resultat"]) if rb["ligne"] == ret["individus_row"] else [rb["resultat"]]
        if rb["ligne"] == ret["individus_row"]:
            if [norm(x) for x in names] != [norm(x) for x in u["individus"]]:
                ctl.append("Divers individus différents du lot envoyé")
            only_info = not names and len(results) == 1 and classify(results[0])[0] == "INFO"
            if len(names) != len(results) and not only_info:
                ctl.append(f"Décalage individus / résultats ({len(names)} / {len(results)})")
        for k in range(max(len(names), len(results), 1)):
            raw = results[k] if k < len(results) else ""
            st, action = classify(raw)
            if not (names[k] if k < len(names) else "") and st == "MANQUANT":
                st, action = "INFO", "Rubrique vide"
            detail.append([rb["rubrique"], k + 1, names[k] if k < len(names) else "", raw, st, action, rb["commentaire"]])
    c = Counter(x[4] for x in detail)
    if c["HIT"]:
        statut = "HIT"
    elif c["KO"]:
        statut = "KO_A_REFAIRE"
    elif c["MANQUANT"] or c["INCONNU"] or ctl:
        statut = "A_VERIFIER"
    else:
        statut = "RAS"
    return statut, detail, ctl


def import_returns(uploads: list[tuple[str, bytes]], journal: list[dict], jpath: Path, user: str,
                   out_dir: Path) -> dict:
    """Lit les retours, écrit un événement RETOUR par unité reconnue et un classeur de résultats."""
    units = {}
    for ev in journal:
        if ev.get("action") == "ENVOI":
            for u in ev.get("unites", []):
                units[u["uid"]] = dict(u, lot=ev.get("lot"))
    synth, details, ctls, seen = [], [], [], Counter()
    for name, data in iter_return_files(uploads):
        if not name.lower().endswith(".xlsx"):
            ctls.append([name, "", "Fichier ignoré (pas un .xlsx)"])
            continue
        try:
            ret = read_return(data)
        except Exception as exc:  # noqa: BLE001
            ctls.append([name, "", f"Illisible : {exc}"])
            continue
        u, how = match_unit(name, ret, units)
        if u is None:
            ctls.append([name, ret["ref"], f"Retour non rattaché à un envoi ({how})"])
            continue
        seen[u["uid"]] += 1
        statut, detail, ctl = analyse(u, ret)
        if seen[u["uid"]] > 1:
            ctl.append("Retour reçu plusieurs fois pour cette unité")
        c = Counter(x[4] for x in detail)
        keep = out_dir / "fichiers"                     # fiche retournée conservée : consultation et annotation
        keep.mkdir(parents=True, exist_ok=True)
        chemin = keep / f"{u['uid']}_{datetime.now():%Y%m%d_%H%M%S}_{safe_name(Path(name).stem)}.xlsx"
        chemin.write_bytes(data)
        append(jpath, user, "RETOUR", uid=u["uid"], lot=u.get("lot"), statut=statut, fichier_retour=name,
               detail={k: c[k] for k in ("RAS", "HIT", "KO", "INFO", "MANQUANT", "INCONNU")}, controles=ctl,
               hits=[x for x in detail if x[4] == "HIT"], chemin=str(chemin),
               points=[x for x in detail if x[4] in ("HIT", "MANQUANT", "INCONNU")])
        synth.append({"Réf dossier": u["ref"], "Client": u["client"], "Bénéficiaire": u["beneficiaire"] or "(client)",
                      "Statut": statut, "HIT": c["HIT"], "KO": c["KO"], "À vérifier": c["MANQUANT"] + c["INCONNU"],
                      "RAS": c["RAS"], "TL concernés": len(u["fichiers"]), "Contrôles": " | ".join(ctl),
                      "Fichier retour": name, "Rapprochement": how, "UID": u["uid"]})
        details += [[u["ref"], u["client"], name, *x] for x in detail]
        ctls += [[name, u["ref"], x] for x in ctl]
    sy = pd.DataFrame(synth)
    de = pd.DataFrame(details, columns=["Réf dossier", "Client", "Fichier retour", "Rubrique", "Rang", "Valeur",
                                        "Résultat brut", "Statut", "Action", "Commentaire"])
    co = pd.DataFrame(ctls, columns=["Fichier", "Réf dossier", "Contrôle"])
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / datetime.now().strftime("resultats_vigilance_%Y%m%d_%H%M%S.xlsx")
    with pd.ExcelWriter(path) as xw:
        sy.to_excel(xw, sheet_name="Synthese", index=False)
        de.to_excel(xw, sheet_name="Detail", index=False)
        co.to_excel(xw, sheet_name="Controles", index=False)
    return {"synthese": sy, "detail": de, "controles": co, "classeur": path}


def decide(jpath: Path, user: str, fichiers: list[str], decision: str, motif: str, uid: str | None = None) -> dict:
    if decision not in ("LEVEE", "CONFIRMEE"):
        raise ValueError("décision inconnue")
    if not text(motif):
        raise ValueError("Motif obligatoire")
    return append(jpath, user, "DECISION", uid=uid, fichiers=fichiers, decision=decision, motif=motif.strip())


VERDICTS = {"NON_AVERE": "Non avéré (faux positif)", "AVERE": "Avéré"}


def annotate(data: bytes, verdicts: list[dict], user: str, ts: str) -> bytes:
    """Fiche retournée + analyse du valideur dans la colonne COMMENTAIRES, sur la ligne de chaque point analysé."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font
    wb = load_workbook(io.BytesIO(data))
    ws = wb.worksheets[0]
    rows = {norm(ws.cell(r, 1).value): r for r in range(1, ws.max_row + 1) if text(ws.cell(r, 1).value)}
    par_ligne: dict[int, list[str]] = {}
    for v in verdicts:
        r = rows.get(norm(v["rubrique"]))
        if not r:
            continue
        qui = f" — {v.get('valeur')}" if text(v.get("valeur")) else ""
        par_ligne.setdefault(r, []).append(f"{'AVÉRÉ' if v['verdict'] == 'AVERE' else 'NON AVÉRÉ'}{qui} : "
                                           f"{text(v['commentaire'])}")
    for r, notes in par_ligne.items():
        c = ws.cell(r, 5)
        old = text(c.value)
        c.value = "\n".join(([old] if old else []) + notes + [f"({user}, {ts[:16].replace('T', ' ')})"])
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(bold=True, color="B3261E" if any(n.startswith("AVÉRÉ") for n in notes) else "1E6B3A")
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def analyse_points(jpath: Path, user: str, uid: str, fichiers: list[str], verdicts: list[dict],
                   retour: Path | None, out_dir: Path) -> dict:
    """Analyse des HIT / points à vérifier saisie dans TL-One.
    Chaque point : verdict AVERE / NON_AVERE + commentaire obligatoire. Tous non avérés -> TL libéré (LEVEE) ;
    au moins un avéré -> alerte confirmée (TL rejeté). La fiche annotée est conservée pour l'archivage."""
    if not verdicts:
        raise ValueError("Aucun point à analyser")
    for v in verdicts:
        if v.get("verdict") not in VERDICTS:
            raise ValueError(f"{v.get('rubrique')} : choisir avéré ou non avéré")
        if not text(v.get("commentaire")):
            raise ValueError(f"{v.get('rubrique')} (rang {v.get('rang')}) : commentaire obligatoire")
    decision = "CONFIRMEE" if any(v["verdict"] == "AVERE" for v in verdicts) else "LEVEE"
    motif = " | ".join(f"{v['rubrique']}{' ' + str(v['valeur']) if v.get('valeur') else ''} : "
                       f"{'avéré' if v['verdict'] == 'AVERE' else 'non avéré'} — {text(v['commentaire'])}"
                       for v in verdicts)
    ts = datetime.now().isoformat(timespec="seconds")
    annotee = None
    if retour and Path(retour).exists():
        out_dir.mkdir(parents=True, exist_ok=True)
        annotee = out_dir / f"{Path(retour).stem}_analyse.xlsx"
        annotee.write_bytes(annotate(Path(retour).read_bytes(), verdicts, user, ts))
    return append(jpath, user, "DECISION", uid=uid, fichiers=fichiers, decision=decision, motif=motif,
                  verdicts=verdicts, fiche_annotee=str(annotee) if annotee else None)


def candidates(decisions: pd.DataFrame, st: dict[str, dict]) -> list[str]:
    """TL validés à soumettre : jamais envoyés, ou plateforme indisponible (à refaire)."""
    if decisions.empty:
        return []
    ok = decisions[decisions["DECISION"] == "VALIDE"]["FICHIER"].astype(str)
    return [f for f in ok if f not in st or st[f]["statut"] == "KO_A_REFAIRE"]


def eligible(st: dict[str, dict]) -> set[str]:
    return {f for f, v in st.items() if v["statut"] in STATUTS_OK}
