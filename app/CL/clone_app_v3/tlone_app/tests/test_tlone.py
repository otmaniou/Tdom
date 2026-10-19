import sys
from datetime import date
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tests import fixtures  # noqa: E402
from tlone import domvalid, nbrunner, planning, transfers  # noqa: E402
from tlone.config import load_settings  # noqa: E402
from tlone.montant_lettres import lettres_to_number  # noqa: E402


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("TLONE_DATA_ROOT", str(tmp_path))
    monkeypatch.setenv("TLONE_RUN_MODE", "local")
    p = fixtures.build(tmp_path)
    return p, load_settings()


# ------------------------------------------------------------ montant en lettres
@pytest.mark.parametrize("txt,val", [
    ("UN MILLION DEUX CENT SOIXANTE ET UN MILLE SIX CENTS DINARS ALGERIENS", 1261600),
    ("TROIS CENT DOUZE MILLE DEUX CENT CINQUANTE CINQ DINARS", 312255),
    ("CENT QUATRE VINGT MILLE DINARS", 180000),
    ("QUATRE-VINGT-DIX-NEUF MILLE NEUF CENT QUATRE-VINGT-DIX-NEUF DA", 99999),
    ("DEUX CENT MILLE DINARS ET 50 CTS", 200000.50),
    ("SOIXANTE DIX SEPT MILLE DINARS ALGÉRIENS ET QUARANTE CENTIMES", 77000.40),
    ("MILLE DINARS", 1000),
])
def test_lettres(txt, val):
    assert lettres_to_number(txt) == pytest.approx(val)


def test_lettres_illisible():
    assert lettres_to_number("TROIS CENT BLURP MILLE") is None


# ------------------------------------------------------------ runner
def test_patch_source_root_and_glob():
    src = "INPUT_DIR = Path('/mnt/data/domiciliations_in')\nMAX_PDFS = None\npdfs = sorted(INPUT_DIR.glob('*.pdf'))\n" \
          "assert RAW_JSON_DIR == Path('/mnt/data/x')\nMODEL_PATH = '/domino/edv/old'\n"
    out, applied = nbrunner.patch_source(src, "dom_extract", "/domino/datasets/local/tlone", 5, "/models/new")
    assert "/mnt/data" not in out and "/domino/datasets/local/tlone/domiciliations_in" in out
    assert "MAX_PDFS = 5" in out and "suffix.lower() == '.pdf'" in out
    assert 'MODEL_PATH = "/models/new"' in out
    assert any(a.startswith("DATA_ROOT") for a in applied)


def test_patch_default_root_untouched():
    src = "OUTPUT_ROOT = Path('/mnt/data/document_pipeline/V15_A')\n"
    out, applied = nbrunner.patch_source(src, "dom_post", "/mnt/data", None, None)
    assert out == src and applied == []


def test_real_notebooks_patchable():
    s = load_settings()
    for step, nb in s.notebooks.items():
        prepared, applied = nbrunner.prepare_notebook(nb, step, "/tmp/x", 3, None)
        assert any(a.startswith("DATA_ROOT") for a in applied), step
        code = "\n".join(c.source for c in prepared.cells if c.cell_type == "code")
        assert "/mnt/data" not in code
        if step in ("dom_extract", "ov_extract"):
            assert "GLOB_PDF_CASE_INSENSITIVE" in applied
        assert any(a.startswith(("MAX_", "pdfs[")) for a in applied), (step, applied)


# ------------------------------------------------------------ module 1
def test_domvalid_roundtrip(env):
    p, _ = env
    man = domvalid.load_manual(p.validation_xlsx)
    assert list(man["_ROW"]) == [2, 3, 4, 5]
    mt = p.validation_xlsx.stat().st_mtime
    ed = man.copy()
    ed.loc[ed["FICHIER"] == "CELIK_MEHMET_CNT.pdf", ["DECISION_MANUELLE", "VALEUR_VALIDEE"]] = ["CORRIGEE", ""]
    assert domvalid.check_decisions(ed)  # CORRIGEE sans valeur bloqué
    ed.loc[ed["FICHIER"] == "CELIK_MEHMET_CNT.pdf", "VALEUR_VALIDEE"] = "02700731010830000161"
    assert not domvalid.check_decisions(ed)
    res = domvalid.save_decisions(p.validation_xlsx, ed, expected_mtime=mt)
    assert len(res["changes"]) == 2 and Path(res["backup"]).exists()
    re = domvalid.load_manual(p.validation_xlsx)
    assert re.loc[re["FICHIER"] == "CELIK_MEHMET_CNT.pdf", "DECISION_MANUELLE"].iloc[0] == "CORRIGEE"
    # Les autres feuilles sont conservées
    assert "DOSSIERS_A_VALIDER" in domvalid.sheet_names(p.validation_xlsx)
    with pytest.raises(RuntimeError):
        domvalid.save_decisions(p.validation_xlsx, ed, expected_mtime=mt)  # fichier modifié entre-temps


def test_progress(env):
    p, _ = env
    pr = domvalid.progress(domvalid.load_manual(p.validation_xlsx)).set_index("FICHIER")
    assert pr.loc["DEMIR_ALI_CNT.pdf", "STATUT_VALIDATION"] == "REJETE"
    assert pr.loc["CELIK_MEHMET_CNT.pdf", "RESTANT"] == 1


# ------------------------------------------------------------ module 2
def _planning(p, excl=False):
    dos = domvalid.load_sheet(p.validation_xlsx, "DOSSIERS_A_VALIDER")
    man = domvalid.load_sheet(p.validation_xlsx, "A_VALIDER_MANUELLEMENT")
    pa = pd.read_excel(p.param_aug_xlsx, dtype=object)
    return planning.build_planning(dos, man, pa, exclure_validation_incomplete=excl)


def test_planning_rules(env):
    p, _ = env
    r = _planning(p)
    pl, rej = r["Planning_TL"], r["REJETS_PLANNING"].set_index("FICHIER")
    assert r["_dq"] == []
    assert rej.loc["DEMIR_ALI_CNT.pdf", "MOTIF"] == "REJETE_EN_PARTIE_2A"
    assert "COMPTE_FINAL_ABSENT" in rej.loc["CELIK_MEHMET_CNT.pdf", "MOTIF"]
    kaya = pl[pl["NOM_TRAVAILLEUR"] == "KAYA AHMET"]
    assert kaya["DOM_ID"].nunique() == 1  # contrat + augmentation sur la même DOM
    jan = kaya[kaya["MOIS_BASE"] == "2026-01"].iloc[0]
    assert jan["PERIODE_TL"] == "2026-01P2" and jan["NB_JOURS_SEGMENT"] == 29
    assert jan["MONTANT_MAX_THEORIQUE"] == round(312255 * 29 / 31, 2)
    mar = kaya[kaya["MOIS_BASE"] == "2026-03"].iloc[0]
    apr = kaya[kaya["MOIS_BASE"] == "2026-04"].iloc[0]
    assert mar["TYPE_VERSION"] == "NOUVEAU_CONTRAT" and mar["MONTANT_MAX_THEORIQUE"] == 312255
    assert apr["TYPE_VERSION"] == "AUGMENTATION" and apr["NUMERO_AUGMENTATION"] == 1
    assert apr["MONTANT_MAX_THEORIQUE"] == 332500
    last = kaya.iloc[-1]
    assert last["PERIODE_TL"] == "2028-01P1" and last["NB_JOURS_SEGMENT"] == 2


def test_planning_matches_notebook(env, monkeypatch):
    """Exécute le VRAI notebook Planning_TL_V15A sur le même jeu et compare."""
    p, _ = env
    import nbformat
    from nbclient import NotebookClient
    nb = nbformat.read(str(Path(__file__).resolve().parent.parent / "notebooks" / "Planning_TL_V15A.ipynb"), 4)
    for c in nb.cells:
        if c.cell_type == "code":
            c.source = c.source.replace("/mnt/data", str(p.data_root))
    NotebookClient(nb, timeout=120, kernel_name="python3").execute()
    ref = pd.read_csv(p.planning_csv, dtype=str).fillna("")
    mine = _planning(p)["Planning_TL"].astype(str).replace({"None": "", "nan": ""})
    ref = ref[mine.columns]
    for c in ("MONTANT_MAX_THEORIQUE", "COEFFICIENT_PRORATA", "PART_TRANSFERABLE_CTS"):
        assert (ref[c].astype(float).round(6) == mine[c].replace("", "nan").astype(float).round(6)).all(), c
    for c in ("PERIODE_ID", "DOM_ID", "VERSION_ID", "PERIODE_TL", "TYPE_VERSION", "DATE_DEBUT_SEGMENT"):
        assert list(ref[c]) == list(mine[c]), c


# ------------------------------------------------------------ module 3
def test_transfer_controls(env):
    p, s = env
    r = _planning(p)
    planning.export(r, p.planning_xlsx, p.planning_csv)
    pl = planning.load_planning(p.planning_csv)
    audit = transfers.read_audit_xlsx(transfers.latest_audit(p.tr_out)[0])
    assert len(audit) == 4 and "OV.montant_chiffres" in audit.columns
    atf = transfers.load_referential(p.atf_xlsx, transfers.ATF_COLUMNS)
    aut = transfers.load_referential(p.autorisations_xlsx, transfers.AUT_COLUMNS)
    reg = transfers.load_registre(p.registre_csv)
    sy, ch = transfers.run_controls(audit, pl, reg, atf, aut, s, today=date(2026, 4, 10))
    sy = sy.set_index("FICHIER")
    mar = sy.loc["07310-108259-202603.pdf"]
    assert mar["STATUT_CONTROLE"] == "CONFORME", ch[ch.FICHIER == "07310-108259-202603.pdf"].to_dict("records")
    assert mar["PERIODE_TL"] == "2026-03" and mar["PROVISION_ATLAS"] == 317255
    feb = sy.loc["07310-108259-202602.pdf"]
    assert feb["STATUT_CONTROLE"] == "NON_CONFORME" and "T09" in feb["MOTIFS"] and "T10" in feb["MOTIFS"]
    jan = sy.loc["07310-108259-202601P2.pdf"]
    assert jan["PERIODE_TL"] == "2026-01P2" and jan["STATUT_CONTROLE"] in ("A_REVOIR", "NON_CONFORME")
    demir = sy.loc["07310-108262-202605.pdf"]
    assert demir["STATUT_CONTROLE"] == "NON_CONFORME" and "T04" in demir["MOTIFS"]

    # Décision + registre + consommation
    ed = sy.reset_index()[["FICHIER", "AUDIT_SOURCE", "STATUT_CONTROLE", "PERIODE_ID"]].copy()
    ed["DECISION"], ed["COMMENTAIRE"] = "", ""
    ed.loc[ed.FICHIER == "07310-108259-202603.pdf", "DECISION"] = "VALIDE"
    ed.loc[ed.FICHIER == "07310-108262-202605.pdf", "DECISION"] = "VALIDE"
    assert transfers.check_decisions(ed)  # non rattaché ne peut être validé
    ed.loc[ed.FICHIER == "07310-108262-202605.pdf", "DECISION"] = "REJETE"
    ed.loc[ed.FICHIER == "07310-108262-202605.pdf", "COMMENTAIRE"] = "DOM rejetée"
    assert not transfers.check_decisions(ed)
    assert transfers.save_decisions(p.decisions_csv, ed, "testeur") == 2
    dec = transfers.load_decisions(p.decisions_csv)
    assert transfers.register(p.registre_csv, sy.reset_index(), dec, "testeur") == 1
    assert transfers.register(p.registre_csv, sy.reset_index(), dec, "testeur") == 0  # idempotent
    reg = transfers.load_registre(p.registre_csv)
    cons = transfers.consumption(pl, reg, today=date(2026, 4, 10)).set_index("PERIODE_TL")
    assert cons.loc["2026-03", "STATUT_PERIODE"] == "TRANSFERE"
    assert cons.loc["2026-02", "STATUT_PERIODE"] == "NON_TRANSFERE"
    # Un 2e OV sur mars est désormais bloqué (déjà transféré)
    sy2, ch2 = transfers.run_controls(audit.iloc[[0]].assign(fichier="BIS-202603.pdf"), pl, reg, atf, aut, s,
                                      today=date(2026, 4, 10))
    assert sy2.iloc[0]["STATUT_CONTROLE"] in ("NON_CONFORME",)


def test_periodes():
    assert transfers.parse_periode("MARS 2026 PART 1") == ("2026-03", "P1")
    assert transfers.parse_periode("FÉVRIER.2026") == ("2026-02", None)
    assert transfers.parse_periode("03/2026") == ("2026-03", None)
    assert transfers.periode_from_filename("02700731011155100123-202510P2.PDF") == ("2025-10", "P2")
    assert transfers.periode_from_filename("07310-108262-202605.pdf") == ("2026-05", None)
    assert transfers.periode_from_filename("07310-109655-042026p2.pdf") == ("2026-04", "P2")
    assert transfers.periode_from_filename("02700731010854700144-202606-OV-PART-2_COKLUK-AHMET.pdf") == ("2026-06", "P2")
    assert transfers.periode_from_filename("02700731010854700144-202607-OV_COKLUK-AHMET.pdf") == ("2026-07", None)
    assert transfers.periode_from_filename("07310-109655-04 p1.pdf") == (None, "P1")
    from tlone.common import siege_racine
    assert siege_racine("02700731010854700144-202607-OV_COKLUK") == "07310108547"
    assert siege_racine("07310-109655-04 p1") == "07310109655"


def test_carry_over(env):
    p, _ = env
    from tlone.common import backup
    prev = backup(p.validation_xlsx)
    # simulate a Part 2 rerun : decisions wiped
    man = domvalid.load_manual(p.validation_xlsx)
    wiped = man.copy()
    for c in domvalid.EDITABLE:
        wiped[c] = ""
    domvalid.save_decisions(p.validation_xlsx, wiped)
    todo = domvalid.carry_over(prev, p.validation_xlsx)
    assert len(todo) == 3  # 3 lignes décidées dans le jeu, la 4e était vide
    domvalid.save_decisions(p.validation_xlsx, todo)
    again = domvalid.load_manual(p.validation_xlsx)
    assert list(again["DECISION_MANUELLE"]) == list(man["DECISION_MANUELLE"])


def test_jobs_local_launch(env, monkeypatch):
    p, s = env
    from tlone import jobs
    from tlone import preflight
    monkeypatch.setattr(jobs.subprocess, "Popen", lambda *a, **k: type("P", (), {"pid": 4242})())
    monkeypatch.setattr(preflight, "check_step", lambda s, step: [])
    rec = jobs.launch(s, "dom_post", 3, "testeur")
    assert rec["backup_validation"] and Path(rec["backup_validation"]).exists()
    runs = jobs.list_runs(s)
    assert runs[0]["run_id"] == rec["run_id"] and runs[0]["status"] == "EN_FILE"


def test_step_runtime(monkeypatch):
    monkeypatch.setenv("TLONE_GPU_TIER", "gpu-commun")
    monkeypatch.setenv("TLONE_DOM_EXTRACT_TIER", "h100-80g")
    monkeypatch.setenv("TLONE_DOM_EXTRACT_ENV", "env-qwen38-tf5")
    monkeypatch.setenv("TLONE_OV_EXTRACT_ENV", "env-qwen25-tf4")
    monkeypatch.setenv("TLONE_OV_EXTRACT_PYTHON", "/opt/venv-qwen25/bin/python")
    s = load_settings()
    a, b, c = s.step_runtime("dom_extract"), s.step_runtime("ov_extract"), s.step_runtime("dom_post")
    assert a["tier"] == "h100-80g" and a["environment"] == "env-qwen38-tf5"
    assert b["tier"] == "gpu-commun" and b["environment"] == "env-qwen25-tf4"
    assert b["python"] == "/opt/venv-qwen25/bin/python"
    assert c["environment"] == "" and c["model"] == ""


def test_domino_payload(env, monkeypatch):
    p, _ = env
    from tlone import jobs
    monkeypatch.setenv("TLONE_RUN_MODE", "domino")
    monkeypatch.setenv("DOMINO_PROJECT_ID", "proj1")
    monkeypatch.setenv("TLONE_OV_EXTRACT_ENV", "env-qwen25")
    monkeypatch.setenv("TLONE_OV_EXTRACT_TIER", "gpu-a10")
    from tlone import preflight
    sent = {}
    monkeypatch.setattr(preflight, "check_step", lambda s, step: [])
    monkeypatch.setattr(jobs, "_domino_request", lambda m, path, payload=None: sent.update(payload) or {"id": "J1"})
    rec = jobs.launch(load_settings(), "ov_extract", 2, "testeur")
    assert sent["environmentId"] == "env-qwen25" and sent["overrideHardwareTierId"] == "gpu-a10"
    assert sent["commandToRun"].endswith(f"ov_extract --run-id {rec['run_id']} --data-root {p.data_root} --max 2")


# ------------------------------------------------------------ contrôle de démarrage
def test_preflight_bloque_sans_pdf_ni_tier(env, monkeypatch):
    p, _ = env
    from tlone import jobs, preflight
    monkeypatch.setenv("TLONE_RUN_MODE", "domino")
    monkeypatch.setenv("DOMINO_PROJECT_ID", "proj1")
    monkeypatch.setenv("DOMINO_API_PROXY", "http://localhost:8899")
    monkeypatch.delenv("TLONE_GPU_TIER", raising=False)
    s = load_settings()
    for f in p.dom_in.glob("*"):
        f.unlink()
    ko = {r["contrôle"] for r in preflight.blocking(preflight.check_step(s, "dom_extract"))}
    assert "PDF DOM à traiter" in ko and "Tier GPU du job" in ko
    called = []
    monkeypatch.setattr(jobs, "_domino_request", lambda *a, **k: called.append(a) or {"id": "X"})
    with pytest.raises(RuntimeError, match="prérequis bloquants"):
        jobs.launch(s, "dom_extract", 2, "testeur")
    assert not called  # aucun job Domino n'est parti


def test_preflight_data_root_et_double_lancement(env, monkeypatch):
    p, s = env
    from tlone import jobs, preflight
    assert not [r for r in preflight.blocking(preflight.check_general(s)) if r["contrôle"] == "TLONE_DATA_ROOT"]
    monkeypatch.delenv("TLONE_DATA_ROOT")
    assert any(r["contrôle"] == "TLONE_DATA_ROOT" for r in preflight.blocking(preflight.check_general(load_settings())))
    # un run en attente bloque un 2e lancement, jusqu'à « Marquer abandonné »
    monkeypatch.setattr(preflight, "check_step", lambda s, step: [])
    monkeypatch.setattr(jobs.subprocess, "Popen", lambda *a, **k: type("P", (), {"pid": 1})())
    rec = jobs.launch(s, "dom_post", None, "testeur")
    monkeypatch.undo()
    monkeypatch.setenv("TLONE_DATA_ROOT", str(p.data_root))
    monkeypatch.setenv("TLONE_RUN_MODE", "local")
    assert [r["run_id"] for r in preflight.pending_runs(s, "dom_post")] == [rec["run_id"]]
    jobs.mark_abandoned(s, rec["run_id"], "testeur")
    assert preflight.pending_runs(s, "dom_post") == []


def test_preflight_carthago_entetes(tmp_path):
    from tlone import preflight
    good, bad = tmp_path / "ok.xlsx", tmp_path / "ko.xlsx"
    pd.DataFrame({"CLIENT ": ["07310104411"], "Nom complet / Raison sociale": ["X"]}).to_excel(good, index=False)
    pd.DataFrame({"Matricule": ["07310104411"], "Intitulé": ["X"]}).to_excel(bad, index=False)
    assert not preflight.blocking(preflight.check_carthago(good))
    ko = preflight.blocking(preflight.check_carthago(bad))
    assert len(ko) == 2 and "Matricule" in ko[0]["action"]


def test_carthago_entetes_reels_et_patch_partie2(tmp_path, monkeypatch):
    """En-têtes de l'export CarthagoDom réel (18/10/2026) : la Partie 2 reçoit les bons noms de colonnes."""
    from tlone import carthago, nbrunner, preflight
    reel = ["Accord pré-dom", "Agence", "Client", "Code bic", "Date apurement", "Date domiciliation",
            "Date déclaration", "Description", "Devise", "Dom Amount Tolerance", "Dom Documents Received",
            "Docs Received Comment", "Docs Received Date", "Invoices Date", "Invoices Reference",
            "Last Expedition Date", "Last Financing Settlement Date", "Mode de règlement", "Montant final",
            "Montant initial", "Name", "Nature opération", "Nom complet", "Numéro de domiciliation"]
    x = tmp_path / "Dossier de domiciliation.xlsx"
    pd.DataFrame([{c: None for c in reel}]).to_excel(x, index=False)
    c = carthago.detect(carthago.read_headers(x))
    assert (c["client"], c["name"], c["dom_ref"], c["agence"]) == ("Client", "Nom complet", "Numéro de domiciliation", "Agence")
    assert not preflight.blocking(preflight.check_carthago(x))
    src = "CARTHAGO_CLIENT_COL = 'Client'\nCARTHAGO_NAME_COL = 'Nom complet/Raison sociale'\nCARTHAGO_DOM_REF_COL = None\n"
    out, applied = nbrunner.patch_source(src, "dom_post", "/mnt/data", None, None, c)
    assert 'CARTHAGO_NAME_COL = "Nom complet"' in out and 'CARTHAGO_DOM_REF_COL = "Numéro de domiciliation"' in out
    assert "CARTHAGO_CLIENT_COL = 'Client'" in out and len(applied) == 2
    monkeypatch.setenv("TLONE_CARTHAGO_NAME_COL", "Name")
    assert carthago.detect(reel)["name"] == "Name"
    # le vrai notebook Partie 2 contient bien les lignes patchées
    import nbformat
    nb = nbformat.read(str(load_settings().notebooks["dom_post"]), as_version=4)
    full = "\n".join(cl.source for cl in nb.cells if cl.cell_type == "code")
    out, applied = nbrunner.patch_source(full, "dom_post", "/mnt/data", None, None, c)
    assert {"CARTHAGO_NAME_COL=Nom complet", "CARTHAGO_DOM_REF_COL=Numéro de domiciliation"} <= set(applied)


def test_job_cmd_python_override():
    from tlone.jobs import _job_cmd
    assert _job_cmd("python /mnt/tlone_app/run_step.py", "/opt/venv-tf4/bin/python") == \
        "/opt/venv-tf4/bin/python /mnt/tlone_app/run_step.py"
    assert _job_cmd("python /mnt/tlone_app/run_step.py", "") == "python /mnt/tlone_app/run_step.py"


def test_ensure_kernel():
    import sys
    from jupyter_client.kernelspec import KernelSpecManager
    name = nbrunner.ensure_kernel()
    spec = KernelSpecManager().get_kernel_spec(name)
    assert spec.argv[0] == sys.executable


def test_job_cmd_follows_app_dir(monkeypatch):
    monkeypatch.delenv("TLONE_JOB_CMD", raising=False)
    from tlone.config import APP_DIR
    assert load_settings().job_cmd_prefix == f"python {APP_DIR / 'run_step.py'}"


# ------------------------------------------------------------ matrice de validation DOM
def _mx(p, journal=None, sources=None):
    from tlone import matrice
    lot = matrice.load_lot(p)
    ms = matrice.build_all(lot, journal or matrice.read_journal(p.dom_journal), sources or matrice.default_sources())
    return {m["fichier"]: m for m in ms}


def _open(m):
    return {(x["ligne"].key, c) for x in m["rows"] for c, cell in x["cells"].items() if cell["open"]}


def test_matrice_anomalies(env):
    p, _ = env
    ms = _mx(p)
    assert _open(ms["KAYA_AHMET_CNT.pdf"]) == {("fin_permis", "TTR"), ("part", "DOM")}
    assert _open(ms["KAYA_AHMET_AUG_26.pdf"]) == {("sig_emp", "CTR")}
    assert ("nom", "TTR") in _open(ms["DEMIR_ALI_CNT.pdf"]) and ("domref", "CARTHAGO") in _open(ms["DEMIR_ALI_CNT.pdf"])
    demir = {x["ligne"].key: x for x in ms["DEMIR_ALI_CNT.pdf"]["rows"]}
    assert demir["domref"]["missing_ref"]  # Carthago ne connaît pas le client : rien à retenir
    assert ("compte", "DOM") in _open(ms["CELIK_MEHMET_CNT.pdf"])
    # salaire ancien seulement pour une augmentation
    assert "salaire_ancien" in {x["ligne"].key for x in ms["KAYA_AHMET_AUG_26.pdf"]["rows"]}
    assert "salaire_ancien" not in {x["ligne"].key for x in ms["KAYA_AHMET_CNT.pdf"]["rows"]}


def test_matrice_source_confirmation_correction(env):
    p, _ = env
    from tlone import matrice
    f = "KAYA_AHMET_CNT.pdf"
    # le valideur prend le TTR comme référence pour la fin de permis : c'est alors le CTS qui est en écart
    matrice.append_event(p.dom_journal, "v1", f, "SOURCE", ligne="fin_permis", col="TTR")
    m = _mx(p)[f]
    fp = next(x for x in m["rows"] if x["ligne"].key == "fin_permis")
    assert fp["retenu_txt"] == "02/01/2027" and fp["cells"]["CTS"]["open"]
    # confirmation de la lecture DOM (confiance faible) puis correction du CTS
    matrice.append_event(p.dom_journal, "v1", f, "CONFIRME", ligne="part", col="DOM", valeur_vue="312255.00")
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="fin_permis", col="CTS", valeur="02/01/2027")
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="fin_permis", col="CTR", valeur="02/01/2027")
    # la nouvelle fin rend le DOM incohérent et la durée CTS (24 mois) n'est plus l'écart des dates (12 mois)
    assert _open(_mx(p)[f]) == {("fin_permis", "DOM"), ("duree", "CTS")}
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="fin_permis", col="DOM", valeur="02/01/2027")
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="duree", col="CTS", valeur="12")
    # durée CTS 12 : DOM (24), CTR (24) et le montant DOM ne collent plus
    assert _open(_mx(p)[f]) == {("duree", "DOM"), ("duree", "CTR"), ("montant", "DOM")}
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="duree", col="DOM", valeur="12")
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="duree", col="CTR", valeur="12")
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="montant", col="DOM", valeur="3747060.00")
    m = _mx(p)[f]
    assert m["n_open"] == 0, _open(m)
    # une confirmation ne vaut que pour la valeur vue : si la Partie 2 relit autre chose, la case se rouvre
    matrice.append_event(p.dom_journal, "v1", f, "CONFIRME", ligne="part", col="DOM", valeur_vue="999")
    assert ("part", "DOM") in _open(_mx(p)[f])


def test_matrice_vers_planning(env):
    p, _ = env
    from tlone import domvalid, matrice
    f = "KAYA_AHMET_CNT.pdf"
    j = p.dom_journal
    matrice.append_event(j, "v1", f, "SOURCE", ligne="fin_permis", col="TTR")
    for col in ("CTS", "CTR"):
        matrice.append_event(j, "v1", f, "CORRIGE", ligne="fin_permis", col=col, valeur="02/01/2027")
    matrice.append_event(j, "v1", f, "CONFIRME", ligne="part", col="DOM", valeur_vue="312255.00")
    matrice.append_event(j, "v1", f, "VALIDE")
    matrice.append_event(j, "v2", "DEMIR_ALI_CNT.pdf", "REJETE", motif="Domiciliation introuvable dans Carthago")
    ms = list(_mx(p).values())
    rows = matrice.planning_rows(ms)
    fin = rows[(rows["FICHIER"] == f) & (rows["CHAMP"] == "CTS_DATE_FIN_VALIDITE_PERMIS")].iloc[0]
    assert fin["VALEUR_VALIDEE"] == "02/01/2027" and fin["DECISION_MANUELLE"] == "CORRIGEE"
    assert rows[(rows["FICHIER"] == "DEMIR_ALI_CNT.pdf")]["DECISION_MANUELLE"].tolist() == ["REJETEE"]
    man = matrice.merge_manual(domvalid.load_sheet(p.validation_xlsx, domvalid.MANUAL_SHEET), rows)
    dos = domvalid.load_sheet(p.validation_xlsx, domvalid.DOSSIERS_SHEET)
    pa = pd.read_excel(p.param_aug_xlsx, dtype=object)
    val = {m["fichier"] for m in ms if m["decision"] and m["decision"]["action"] == "VALIDE"}
    res = planning.build_planning(dos, man, pa, exclure_validation_incomplete=True, fichiers_valides=val)
    base = res["BASE_VERSIONS_CTS"]
    assert set(base["FICHIER"]) == {f}
    assert str(base.iloc[0]["CTS_DATE_FIN_VALIDITE_PERMIS"]) == "2027-01-02"
    rej = res["REJETS_PLANNING"].set_index("FICHIER")["MOTIF"]
    assert rej["DEMIR_ALI_CNT.pdf"] == "REJETE_EN_PARTIE_2A"
    assert rej["CELIK_MEHMET_CNT.pdf"] == "VALIDATION_HUMAINE_INCOMPLETE"


def test_sources_reference_parametre(env):
    p, _ = env
    from tlone import matrice
    t = matrice.sources_table()
    t.loc[t["LIGNE"] == "part", "SOURCE_REFERENCE"] = "DOM"
    t.loc[t["LIGNE"] == "nom", "SOURCE_REFERENCE"] = "ATLAS"  # source inconnue : ignorée
    p.sources_reference_xlsx.parent.mkdir(parents=True, exist_ok=True)
    t.to_excel(p.sources_reference_xlsx, index=False)
    src = matrice.default_sources(p.sources_reference_xlsx)
    assert src["part"] == "DOM" and src["nom"] == "CTS"


def test_page_rotation(tmp_path):
    import json
    from tlone import matrice
    (tmp_path / "x.json").write_text(json.dumps({"source_file": "A.pdf", "page_records": [
        {"page_num": 2, "rotation_clockwise_applied": 270}]}), encoding="utf-8")
    (tmp_path / "A.json").write_text(json.dumps({"source_file": "A.pdf", "page_records": [
        {"page_num": 2, "rotation_clockwise_applied": 90}]}), encoding="utf-8")
    assert matrice.page_rotation(tmp_path, "A.pdf", 2) == 90 and matrice.page_rotation(tmp_path, "A.pdf", 1) == 0


def test_matrice_bloquants(env):
    p, _ = env
    from tlone import matrice
    j = p.dom_journal
    # signature absente : une confirmation ne suffit pas (seule une correction ou un rejet)
    f = "KAYA_AHMET_AUG_26.pdf"
    matrice.append_event(j, "v1", f, "CONFIRME", ligne="sig_emp", col="CTR", valeur_vue="False")
    m = _mx(p)[f]
    assert ("sig_emp", "CTR") in _open(m) and m["n_bloq"] == 1
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="sig_emp", col="CTR", valeur="True")
    assert _mx(p)[f]["n_open"] == 0
    # compte : la correction est re-contrôlée (longueur, clé Modulo 97)
    f = "CELIK_MEHMET_CNT.pdf"
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="compte", col="DOM", valeur="02700731010830000155")
    cell = next(x for x in _mx(p)[f]["rows"] if x["ligne"].key == "compte")["cells"]["DOM"]
    assert cell["open"] and "attendue 61" in cell["issues"][0][1]
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="compte", col="DOM", valeur="02700731010830000161")
    assert ("compte", "DOM") not in _open(_mx(p)[f])
    # Carthago ne connaît pas le client : bloquant, non confirmable
    f = "DEMIR_ALI_CNT.pdf"
    matrice.append_event(j, "v1", f, "CONFIRME", ligne="domref", col="CARTHAGO", valeur_vue="None")
    assert ("domref", "CARTHAGO") in _open(_mx(p)[f])
    assert matrice.compte_issues("02700731010910800156")[0][1] == "Clé Modulo 97 invalide : lue 56, attendue 91"


# ------------------------------------------------------------ matrice OV
def test_ov_matrice_et_correction(env):
    p, s = env
    from tlone import ov_matrice as ov
    res = _planning(p)
    planning.export(res, p.planning_xlsx, p.planning_csv)
    pl = planning.load_planning(p.planning_csv)
    a = transfers.read_audit_xlsx(transfers.latest_audit(p.tr_out)[0])
    args = (pl, transfers.load_registre(p.registre_csv),
            transfers.load_referential(p.atf_xlsx, transfers.ATF_COLUMNS),
            transfers.load_referential(p.autorisations_xlsx, transfers.AUT_COLUMNS), s)
    f = "07310-108259-202602.pdf"
    sy, ch = transfers.run_controls(a, *args)
    row = sy.set_index("FICHIER").loc[f]
    assert row["STATUT_CONTROLE"] == "NON_CONFORME"
    m = ov.build_ov(f, a[a["fichier"] == f].iloc[0].to_dict(), sy[sy["FICHIER"] == f].iloc[0].to_dict(), ch, pl, {})
    plaf = next(x for x in m["rows"] if x["ligne"].key == "plafond")
    assert plaf["cells"]["ANNEXE_II"]["open"] and plaf["cells"]["ANNEXE_II"]["bloq"]
    # le valideur corrige le montant lu sur l'OV : les contrôles sont recalculés sur la valeur corrigée
    ov.append_event(ov_journal := p.tr_root / "journal_validation_ov.jsonl", "v1", f, "CORRIGE",
                    champ="OV.montant_chiffres", valeur="300000")
    a2 = ov.apply_corrections(a, ov.read_journal(ov_journal))
    sy2, ch2 = transfers.run_controls(a2, *args)
    t09 = ch2[(ch2["FICHIER"] == f) & (ch2["CODE"] == "T09")].iloc[0]
    assert t09["RESULTAT"] == "OK"
    # échantillon de contrôle reproductible
    fs = [f"F{i}" for i in range(100)]
    assert ov.sample(fs, "lot1") == ov.sample(fs, "lot1") and len(ov.sample(fs, "lot1")) == 5
    assert len(ov.sample(fs[:10], "x")) == 3
    lf = ov.list_frame(sy2, transfers.load_decisions(p.decisions_csv))
    assert len(lf) == len(sy2) and {"FEU", "MARGE", "BLOQUANTS"} <= set(lf.columns)


def test_matrice_regles_agence():
    from tlone import matrice as mx
    assert mx.permis_parts("20-00010423/31-26-000472") == ("20-00010423", "31-26-000472")
    assert mx.permis_parts("20/00010423 / 31/26/000472") == ("20-00010423", "31-26-000472")
    assert mx.permis_parts("31-26-000472") == (None, "31-26-000472")
    assert mx.same("permis", "20-00010423/31-26-000472", "31/26/000472")[0] is True
    assert mx.same("permis", "20-00010423/31-26-000472", "20-00010423/31-26-000473")[0] is False
    assert mx.to_months("2 ANS, 0 JOURS") == 24 and mx.to_months("1 AN 6 MOIS") == 18 and mx.to_months("24") == 24
    from datetime import date
    assert mx.months_between(date(2026, 1, 3), date(2028, 1, 2)) == 24
    assert mx.months_between(date(2026, 2, 1), date(2027, 1, 31)) == 12


def test_nouveaux_champs_pris_en_charge(env):
    """Les 4 champs demandés par l'agence : « à extraire » aujourd'hui, contrôlés dès qu'ils sont dans les JSON."""
    p, _ = env
    f = "KAYA_AHMET_CNT.pdf"

    def cell(m, k, c):
        return next(x for x in m["rows"] if x["ligne"].key == k)["cells"][c]
    m = _mx(p)[f]
    assert cell(m, "sig_trav", "DOM")["a_extraire"] and not cell(m, "sig_trav", "DOM")["open"]
    # simulation : la Partie 2 produit maintenant les nouveaux champs
    fld = pd.read_csv(p.fields_csv, dtype=object, encoding="utf-8-sig")
    base = fld[fld["FICHIER"] == f].iloc[0].to_dict()
    new = [("ENGAGEMENT_DOMICILIATION", 1, "DOM_SIGNATURE_TRAVAILLEUR_PRESENTE", "False"),
           ("ENGAGEMENT_DOMICILIATION", 1, "DOM_LU_APPROUVEE_TRAVAILLEUR_PRESENTE", "True"),
           ("CONTRAT_SPECIFIQUE", 3, "CTS_NUMERO_CONTRAT_INSPECTION_TRAVAIL", "CT-2026-0002"),
           ("CONTRAT_SPECIFIQUE", 3, "CTS_DATE_VISA_INSPECTION_TRAVAIL", "05/01/2026")]
    rows = [dict(base, TYPE_DOCUMENT=d, PAGE=pg, CHAMP=c, RAW=v, NORMALIZED=v, DQ_STATUS_FIELD="OK") for d, pg, c, v in new]
    pd.concat([fld, pd.DataFrame(rows)]).to_csv(p.fields_csv, index=False, encoding="utf-8-sig")
    m = _mx(p)[f]
    assert not cell(m, "sig_trav", "DOM")["a_extraire"] and cell(m, "sig_trav", "DOM")["bloq"]   # signature absente
    assert cell(m, "lu_approuve", "DOM")["text"] == "Présente"
    assert cell(m, "contrat", "CTS")["open"]                       # CT-2026-0002 ≠ n° contrat du DOM
    assert cell(m, "date_visa", "CTS")["text"] == "05/01/2026"
    # un autre dossier sans ces champs n'affiche plus « à extraire » mais « — » (champ connu, absent de la pièce)
    assert not cell(_mx(p)["DEMIR_ALI_CNT.pdf"], "sig_trav", "DOM")["a_extraire"]


def test_partie2_attend_partie1(env, monkeypatch):
    p, s = env
    from tlone import jobs, preflight
    monkeypatch.setattr(jobs.subprocess, "Popen", lambda *a, **k: type("P", (), {"pid": 1})())
    real = preflight.check_step
    monkeypatch.setattr(preflight, "check_step", lambda s_, step: [])
    jobs.launch(s, "dom_extract", 1, "testeur")
    monkeypatch.setattr(preflight, "check_step", real)
    ko = {r["contrôle"] for r in preflight.blocking(preflight.check_step(s, "dom_post"))}
    assert "Extraction DOM en cours" in ko


# ------------------------------------------------------------ worker GPU (modèle chargé une fois)
def test_worker_modele_charge_une_fois(env, tmp_path, monkeypatch):
    import dataclasses
    import sys as _sys
    import nbformat
    from tlone import jobs, preflight, worker
    p, s = env
    nb = nbformat.v4.new_notebook()
    nb.cells = [
        nbformat.v4.new_code_cell("MAX_PDFS = None\nfrom pathlib import Path\n"
                                  "INPUT_DIR = Path('/mnt/data/domiciliations_in')\n"
                                  "pdfs = sorted(INPUT_DIR.glob('*.pdf'))\n"
                                  "if MAX_PDFS is not None:\n    pdfs = pdfs[:int(MAX_PDFS)]"),
        nbformat.v4.new_code_cell("# simule AutoModel.from_pretrained(...)\nLOADS = globals().get('LOADS', 0) + 1\nCOUNT = 0"),
        nbformat.v4.new_code_cell("if not pdfs:\n    print('rien')\nelse:\n    COUNT += 1\n"
                                  "    (INPUT_DIR.parent / 'out.txt').write_text(f'{LOADS} {COUNT} {len(pdfs)}')"),
    ]
    fake = tmp_path / "fake_part1.ipynb"
    nbformat.write(nb, str(fake))
    s = dataclasses.replace(s, notebooks={**s.notebooks, "dom_extract": fake}, run_mode="worker")
    # sans worker : lancement bloqué
    assert any(r["contrôle"] == "Worker GPU" for r in preflight.blocking(preflight.check_step(s, "dom_extract")))
    w = worker.Worker(s, ["dom_extract"], pythons={"dom_extract": _sys.executable})
    w.beat()
    assert worker.worker_state(s)["state"] == "ATTENTE"
    monkeypatch.setattr(preflight, "check_step", lambda s_, step: [])
    r1 = jobs.launch(s, "dom_extract", 1, "testeur")
    w.loop(once=True)
    assert (p.data_root / "out.txt").read_text() == "1 1 1"
    r2 = jobs.launch(s, "dom_extract", None, "testeur")
    w.loop(once=True)
    assert (p.data_root / "out.txt").read_text() == "1 2 4"   # modèle chargé 1 fois, 2 traitements, 4 PDF
    runs = {r["run_id"]: r for r in jobs.list_runs(s)}
    assert runs[r1["run_id"]]["status"] == runs[r2["run_id"]]["status"] == "SUCCEEDED"
    w.stop()
    assert worker.worker_state(s) is None


def test_worker_reconnait_les_vrais_notebooks(env):
    """Sur les vrais notebooks : cellule de config, chargement du modèle et traitement bien repérés."""
    import sys as _sys
    from tlone import worker
    _, s = env
    for step, n_proc in (("dom_extract", 1), ("ov_extract", 2)):
        w = worker.WarmNotebook(s, step, _sys.executable)
        model = [i for i, c in enumerate(w.nb.cells) if c.cell_type == "code" and worker.MODEL_MARKER in c.source]
        assert model and w.cfg < model[0] < min(w.process), step     # config -> modèle -> traitement
        assert len(w.process) == n_proc and "GLOB_PDF_CASE_INSENSITIVE" in w.patches
        assert "MAX_PDFS" in w.nb.cells[w.cfg].source or step == "ov_extract"
        if step == "ov_extract":       # v27 : la boucle des PDF fait partie du traitement ; INPUT_DIR redirigé
            assert any("pdf_path in enumerate(pdfs" in w.nb.cells[i].source for i in w.process)
            assert "INPUT_DIR=transferts_in" in w.patches and "NO_PIP_INSTALL" in w.patches


def test_worker_signal_independant_du_fuseau(env):
    """L'App et le workspace peuvent avoir des fuseaux différents : l'âge du signal vient de l'epoch."""
    import json as _json
    import time as _time
    from tlone import worker
    _, s = env
    hb = worker.heartbeat_path(s)
    hb.parent.mkdir(parents=True, exist_ok=True)
    hb.write_text(_json.dumps({"ts": "2026-10-18T10:34:00+05:00", "epoch": _time.time(), "host": "ws",
                               "state": "ATTENTE", "current": None, "warm": {}}))
    assert worker.worker_state(s)["host"] == "ws"
    hb.write_text(_json.dumps({"ts": "x", "epoch": _time.time() - 600, "host": "ws", "state": "ATTENTE"}))
    assert worker.worker_state(s) is None
    assert "il y a 600 s" in worker.worker_absent_text(s)
    hb.unlink()
    assert worker.worker_absent_text(s).startswith("non démarré")


def test_domref_recontrole_apres_lecture_ctr_vide():
    """Partie 2 « n° illisible » (CTR vide au moment du run) : l'App recompare avec la valeur CTR courante."""
    from tlone import matrice as mx
    def cells(ctr, status="NUMERO_DOMICILIATION_NOT_NORMALIZABLE", cands=("271901|2019.3|40|002031|DZD",)):
        return {"CTR": {"value": ctr, "issues": []},
                "CARTHAGO": {"value": cands[0], "carthago_status": status, "candidats": list(cands),
                             "issues": [("bloq", "Carthago : n° de domiciliation illisible")]}}
    c = cells("27-19-01 | 2019.3 | 40 | 002031 | DZD")          # n° du contrat = Carthago
    mx._domref_recalcul(c)
    assert c["CARTHAGO"]["issues"] == [] and c["CARTHAGO"]["recalcul"]
    c = cells(None)                                             # contrat toujours vide
    mx._domref_recalcul(c)
    assert "contrat de travail (CTR)" in c["CARTHAGO"]["issues"][0][1]
    c = cells("271901201934000999DZD", status="NO_EXACT_DOMICILIATION_REFERENCE_MATCH")
    mx._domref_recalcul(c)
    assert c["CARTHAGO"]["issues"][0][0] == "bloq" and "Aucune domiciliation" in c["CARTHAGO"]["issues"][0][1]
    c = cells("x", status="CLIENT_NOT_FOUND")                   # autres statuts : inchangés
    mx._domref_recalcul(c)
    assert c["CARTHAGO"]["issues"][0][1].startswith("Carthago :")


def test_preflight_worker_utilise_l_environnement_du_worker(env, monkeypatch):
    """En mode worker, les versions contrôlées sont celles du workspace du worker, pas du conteneur de l'App."""
    import dataclasses
    import json as _json
    import time as _time
    from tlone import preflight, worker
    p, s = env
    s = dataclasses.replace(s, run_mode="worker")
    monkeypatch.setattr(preflight, "probe_python", lambda py, mods: {m: "5.20.1" for m in mods})  # conteneur App
    mods = preflight.STEP_MODULES["ov_extract"]
    hb = worker.heartbeat_path(s)
    hb.parent.mkdir(parents=True, exist_ok=True)
    hb.write_text(_json.dumps({"epoch": _time.time(), "ts": "x", "host": "ws", "state": "ATTENTE",
                               "steps": ["dom_extract", "dom_post"], "warm": {},
                               "env": {"ov_extract": {"python": "/opt/tf4/python",
                                                      "versions": {**{m: "ok" for m in mods},
                                                                   "transformers": "4.57.0"}}}}))
    rows = {r["contrôle"]: r for r in preflight.check_step(s, "ov_extract")}
    assert rows["transformers ov_extract (worker)"]["niveau"] == preflight.OK
    assert "transformers ov_extract" not in rows                      # pas de sondage local
    assert rows["Worker ov_extract"]["niveau"] == preflight.WARN      # modèle non préchargé : alerte seulement
    assert not preflight.blocking(list(rows.values())) or all(
        r["contrôle"] in ("PDF OV à traiter",) for r in preflight.blocking(list(rows.values())))


def test_worker_echec_chargement_n_arrete_pas_le_worker(env):
    """Un notebook qui échoue au préchargement (ex. ImportError transformers) n'arrête pas le worker."""
    import dataclasses
    import sys as _sys
    import nbformat
    from tlone import preflight, worker
    p, s = env
    nb = nbformat.v4.new_notebook()
    nb.cells = [nbformat.v4.new_code_cell("from transformers_inexistant import AutoModelForVision2Seq"),
                nbformat.v4.new_code_cell("pdfs = sorted([])"), nbformat.v4.new_code_cell("for pdf_path in pdfs: pass")]
    bad = p.data_root / "bad.ipynb"
    nbformat.write(nb, str(bad))
    s = dataclasses.replace(s, notebooks={**s.notebooks, "ov_extract": bad}, run_mode="worker")
    w = worker.Worker(s, ["ov_extract"], prewarm=True, pythons={"ov_extract": _sys.executable})
    w.loop(once=True)                                   # ne lève pas
    assert "ov_extract" in w.errors and "ov_extract" not in w.warm
    w.beat()
    rows = {r["contrôle"]: r for r in preflight.check_step(s, "ov_extract")}
    assert rows["Worker ov_extract"]["niveau"] == preflight.KO
    w.stop()


def test_patch_desactive_pip_install():
    out, applied = nbrunner.patch_source("%pip install -q pymupdf transformers\nif True:\n    !pip install x\nprint(1)",
                                         "ov_extract", "", None, None)
    assert "NO_PIP_INSTALL" in applied and "%pip" not in out.split("#")[0]
    compile(out, "cellule", "exec")


def _v27_excel(path):
    """Classeur produit par la cellule « Export Excel » du vrai notebook v27 (code exécuté tel quel)."""
    import json as _json
    nb = _json.load(open(Path(__file__).resolve().parents[1] / "notebooks" / "pipeline_ocr_vl27.ipynb", encoding="utf-8"))
    src = next("".join(c["source"]) for c in nb["cells"] if "def create_excel" in "".join(c["source"]))
    ns: dict = {}
    exec("from pathlib import Path\nfrom collections import defaultdict\nfrom openpyxl import Workbook\n"
         "from openpyxl.styles import Font, PatternFill, Alignment, Border, Side\n"
         "from openpyxl.utils import get_column_letter\n" + src, ns)
    row = {"fichier": "02700731010825900165-202603P1.pdf", "date_traitement": "2026-10-19 01:00:00",
           "temps_total_s": 12.0, "tokens_total": 1000, "pages_trouvees": "ANNEXE_I, ANNEXE_II, BULLETIN, OV",
           "pages_manquantes": None, "anomalies": None,
           "OV": {"monnaie": "DZD", "montant_chiffres": 312255.0, "montant_lettres": "trois cent douze mille deux cent cinquante-cinq dinars",
                  "periode": "MARS 2026 P1", "mois": "MARS", "annee": "2026", "tranche": "P1", "complement_ov": None,
                  "date_demande": "05/03/2026", "compte_donneur_ordre": "02700731010825900165",
                  "nature_paiement_autre_libelle": "EUR", "beneficiaire_nom": "KAYA AHMET",
                  "beneficiaire_compte": "TR33", "beneficiaire_adresse": "X",
                  "code_swift_banque_beneficiaire": "TEBUTRIS", "nom_banque_beneficiaire": "TEB"},
           "ANNEXE_I": {"nom_prenom_employe": "KAYA AHMET", "compte_bancaire_local": "02700731010825900165"},
           "ANNEXE_II": {"mois_transfert": "MARS 2026", "nom_prenom_travailleur": "KAYA AHMET",
                         "compte_bancaire_local": "02700731010825900165", "salaire_mensuel": 300000.0,
                         "part_transferable": 312255.0, "numero_domiciliation": "999"},
           "BULLETIN": {"nom_prenom_salarie": "KAYA AHMET", "periode_bulletin": "MARS 2026", "net_a_payer": 340000.0}}
    ns["create_excel"](path, [row])


def test_lecture_classeur_v27_et_controles_retires(tmp_path):
    """v27 : libellés différents -> mêmes clés que v26 ; T16 (n° DOM) et T08 (net/salaire) ne sont plus contrôlés."""
    _v27_excel(tmp_path / "audit_transferts_20261019_0100.xlsx")
    a = transfers.read_audit_xlsx(tmp_path / "audit_transferts_20261019_0100.xlsx")
    for col in ("OV.compte_debiteur", "OV.swift", "OV.banque_nom", "OV.periode", "OV.mois", "ANNEXE_II.nom_prenom",
                "ANNEXE_II.compte_bancaire", "ANNEXE_II.numero_domiciliation", "BULLETIN.nom_prenom",
                "BULLETIN.mois_bulletin", "ANNEXE_I.nom_prenom"):
        assert col in a.columns, col
    s = load_settings()
    synth, rows = transfers.control_dossier(a.iloc[0].to_dict(), pd.DataFrame(), pd.DataFrame(columns=["PERIODE_ID", "DOM_ID", "FICHIER", "_MONTANT"]),
                                            pd.DataFrame(), pd.DataFrame(), s)
    codes = {r["CODE"]: r for r in rows}
    assert "T16" not in codes and "T08" not in codes
    assert codes["T07"]["RESULTAT"] == "OK"               # MARS 2026 partout, malgré « P1 »
    assert codes["T03"]["RESULTAT"] == "OK"


def test_periode_annee_deux_chiffres():
    assert transfers.parse_periode("JANV 26 P2") == ("2026-01", "P2")
    assert transfers.parse_periode("MARS 2026") == ("2026-03", None)


# ------------------------------------------------------------------ vigilance
def _modele_vigilance(path):
    """Réplique du modèle « CHECK LIST TRANSFERT LIBRE » (mêmes libellés / lignes que le modèle BNP fourni)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Feuil1"
    ws["A2"] = "CHECK LIST TRANSFERT LIBRE"
    for r, lab in ((5, "Client:"), (6, "Réf Dossier:"), (7, "Circuit:"), (8, "Type de dossier:")):
        ws.cell(r, 1, lab)
    for c, lab in zip("ABCDE", ["Contrôle", "Type", "Valeur", "HIT / No HIT / KO", "COMMENTAIRES"]):
        ws[f"{c}10"] = lab
    ws["A11"] = "VIGILANCE"
    ws.merge_cells("A11:E11")
    rub = ["Banque du bénéficiaire", "Pays de la banque du bénéficiaire", "Banque intermédiaire",
           "Pays de la banque intermédiaire", "Bénéficiaire", "Pays du bénéficiaire", "Pays d’origine et provenance",
           "Nom du fabricant", "Pays du fabricant", "Compagnie maritime ou aérienne", "Navire", "Expéditeur sur le BL",
           "Agent signataire P/C compagnie maritime", "Transitaire", "Transporteur LVI", "Port ou aéroport d’embarquement",
           "Port ou aéroport de débarquement", "Lieu de prise en charge sur le BL", "Lieu établissement document transport",
           "Lieu de travail", "Divers individus sur les documents", "Divers entités sur les documents", "Divers lieux"]
    for i, lab in enumerate(rub):
        ws.cell(12 + i, 1, lab)
        ws.cell(12 + i, 2, "I" if "individus" in lab else "C")
    ws["A37"] = "Valideur :"
    wb.save(path)


def _retour(path_in, results: dict, individus=None) -> bytes:
    import io
    from openpyxl import load_workbook
    wb = load_workbook(path_in)
    ws = wb.worksheets[0]
    for row in range(12, 35):
        ws[f"D{row}"] = "OK ==> RAS" if ws[f"C{row}"].value else "Rubrique non renseignée"
    if ws["C32"].value:
        ws["D32"] = "\n".join(["OK ==> RAS"] * len(str(ws["C32"].value).split("\n")))
    for cell, val in results.items():
        ws[cell] = val
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_vigilance_regroupement_preparation_retours_decisions(tmp_path):
    import io
    import zipfile
    from openpyxl import load_workbook
    from tlone import vigilance as v
    tpl = tmp_path / "modele.xlsx"
    _modele_vigilance(tpl)
    assert v.check_template(tpl) == []
    ds = [dict(fichier="f1", ref_client="07310108547", client="KAYA AHMET", beneficiaire="AHMET KAYA", iban="TR33 0006",
               swift="TEBUTRIS", banque="TEB", pays_banque="TURQUIE", individus=["MEHMET KAYA", "AYSE KAYA"]),
          dict(fichier="f2", ref_client="07310108547", client="KAYA AHMET", beneficiaire="AHMET KAYA", iban="TR330006",
               swift="TEBUTRIS", banque="TEB", pays_banque="TURQUIE", individus=["AYSE KAYA", "SIGNATAIRE X"]),
          dict(fichier="f3", ref_client="07310262001", client="DEMIR ALI", beneficiaire="DEMIR FATMA", iban="TR99",
               swift="ZIRATRIS", banque="ZIRAAT", pays_banque="TURQUIE", individus=[]),
          dict(fichier="f4", ref_client="07310262001", client="DEMIR ALI", beneficiaire="DEMIR ALI", iban="TR11",
               swift="ZIRATRIS", banque="ZIRAAT", pays_banque="TURQUIE", individus=[])]
    units = v.build_units(ds)
    by = {tuple(u["fichiers"]): u for u in units}
    assert len(units) == 3                                     # KAYA : même bénéficiaire -> 1 fichier ; DEMIR : 2
    assert by[("f1", "f2")]["beneficiaire"] == ""              # bénéficiaire = client (ordre inversé) : omis
    assert by[("f1", "f2")]["individus"] == ["MEHMET KAYA", "AYSE KAYA", "SIGNATAIRE X"]
    assert by[("f3",)]["beneficiaire"] == "DEMIR FATMA"
    m = v.prepare(units, tpl, tmp_path / "lot", "TOSYALI IRON STEEL INDUSTRY ALGERIE")
    ws = load_workbook(tmp_path / "lot" / by[("f1", "f2")]["fichier"]).worksheets[0]
    assert (ws["C5"].value, ws["C6"].value, ws["C7"].value, ws["C8"].value) == ("KAYA AHMET", "07310 108547", "TL", "Création")
    assert ws["C12"].value == "TEB" and ws["C13"].value == "TURQUIE" and ws["C31"].value.startswith("TOSYALI")
    assert ws["C32"].value == "MEHMET KAYA\nAYSE KAYA\nSIGNATAIRE X" and ws["C16"].value is None
    assert zipfile.ZipFile(m["zip"]).namelist() == [u["fichier"] for u in units]

    jp = tmp_path / "journal.jsonl"
    v.append(jp, "u", "ENVOI", lot=m["lot"], unites=units)
    st = v.statuses(v.read_journal(jp))
    assert {f: x["statut"] for f, x in st.items()} == dict.fromkeys(["f1", "f2", "f3", "f4"], "EN_ATTENTE")
    ret = [("Recherche_Vigilance_Data_A.xlsx", _retour(tmp_path / "lot" / by[("f1", "f2")]["fichier"], {})),
           ("Recherche_Vigilance_Data_B.xlsx", _retour(tmp_path / "lot" / by[("f3",)]["fichier"], {"D16": "OK ==> HIT"})),
           ("Recherche_Vigilance_Data_C.xlsx", _retour(tmp_path / "lot" / by[("f4",)]["fichier"],
                                                       {"D12": "KO ==> La Plateforme Vigilance est indisponible"}))]
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as zz:
        for n, b in ret:
            zz.writestr(n, b)
    res = v.import_returns([("retours.zip", z.getvalue())], v.read_journal(jp), jp, "u", tmp_path / "res")
    assert set(res["synthese"]["Statut"]) == {"RAS", "HIT", "KO_A_REFAIRE"} and res["controles"].empty
    st = v.statuses(v.read_journal(jp))
    assert [st[f]["statut"] for f in ("f1", "f2", "f3", "f4")] == ["RAS", "RAS", "HIT", "KO_A_REFAIRE"]
    assert v.eligible(st) == {"f1", "f2"}
    dec = pd.DataFrame({"FICHIER": ["f1", "f2", "f3", "f4", "f5"], "DECISION": ["VALIDE"] * 5})
    assert v.candidates(dec, st) == ["f4", "f5"]               # KO à refaire + jamais envoyé
    with pytest.raises(ValueError):
        v.decide(jp, "u", ["f3"], "LEVEE", "  ")               # motif obligatoire
    v.decide(jp, "u", ["f3"], "LEVEE", "Homonyme : date de naissance différente")
    assert v.statuses(v.read_journal(jp))["f3"]["statut"] == "HIT_LEVE"
    assert v.eligible(v.statuses(v.read_journal(jp))) == {"f1", "f2", "f3"}


def test_vigilance_classement_et_controles():
    from tlone import vigilance as v
    assert v.classify("OK ==> RAS")[0] == "RAS" and v.classify("OK ==> HIT")[0] == "HIT"
    assert v.classify("KO ==> La Plateforme Vigilance est indisponible")[0] == "KO"
    assert v.classify("Rubrique non renseingée")[0] == "INFO" and v.classify("")[0] == "MANQUANT"
    assert v.classify("peut-être")[0] == "INCONNU"
    u = {"client": "A B", "individus": ["X", "Y"]}
    ret = {"client": "A B", "individus_row": 32, "rubriques": [
        {"ligne": 32, "rubrique": "Divers individus", "valeur": "X\nY", "resultat": "OK ==> RAS", "commentaire": ""}]}
    statut, _, ctl = v.analyse(u, ret)
    assert statut == "A_VERIFIER" and "Décalage" in ctl[0]     # 2 individus, 1 résultat : appariement à vérifier
    assert v.format_ref("07310108547") == "07310 108547" and v.pays_banque("TEBUTRIS", None) == "TURQUIE"
    assert v.safe("=HYPERLINK(1)") == "HYPERLINK(1)"


def test_registre_exige_vigilance():
    import tempfile
    sy = pd.DataFrame([{"FICHIER": f, "PERIODE_ID": f"P{f}", "MONTANT_OV": 1.0} for f in ("a", "b")])
    dec = pd.DataFrame({"FICHIER": ["a", "b"], "DECISION": ["VALIDE", "VALIDE"], "COMMENTAIRE": ["", ""]})
    with tempfile.TemporaryDirectory() as d:
        n = transfers.register(Path(d) / "reg.csv", sy, dec, "u", vigilance={"a": "RAS", "b": "HIT"})
        reg = transfers.load_registre(Path(d) / "reg.csv")
    assert n == 1 and list(reg["FICHIER"]) == ["a"] and list(reg["STATUT_VIGILANCE"]) == ["RAS"]


def test_vigilance_analyse_hit_dans_tlone(tmp_path):
    """HIT commenté dans l'app : non avéré -> TL libéré + fiche annotée ; avéré -> alerte confirmée."""
    from openpyxl import load_workbook
    from tlone import vigilance as v
    tpl = tmp_path / "modele.xlsx"
    _modele_vigilance(tpl)
    ds = [dict(fichier="f1", ref_client="07310108547", client="KAYA AHMET", beneficiaire="KAYA AHMET", iban="TR1",
               swift="TEBUTRIS", banque="TEB", pays_banque="TURQUIE", individus=["MEHMET KAYA", "AYSE KAYA"]),
          dict(fichier="f2", ref_client="07310262001", client="DEMIR ALI", beneficiaire="DEMIR ALI", iban="TR2",
               swift="ZIRATRIS", banque="ZIRAAT", pays_banque="TURQUIE", individus=[])]
    units = v.build_units(ds)
    v.prepare(units, tpl, tmp_path / "lot", "TOSYALI")
    jp = tmp_path / "j.jsonl"
    v.append(jp, "u", "ENVOI", lot="L1", unites=units)
    u1 = next(u for u in units if u["fichiers"] == ["f1"])
    u2 = next(u for u in units if u["fichiers"] == ["f2"])
    r1 = _retour(tmp_path / "lot" / u1["fichier"], {"D32": "OK ==> RAS\nOK ==> HIT"})
    r2 = _retour(tmp_path / "lot" / u2["fichier"], {"D12": "OK ==> HIT"})
    v.import_returns([(u1["fichier"], r1), (u2["fichier"], r2)], v.read_journal(jp), jp, "u", tmp_path / "res")
    j = v.read_journal(jp)
    ret1 = next(e for e in j if e["action"] == "RETOUR" and e["uid"] == u1["uid"])
    assert Path(ret1["chemin"]).exists() and ret1["points"][0][2] == "AYSE KAYA"   # HIT sur la mère (rang 2)
    pt = ret1["points"][0]
    with pytest.raises(ValueError):                                                  # commentaire obligatoire
        v.analyse_points(jp, "u", u1["uid"], ["f1"], [{"rubrique": pt[0], "rang": pt[1], "valeur": pt[2],
                                                       "verdict": "NON_AVERE", "commentaire": ""}], None, tmp_path)
    ev = v.analyse_points(jp, "u", u1["uid"], ["f1"],
                          [{"rubrique": pt[0], "rang": pt[1], "valeur": pt[2], "verdict": "NON_AVERE",
                            "commentaire": "Homonyme : date de naissance différente"}],
                          Path(ret1["chemin"]), tmp_path / "analyses")
    assert ev["decision"] == "LEVEE"
    ws = load_workbook(ev["fiche_annotee"]).worksheets[0]
    assert "NON AVÉRÉ — AYSE KAYA : Homonyme" in ws["E32"].value
    ret2 = next(e for e in v.read_journal(jp) if e["action"] == "RETOUR" and e["uid"] == u2["uid"])
    p2 = ret2["points"][0]
    ev2 = v.analyse_points(jp, "u", u2["uid"], ["f2"], [{"rubrique": p2[0], "rang": 1, "valeur": p2[2],
                                                         "verdict": "AVERE", "commentaire": "Banque sous sanctions"}],
                           Path(ret2["chemin"]), tmp_path / "analyses")
    assert ev2["decision"] == "CONFIRMEE"
    st = v.statuses(v.read_journal(jp))
    assert st["f1"]["statut"] == "HIT_LEVE" and st["f2"]["statut"] == "HIT_CONFIRME"
    assert v.eligible(st) == {"f1"}


def test_preflight_worker_autre_dossier(env):
    """Worker lancé depuis un autre dossier de version (ex. tlone_app_v2 alors que l'App tourne en v3) : bloquant."""
    import dataclasses
    import json as _json
    import time as _time
    from tlone import preflight, worker
    _, s = env
    s = dataclasses.replace(s, run_mode="worker")
    hb = worker.heartbeat_path(s)
    hb.parent.mkdir(parents=True, exist_ok=True)
    hb.write_text(_json.dumps({"epoch": _time.time(), "ts": "x", "host": "ws", "state": "ATTENTE", "warm": {},
                               "steps": ["ov_extract"], "app_dir": "/mnt/code/tlone_app_v2", "env": {}}))
    rows = {r["contrôle"]: r for r in preflight.check_step(s, "ov_extract")}
    assert rows["Dossier du worker"]["niveau"] == preflight.KO
