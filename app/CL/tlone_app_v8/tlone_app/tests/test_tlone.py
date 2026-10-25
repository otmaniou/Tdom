import sys
from datetime import date
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tests import fixtures  # noqa: E402
from tlone import domvalid, nbrunner, planning, transfers  # noqa: E402
from tlone.config import load_settings  # noqa: E402
from tlone.montant_lettres import lettres_to_number  # noqa: E402


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("TLONE_DATA_ROOT", str(tmp_path))
    monkeypatch.setenv("TLONE_RUN_MODE", "local")
    p = fixtures.build(tmp_path)
    return p, load_settings()


# ------------------------------------------------------------ montant en lettres
@pytest.mark.parametrize("txt,val", [
    ("UN MILLION DEUX CENT SOIXANTE ET UN MILLE SIX CENTS DINARS ALGERIENS", 1261600),
    ("TROIS CENT DOUZE MILLE DEUX CENT CINQUANTE CINQ DINARS", 312255),
    ("CENT QUATRE VINGT MILLE DINARS", 180000),
    ("QUATRE-VINGT-DIX-NEUF MILLE NEUF CENT QUATRE-VINGT-DIX-NEUF DA", 99999),
    ("DEUX CENT MILLE DINARS ET 50 CTS", 200000.50),
    ("SOIXANTE DIX SEPT MILLE DINARS ALGÉRIENS ET QUARANTE CENTIMES", 77000.40),
    ("MILLE DINARS", 1000),
])
def test_lettres(txt, val):
    assert lettres_to_number(txt) == pytest.approx(val)


def test_lettres_illisible():
    assert lettres_to_number("TROIS CENT BLURP MILLE") is None


# ------------------------------------------------------------ runner
def test_patch_source_root_and_glob():
    src = "INPUT_DIR = Path('/mnt/data/domiciliations_in')\nMAX_PDFS = None\npdfs = sorted(INPUT_DIR.glob('*.pdf'))\n" \
          "assert RAW_JSON_DIR == Path('/mnt/data/x')\nMODEL_PATH = '/domino/edv/old'\n"
    out, applied = nbrunner.patch_source(src, "dom_extract", "/domino/datasets/local/tlone", 5, "/models/new")
    assert "/mnt/data" not in out and "/domino/datasets/local/tlone/domiciliations_in" in out
    assert "MAX_PDFS = 5" in out and "suffix.lower() == '.pdf'" in out
    assert 'MODEL_PATH = "/models/new"' in out
    assert any(a.startswith("DATA_ROOT") for a in applied)


def test_patch_default_root_untouched():
    src = "OUTPUT_ROOT = Path('/mnt/data/document_pipeline/V15_A')\n"
    out, applied = nbrunner.patch_source(src, "dom_post", "/mnt/data", None, None)
    assert out == src and applied == []


def test_real_notebooks_patchable():
    s = load_settings()
    for step, nb in s.notebooks.items():
        prepared, applied = nbrunner.prepare_notebook(nb, step, "/tmp/x", 3, None)
        assert any(a.startswith("DATA_ROOT") for a in applied), step
        code = "\n".join(c.source for c in prepared.cells if c.cell_type == "code")
        assert "/mnt/data" not in code
        if step in ("dom_extract", "ov_extract"):
            assert "GLOB_PDF_CASE_INSENSITIVE" in applied
        assert any(a.startswith(("MAX_", "pdfs[")) for a in applied), (step, applied)


# ------------------------------------------------------------ module 1
def test_domvalid_roundtrip(env):
    p, _ = env
    man = domvalid.load_manual(p.validation_xlsx)
    assert list(man["_ROW"]) == [2, 3, 4, 5]
    mt = p.validation_xlsx.stat().st_mtime
    ed = man.copy()
    ed.loc[ed["FICHIER"] == "CELIK_MEHMET_CNT.pdf", ["DECISION_MANUELLE", "VALEUR_VALIDEE"]] = ["CORRIGEE", ""]
    assert domvalid.check_decisions(ed)  # CORRIGEE sans valeur bloqué
    ed.loc[ed["FICHIER"] == "CELIK_MEHMET_CNT.pdf", "VALEUR_VALIDEE"] = "02700731010830000161"
    assert not domvalid.check_decisions(ed)
    res = domvalid.save_decisions(p.validation_xlsx, ed, expected_mtime=mt)
    assert len(res["changes"]) == 2 and Path(res["backup"]).exists()
    re = domvalid.load_manual(p.validation_xlsx)
    assert re.loc[re["FICHIER"] == "CELIK_MEHMET_CNT.pdf", "DECISION_MANUELLE"].iloc[0] == "CORRIGEE"
    # Les autres feuilles sont conservées
    assert "DOSSIERS_A_VALIDER" in domvalid.sheet_names(p.validation_xlsx)
    with pytest.raises(RuntimeError):
        domvalid.save_decisions(p.validation_xlsx, ed, expected_mtime=mt)  # fichier modifié entre-temps


def test_progress(env):
    p, _ = env
    pr = domvalid.progress(domvalid.load_manual(p.validation_xlsx)).set_index("FICHIER")
    assert pr.loc["DEMIR_ALI_CNT.pdf", "STATUT_VALIDATION"] == "REJETE"
    assert pr.loc["CELIK_MEHMET_CNT.pdf", "RESTANT"] == 1


# ------------------------------------------------------------ module 2
def _planning(p, excl=False):
    dos = domvalid.load_sheet(p.validation_xlsx, "DOSSIERS_A_VALIDER")
    man = domvalid.load_sheet(p.validation_xlsx, "A_VALIDER_MANUELLEMENT")
    pa = pd.read_excel(p.param_aug_xlsx, dtype=object)
    return planning.build_planning(dos, man, pa, exclure_validation_incomplete=excl)


def test_planning_rules(env):
    p, _ = env
    r = _planning(p)
    pl, rej = r["Planning_TL"], r["REJETS_PLANNING"].set_index("FICHIER")
    assert r["_dq"] == []
    assert rej.loc["DEMIR_ALI_CNT.pdf", "MOTIF"] == "REJETE_EN_PARTIE_2A"
    assert "COMPTE_FINAL_ABSENT" in rej.loc["CELIK_MEHMET_CNT.pdf", "MOTIF"]
    kaya = pl[pl["NOM_TRAVAILLEUR"] == "KAYA AHMET"]
    assert kaya["DOM_ID"].nunique() == 1  # contrat + augmentation sur la même DOM
    jan = kaya[kaya["MOIS_BASE"] == "2026-01"].iloc[0]
    assert jan["PERIODE_TL"] == "2026-01P2" and jan["NB_JOURS_SEGMENT"] == 29
    assert jan["MONTANT_MAX_THEORIQUE"] == round(312255 * 29 / 31, 2)
    mar = kaya[kaya["MOIS_BASE"] == "2026-03"].iloc[0]
    apr = kaya[kaya["MOIS_BASE"] == "2026-04"].iloc[0]
    assert mar["TYPE_VERSION"] == "NOUVEAU_CONTRAT" and mar["MONTANT_MAX_THEORIQUE"] == 312255
    assert apr["TYPE_VERSION"] == "AUGMENTATION" and apr["NUMERO_AUGMENTATION"] == 1
    assert apr["MONTANT_MAX_THEORIQUE"] == 332500
    last = kaya.iloc[-1]
    assert last["PERIODE_TL"] == "2028-01P1" and last["NB_JOURS_SEGMENT"] == 2


def test_planning_matches_notebook(env, monkeypatch):
    """Exécute le VRAI notebook Planning_TL_V15A sur le même jeu et compare."""
    p, _ = env
    import nbformat
    from nbclient import NotebookClient
    nb = nbformat.read(str(Path(__file__).resolve().parent.parent / "notebooks" / "Planning_TL_V15A.ipynb"), 4)
    for c in nb.cells:
        if c.cell_type == "code":
            c.source = c.source.replace("/mnt/data", str(p.data_root))
    NotebookClient(nb, timeout=120, kernel_name="python3").execute()
    ref = pd.read_csv(p.planning_csv, dtype=str).fillna("")
    mine = _planning(p)["Planning_TL"].astype(str).replace({"None": "", "nan": ""})
    mine = mine[[c for c in mine.columns if c in ref.columns]]      # colonnes TL-One ajoutées hors comparaison
    ref = ref[mine.columns]
    for c in ("MONTANT_MAX_THEORIQUE", "COEFFICIENT_PRORATA", "PART_TRANSFERABLE_CTS"):
        assert (ref[c].astype(float).round(6) == mine[c].replace("", "nan").astype(float).round(6)).all(), c
    for c in ("PERIODE_ID", "DOM_ID", "VERSION_ID", "PERIODE_TL", "TYPE_VERSION", "DATE_DEBUT_SEGMENT"):
        assert list(ref[c]) == list(mine[c]), c


# ------------------------------------------------------------ module 3
def test_transfer_controls(env):
    p, s = env
    r = _planning(p)
    planning.export(r, p.planning_xlsx, p.planning_csv)
    pl = planning.load_planning(p.planning_csv)
    audit = transfers.read_audit_xlsx(transfers.latest_audit(p.tr_out)[0])
    assert len(audit) == 4 and "OV.montant_chiffres" in audit.columns
    atf = transfers.load_referential(p.atf_xlsx, transfers.ATF_COLUMNS)
    aut = transfers.load_referential(p.autorisations_xlsx, transfers.AUT_COLUMNS)
    reg = transfers.load_registre(p.registre_csv)
    sy, ch = transfers.run_controls(audit, pl, reg, atf, aut, s, today=date(2026, 4, 10))
    sy = sy.set_index("FICHIER")
    mar = sy.loc["07310-108259-202603.pdf"]
    assert mar["STATUT_CONTROLE"] == "CONFORME", ch[ch.FICHIER == "07310-108259-202603.pdf"].to_dict("records")
    assert mar["PERIODE_TL"] == "2026-03" and mar["PROVISION_ATLAS"] == 317255
    feb = sy.loc["07310-108259-202602.pdf"]
    assert feb["STATUT_CONTROLE"] == "NON_CONFORME" and "T09" in feb["MOTIFS"] and "T10" in feb["MOTIFS"]
    jan = sy.loc["07310-108259-202601P2.pdf"]
    assert jan["PERIODE_TL"] == "2026-01P2" and jan["STATUT_CONTROLE"] in ("A_REVOIR", "NON_CONFORME")
    demir = sy.loc["07310-108262-202605.pdf"]
    assert demir["STATUT_CONTROLE"] == "NON_CONFORME" and "T04" in demir["MOTIFS"]

    # Décision + registre + consommation
    ed = sy.reset_index()[["FICHIER", "AUDIT_SOURCE", "STATUT_CONTROLE", "PERIODE_ID"]].copy()
    ed["DECISION"], ed["COMMENTAIRE"] = "", ""
    ed.loc[ed.FICHIER == "07310-108259-202603.pdf", "DECISION"] = "VALIDE"
    ed.loc[ed.FICHIER == "07310-108262-202605.pdf", "DECISION"] = "VALIDE"
    assert transfers.check_decisions(ed)  # non rattaché ne peut être validé
    ed.loc[ed.FICHIER == "07310-108262-202605.pdf", "DECISION"] = "REJETE"
    ed.loc[ed.FICHIER == "07310-108262-202605.pdf", "COMMENTAIRE"] = "DOM rejetée"
    assert not transfers.check_decisions(ed)
    assert transfers.save_decisions(p.decisions_csv, ed, "testeur") == 2
    dec = transfers.load_decisions(p.decisions_csv)
    assert transfers.register(p.registre_csv, sy.reset_index(), dec, "testeur") == 1
    assert transfers.register(p.registre_csv, sy.reset_index(), dec, "testeur") == 0  # idempotent
    reg = transfers.load_registre(p.registre_csv)
    cons = transfers.consumption(pl, reg, today=date(2026, 4, 10)).set_index("PERIODE_TL")
    assert cons.loc["2026-03", "STATUT_PERIODE"] == "TRANSFERE"
    assert cons.loc["2026-02", "STATUT_PERIODE"] == "NON_TRANSFERE"
    # Un 2e OV sur mars est désormais bloqué (déjà transféré)
    sy2, ch2 = transfers.run_controls(audit.iloc[[0]].assign(fichier="BIS-202603.pdf"), pl, reg, atf, aut, s,
                                      today=date(2026, 4, 10))
    assert sy2.iloc[0]["STATUT_CONTROLE"] in ("NON_CONFORME",)


def test_periodes():
    assert transfers.parse_periode("MARS 2026 PART 1") == ("2026-03", "P1")
    assert transfers.parse_periode("FÉVRIER.2026") == ("2026-02", None)
    assert transfers.parse_periode("03/2026") == ("2026-03", None)
    assert transfers.periode_from_filename("02700731011155100123-202510P2.PDF") == ("2025-10", "P2")
    assert transfers.periode_from_filename("07310-108262-202605.pdf") == ("2026-05", None)
    assert transfers.periode_from_filename("07310-109655-042026p2.pdf") == ("2026-04", "P2")
    assert transfers.periode_from_filename("02700731010854700144-202606-OV-PART-2_COKLUK-AHMET.pdf") == ("2026-06", "P2")
    assert transfers.periode_from_filename("02700731010854700144-202607-OV_COKLUK-AHMET.pdf") == ("2026-07", None)
    assert transfers.periode_from_filename("07310-109655-04 p1.pdf") == (None, "P1")
    from tlone.common import siege_racine
    assert siege_racine("02700731010854700144-202607-OV_COKLUK") == "07310108547"
    assert siege_racine("07310-109655-04 p1") == "07310109655"


def test_carry_over(env):
    p, _ = env
    from tlone.common import backup
    prev = backup(p.validation_xlsx)
    # simulate a Part 2 rerun : decisions wiped
    man = domvalid.load_manual(p.validation_xlsx)
    wiped = man.copy()
    for c in domvalid.EDITABLE:
        wiped[c] = ""
    domvalid.save_decisions(p.validation_xlsx, wiped)
    todo = domvalid.carry_over(prev, p.validation_xlsx)
    assert len(todo) == 3  # 3 lignes décidées dans le jeu, la 4e était vide
    domvalid.save_decisions(p.validation_xlsx, todo)
    again = domvalid.load_manual(p.validation_xlsx)
    assert list(again["DECISION_MANUELLE"]) == list(man["DECISION_MANUELLE"])


def test_jobs_local_launch(env, monkeypatch):
    p, s = env
    from tlone import jobs
    from tlone import preflight
    monkeypatch.setattr(jobs.subprocess, "Popen", lambda *a, **k: type("P", (), {"pid": 4242})())
    monkeypatch.setattr(preflight, "check_step", lambda s, step: [])
    rec = jobs.launch(s, "dom_post", 3, "testeur")
    assert rec["backup_validation"] and Path(rec["backup_validation"]).exists()
    runs = jobs.list_runs(s)
    assert runs[0]["run_id"] == rec["run_id"] and runs[0]["status"] == "EN_FILE"


def test_step_runtime(monkeypatch):
    monkeypatch.setenv("TLONE_GPU_TIER", "gpu-commun")
    monkeypatch.setenv("TLONE_DOM_EXTRACT_TIER", "h100-80g")
    monkeypatch.setenv("TLONE_DOM_EXTRACT_ENV", "env-qwen38-tf5")
    monkeypatch.setenv("TLONE_OV_EXTRACT_ENV", "env-qwen25-tf4")
    monkeypatch.setenv("TLONE_OV_EXTRACT_PYTHON", "/opt/venv-qwen25/bin/python")
    s = load_settings()
    a, b, c = s.step_runtime("dom_extract"), s.step_runtime("ov_extract"), s.step_runtime("dom_post")
    assert a["tier"] == "h100-80g" and a["environment"] == "env-qwen38-tf5"
    assert b["tier"] == "gpu-commun" and b["environment"] == "env-qwen25-tf4"
    assert b["python"] == "/opt/venv-qwen25/bin/python"
    assert c["environment"] == "" and c["model"] == ""


def test_domino_payload(env, monkeypatch):
    p, _ = env
    from tlone import jobs
    monkeypatch.setenv("TLONE_RUN_MODE", "domino")
    monkeypatch.setenv("DOMINO_PROJECT_ID", "proj1")
    monkeypatch.setenv("TLONE_OV_EXTRACT_ENV", "env-qwen25")
    monkeypatch.setenv("TLONE_OV_EXTRACT_TIER", "gpu-a10")
    from tlone import preflight
    sent = {}
    monkeypatch.setattr(preflight, "check_step", lambda s, step: [])
    monkeypatch.setattr(jobs, "_domino_request", lambda m, path, payload=None: sent.update(payload) or {"id": "J1"})
    rec = jobs.launch(load_settings(), "ov_extract", 2, "testeur")
    assert sent["environmentId"] == "env-qwen25" and sent["overrideHardwareTierId"] == "gpu-a10"
    assert sent["commandToRun"].endswith(f"ov_extract --run-id {rec['run_id']} --data-root {p.data_root} --max 2")


# ------------------------------------------------------------ contrôle de démarrage
def test_preflight_bloque_sans_pdf_ni_tier(env, monkeypatch):
    p, _ = env
    from tlone import jobs, preflight
    monkeypatch.setenv("TLONE_RUN_MODE", "domino")
    monkeypatch.setenv("DOMINO_PROJECT_ID", "proj1")
    monkeypatch.setenv("DOMINO_API_PROXY", "http://localhost:8899")
    monkeypatch.delenv("TLONE_GPU_TIER", raising=False)
    s = load_settings()
    for f in p.dom_in.glob("*"):
        f.unlink()
    ko = {r["contrôle"] for r in preflight.blocking(preflight.check_step(s, "dom_extract"))}
    assert "PDF DOM à traiter" in ko and "Tier GPU du job" in ko
    called = []
    monkeypatch.setattr(jobs, "_domino_request", lambda *a, **k: called.append(a) or {"id": "X"})
    with pytest.raises(RuntimeError, match="prérequis bloquants"):
        jobs.launch(s, "dom_extract", 2, "testeur")
    assert not called  # aucun job Domino n'est parti


def test_preflight_data_root_et_double_lancement(env, monkeypatch):
    p, s = env
    from tlone import jobs, preflight
    assert not [r for r in preflight.blocking(preflight.check_general(s)) if r["contrôle"] == "TLONE_DATA_ROOT"]
    monkeypatch.delenv("TLONE_DATA_ROOT")
    assert any(r["contrôle"] == "TLONE_DATA_ROOT" for r in preflight.blocking(preflight.check_general(load_settings())))
    # un run en attente bloque un 2e lancement, jusqu'à « Marquer abandonné »
    monkeypatch.setattr(preflight, "check_step", lambda s, step: [])
    monkeypatch.setattr(jobs.subprocess, "Popen", lambda *a, **k: type("P", (), {"pid": 1})())
    rec = jobs.launch(s, "dom_post", None, "testeur")
    monkeypatch.undo()
    monkeypatch.setenv("TLONE_DATA_ROOT", str(p.data_root))
    monkeypatch.setenv("TLONE_RUN_MODE", "local")
    assert [r["run_id"] for r in preflight.pending_runs(s, "dom_post")] == [rec["run_id"]]
    jobs.mark_abandoned(s, rec["run_id"], "testeur")
    assert preflight.pending_runs(s, "dom_post") == []


def test_preflight_carthago_entetes(tmp_path):
    from tlone import preflight
    good, bad = tmp_path / "ok.xlsx", tmp_path / "ko.xlsx"
    pd.DataFrame({"CLIENT ": ["07310104411"], "Nom complet / Raison sociale": ["X"]}).to_excel(good, index=False)
    pd.DataFrame({"Matricule": ["07310104411"], "Intitulé": ["X"]}).to_excel(bad, index=False)
    assert not preflight.blocking(preflight.check_carthago(good))
    ko = preflight.blocking(preflight.check_carthago(bad))
    assert len(ko) == 2 and "Matricule" in ko[0]["action"]


def test_carthago_entetes_reels_et_patch_partie2(tmp_path, monkeypatch):
    """En-têtes de l'export CarthagoDom réel (18/10/2026) : la Partie 2 reçoit les bons noms de colonnes."""
    from tlone import carthago, nbrunner, preflight
    reel = ["Accord pré-dom", "Agence", "Client", "Code bic", "Date apurement", "Date domiciliation",
            "Date déclaration", "Description", "Devise", "Dom Amount Tolerance", "Dom Documents Received",
            "Docs Received Comment", "Docs Received Date", "Invoices Date", "Invoices Reference",
            "Last Expedition Date", "Last Financing Settlement Date", "Mode de règlement", "Montant final",
            "Montant initial", "Name", "Nature opération", "Nom complet", "Numéro de domiciliation"]
    x = tmp_path / "Dossier de domiciliation.xlsx"
    pd.DataFrame([{c: None for c in reel}]).to_excel(x, index=False)
    c = carthago.detect(carthago.read_headers(x))
    assert (c["client"], c["name"], c["dom_ref"], c["agence"]) == ("Client", "Nom complet", "Numéro de domiciliation", "Agence")
    assert not preflight.blocking(preflight.check_carthago(x))
    src = "CARTHAGO_CLIENT_COL = 'Client'\nCARTHAGO_NAME_COL = 'Nom complet/Raison sociale'\nCARTHAGO_DOM_REF_COL = None\n"
    out, applied = nbrunner.patch_source(src, "dom_post", "/mnt/data", None, None, c)
    assert 'CARTHAGO_NAME_COL = "Nom complet"' in out and 'CARTHAGO_DOM_REF_COL = "Numéro de domiciliation"' in out
    assert "CARTHAGO_CLIENT_COL = 'Client'" in out and len(applied) == 2
    monkeypatch.setenv("TLONE_CARTHAGO_NAME_COL", "Name")
    assert carthago.detect(reel)["name"] == "Name"
    # le vrai notebook Partie 2 contient bien les lignes patchées
    import nbformat
    nb = nbformat.read(str(load_settings().notebooks["dom_post"]), as_version=4)
    full = "\n".join(cl.source for cl in nb.cells if cl.cell_type == "code")
    out, applied = nbrunner.patch_source(full, "dom_post", "/mnt/data", None, None, c)
    assert {"CARTHAGO_NAME_COL=Nom complet", "CARTHAGO_DOM_REF_COL=Numéro de domiciliation"} <= set(applied)


def test_job_cmd_python_override():
    from tlone.jobs import _job_cmd
    assert _job_cmd("python /mnt/tlone_app/run_step.py", "/opt/venv-tf4/bin/python") == \
        "/opt/venv-tf4/bin/python /mnt/tlone_app/run_step.py"
    assert _job_cmd("python /mnt/tlone_app/run_step.py", "") == "python /mnt/tlone_app/run_step.py"


def test_ensure_kernel():
    import sys
    from jupyter_client.kernelspec import KernelSpecManager
    name = nbrunner.ensure_kernel()
    spec = KernelSpecManager().get_kernel_spec(name)
    assert spec.argv[0] == sys.executable


def test_job_cmd_follows_app_dir(monkeypatch):
    monkeypatch.delenv("TLONE_JOB_CMD", raising=False)
    from tlone.config import APP_DIR
    assert load_settings().job_cmd_prefix == f"python {APP_DIR / 'run_step.py'}"


# ------------------------------------------------------------ matrice de validation DOM
def _mx(p, journal=None, sources=None):
    from tlone import matrice
    lot = matrice.load_lot(p)
    ms = matrice.build_all(lot, journal or matrice.read_journal(p.dom_journal), sources or matrice.default_sources())
    return {m["fichier"]: m for m in ms}


def _open(m):
    return {(x["ligne"].key, c) for x in m["rows"] for c, cell in x["cells"].items() if cell["open"]}


def test_matrice_anomalies(env):
    p, _ = env
    ms = _mx(p)
    assert _open(ms["KAYA_AHMET_CNT.pdf"]) == {("fin_permis", "TTR"), ("part", "DOM")}
    # signature employeur absente : non obligatoire pour le moment (tolérance ⓘ, ne bloque plus)
    assert _open(ms["KAYA_AHMET_AUG_26.pdf"]) == set()
    sig = next(x for x in ms["KAYA_AHMET_AUG_26.pdf"]["rows"] if x["ligne"].key == "sig_emp")["cells"]["CTR"]
    assert sig["info"] and not sig["open"] and "non obligatoire" in sig["issues"][0][1]
    assert ("nom", "TTR") in _open(ms["DEMIR_ALI_CNT.pdf"]) and ("domref", "CARTHAGO") in _open(ms["DEMIR_ALI_CNT.pdf"])
    demir = {x["ligne"].key: x for x in ms["DEMIR_ALI_CNT.pdf"]["rows"]}
    assert demir["domref"]["missing_ref"]  # Carthago ne connaît pas le client : rien à retenir
    assert ("compte", "DOM") in _open(ms["CELIK_MEHMET_CNT.pdf"])
    # salaire ancien seulement pour une augmentation
    assert "salaire_ancien" in {x["ligne"].key for x in ms["KAYA_AHMET_AUG_26.pdf"]["rows"]}
    assert "salaire_ancien" not in {x["ligne"].key for x in ms["KAYA_AHMET_CNT.pdf"]["rows"]}


def test_matrice_source_confirmation_correction(env):
    p, _ = env
    from tlone import matrice
    f = "KAYA_AHMET_CNT.pdf"
    # le valideur prend le TTR comme référence pour la fin de permis : c'est alors le CTS qui est en écart
    matrice.append_event(p.dom_journal, "v1", f, "SOURCE", ligne="fin_permis", col="TTR")
    m = _mx(p)[f]
    fp = next(x for x in m["rows"] if x["ligne"].key == "fin_permis")
    assert fp["retenu_txt"] == "02/01/2027" and fp["cells"]["CTS"]["open"]
    # confirmation de la lecture DOM (confiance faible) puis correction du CTS
    matrice.append_event(p.dom_journal, "v1", f, "CONFIRME", ligne="part", col="DOM", valeur_vue="312255.00")
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="fin_permis", col="CTS", valeur="02/01/2027")
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="fin_permis", col="CTR", valeur="02/01/2027")
    # la nouvelle fin rend le DOM incohérent et la durée CTS (24 mois) n'est plus l'écart des dates (12 mois)
    assert _open(_mx(p)[f]) == {("fin_permis", "DOM"), ("duree", "CTS")}
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="fin_permis", col="DOM", valeur="02/01/2027")
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="duree", col="CTS", valeur="12")
    # durée CTS 12 : DOM (24), CTR (24) et le montant DOM ne collent plus
    assert _open(_mx(p)[f]) == {("duree", "DOM"), ("duree", "CTR"), ("montant", "DOM")}
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="duree", col="DOM", valeur="12")
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="duree", col="CTR", valeur="12")
    matrice.append_event(p.dom_journal, "v1", f, "CORRIGE", ligne="montant", col="DOM", valeur="3747060.00")
    m = _mx(p)[f]
    assert m["n_open"] == 0, _open(m)
    # une confirmation ne vaut que pour la valeur vue : si la Partie 2 relit autre chose, la case se rouvre
    matrice.append_event(p.dom_journal, "v1", f, "CONFIRME", ligne="part", col="DOM", valeur_vue="999")
    assert ("part", "DOM") in _open(_mx(p)[f])


def test_matrice_vers_planning(env):
    p, _ = env
    from tlone import domvalid, matrice
    f = "KAYA_AHMET_CNT.pdf"
    j = p.dom_journal
    matrice.append_event(j, "v1", f, "SOURCE", ligne="fin_permis", col="TTR")
    for col in ("CTS", "CTR"):
        matrice.append_event(j, "v1", f, "CORRIGE", ligne="fin_permis", col=col, valeur="02/01/2027")
    matrice.append_event(j, "v1", f, "CONFIRME", ligne="part", col="DOM", valeur_vue="312255.00")
    matrice.append_event(j, "v1", f, "VALIDE")
    matrice.append_event(j, "v2", "DEMIR_ALI_CNT.pdf", "REJETE", motif="Domiciliation introuvable dans Carthago")
    ms = list(_mx(p).values())
    rows = matrice.planning_rows(ms)
    fin = rows[(rows["FICHIER"] == f) & (rows["CHAMP"] == "CTS_DATE_FIN_VALIDITE_PERMIS")].iloc[0]
    assert fin["VALEUR_VALIDEE"] == "02/01/2027" and fin["DECISION_MANUELLE"] == "CORRIGEE"
    assert rows[(rows["FICHIER"] == "DEMIR_ALI_CNT.pdf")]["DECISION_MANUELLE"].tolist() == ["REJETEE"]
    man = matrice.merge_manual(domvalid.load_sheet(p.validation_xlsx, domvalid.MANUAL_SHEET), rows)
    dos = domvalid.load_sheet(p.validation_xlsx, domvalid.DOSSIERS_SHEET)
    pa = pd.read_excel(p.param_aug_xlsx, dtype=object)
    val = {m["fichier"] for m in ms if m["decision"] and m["decision"]["action"] == "VALIDE"}
    res = planning.build_planning(dos, man, pa, exclure_validation_incomplete=True, fichiers_valides=val)
    base = res["BASE_VERSIONS_CTS"]
    assert set(base["FICHIER"]) == {f}
    assert str(base.iloc[0]["CTS_DATE_FIN_VALIDITE_PERMIS"]) == "2027-01-02"
    rej = res["REJETS_PLANNING"].set_index("FICHIER")["MOTIF"]
    assert rej["DEMIR_ALI_CNT.pdf"] == "REJETE_EN_PARTIE_2A"
    assert rej["CELIK_MEHMET_CNT.pdf"] == "VALIDATION_HUMAINE_INCOMPLETE"


def test_sources_reference_parametre(env):
    p, _ = env
    from tlone import matrice
    t = matrice.sources_table()
    t.loc[t["LIGNE"] == "part", "SOURCE_REFERENCE"] = "DOM"
    t.loc[t["LIGNE"] == "nom", "SOURCE_REFERENCE"] = "ATLAS"  # source inconnue : ignorée
    p.sources_reference_xlsx.parent.mkdir(parents=True, exist_ok=True)
    t.to_excel(p.sources_reference_xlsx, index=False)
    src = matrice.default_sources(p.sources_reference_xlsx)
    assert src["part"] == "DOM" and src["nom"] == "CTS"


def test_page_rotation(tmp_path):
    import json
    from tlone import matrice
    (tmp_path / "x.json").write_text(json.dumps({"source_file": "A.pdf", "page_records": [
        {"page_num": 2, "rotation_clockwise_applied": 270}]}), encoding="utf-8")
    (tmp_path / "A.json").write_text(json.dumps({"source_file": "A.pdf", "page_records": [
        {"page_num": 2, "rotation_clockwise_applied": 90}]}), encoding="utf-8")
    assert matrice.page_rotation(tmp_path, "A.pdf", 2) == 90 and matrice.page_rotation(tmp_path, "A.pdf", 1) == 0


def test_matrice_bloquants(env):
    p, _ = env
    from tlone import matrice
    j = p.dom_journal
    # signature absente : non obligatoire pour le moment (24/10/2026) -> ni bloquante ni à confirmer
    f = "KAYA_AHMET_AUG_26.pdf"
    m = _mx(p)[f]
    assert ("sig_emp", "CTR") not in _open(m) and m["n_bloq"] == 0 and m["n_open"] == 0
    # compte : la correction est re-contrôlée (longueur, clé Modulo 97)
    f = "CELIK_MEHMET_CNT.pdf"
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="compte", col="DOM", valeur="02700731010830000155")
    cell = next(x for x in _mx(p)[f]["rows"] if x["ligne"].key == "compte")["cells"]["DOM"]
    assert cell["open"] and "attendue 61" in cell["issues"][0][1]
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="compte", col="DOM", valeur="02700731010830000161")
    assert ("compte", "DOM") not in _open(_mx(p)[f])
    # Carthago ne connaît pas le client : bloquant, non confirmable
    f = "DEMIR_ALI_CNT.pdf"
    matrice.append_event(j, "v1", f, "CONFIRME", ligne="domref", col="CARTHAGO", valeur_vue="None")
    assert ("domref", "CARTHAGO") in _open(_mx(p)[f])
    assert matrice.compte_issues("02700731010910800156")[0][1] == "Clé Modulo 97 invalide : lue 56, attendue 91"


# ------------------------------------------------------------ matrice OV
def test_ov_matrice_et_correction(env):
    p, s = env
    from tlone import ov_matrice as ov
    res = _planning(p)
    planning.export(res, p.planning_xlsx, p.planning_csv)
    pl = planning.load_planning(p.planning_csv)
    a = transfers.read_audit_xlsx(transfers.latest_audit(p.tr_out)[0])
    args = (pl, transfers.load_registre(p.registre_csv),
            transfers.load_referential(p.atf_xlsx, transfers.ATF_COLUMNS),
            transfers.load_referential(p.autorisations_xlsx, transfers.AUT_COLUMNS), s)
    f = "07310-108259-202602.pdf"
    sy, ch = transfers.run_controls(a, *args)
    row = sy.set_index("FICHIER").loc[f]
    assert row["STATUT_CONTROLE"] == "NON_CONFORME"
    m = ov.build_ov(f, a[a["fichier"] == f].iloc[0].to_dict(), sy[sy["FICHIER"] == f].iloc[0].to_dict(), ch, pl, {})
    plaf = next(x for x in m["rows"] if x["ligne"].key == "montant")      # montant OV et montants max. consolidés
    assert plaf["cells"]["ANNEXE_II"]["open"] and plaf["cells"]["ANNEXE_II"]["bloq"]
    # le valideur corrige le montant lu sur l'OV : les contrôles sont recalculés sur la valeur corrigée
    ov.append_event(ov_journal := p.tr_root / "journal_validation_ov.jsonl", "v1", f, "CORRIGE",
                    champ="OV.montant_chiffres", valeur="300000")
    a2 = ov.apply_corrections(a, ov.read_journal(ov_journal))
    sy2, ch2 = transfers.run_controls(a2, *args)
    t09 = ch2[(ch2["FICHIER"] == f) & (ch2["CODE"] == "T09")].iloc[0]
    assert t09["RESULTAT"] == "OK"
    # échantillon de contrôle reproductible
    fs = [f"F{i}" for i in range(100)]
    assert ov.sample(fs, "lot1") == ov.sample(fs, "lot1") and len(ov.sample(fs, "lot1")) == 5
    assert len(ov.sample(fs[:10], "x")) == 3
    lf = ov.list_frame(sy2, transfers.load_decisions(p.decisions_csv))
    assert len(lf) == len(sy2) and {"FEU", "MARGE", "BLOQUANTS"} <= set(lf.columns)


def test_matrice_regles_agence():
    from tlone import matrice as mx
    assert mx.permis_parts("20-00010423/31-26-000472") == ("20-00010423", "31-26-000472")
    assert mx.permis_parts("20/00010423 / 31/26/000472") == ("20-00010423", "31-26-000472")
    assert mx.permis_parts("31-26-000472") == (None, "31-26-000472")
    assert mx.same("permis", "20-00010423/31-26-000472", "31/26/000472")[0] is True
    assert mx.same("permis", "20-00010423/31-26-000472", "20-00010423/31-26-000473")[0] is False
    assert mx.to_months("2 ANS, 0 JOURS") == 24 and mx.to_months("1 AN 6 MOIS") == 18 and mx.to_months("24") == 24
    from datetime import date
    assert mx.months_between(date(2026, 1, 3), date(2028, 1, 2)) == 24
    assert mx.months_between(date(2026, 2, 1), date(2027, 1, 31)) == 12


def test_nouveaux_champs_pris_en_charge(env):
    """Les 4 champs demandés par l'agence : « à extraire » aujourd'hui, contrôlés dès qu'ils sont dans les JSON."""
    p, _ = env
    f = "KAYA_AHMET_CNT.pdf"

    def cell(m, k, c):
        return next(x for x in m["rows"] if x["ligne"].key == k)["cells"][c]
    m = _mx(p)[f]
    assert cell(m, "sig_trav", "DOM")["a_extraire"] and not cell(m, "sig_trav", "DOM")["open"]
    # simulation : la Partie 2 produit maintenant les nouveaux champs
    fld = pd.read_csv(p.fields_csv, dtype=object, encoding="utf-8-sig")
    base = fld[fld["FICHIER"] == f].iloc[0].to_dict()
    new = [("ENGAGEMENT_DOMICILIATION", 1, "DOM_SIGNATURE_TRAVAILLEUR_PRESENTE", "False"),
           ("ENGAGEMENT_DOMICILIATION", 1, "DOM_LU_APPROUVEE_TRAVAILLEUR_PRESENTE", "True"),
           ("CONTRAT_SPECIFIQUE", 3, "CTS_NUMERO_CONTRAT_INSPECTION_TRAVAIL", "CT-2026-0002"),
           ("CONTRAT_SPECIFIQUE", 3, "CTS_DATE_VISA_INSPECTION_TRAVAIL", "05/01/2026")]
    rows = [dict(base, TYPE_DOCUMENT=d, PAGE=pg, CHAMP=c, RAW=v, NORMALIZED=v, DQ_STATUS_FIELD="OK") for d, pg, c, v in new]
    pd.concat([fld, pd.DataFrame(rows)]).to_csv(p.fields_csv, index=False, encoding="utf-8-sig")
    m = _mx(p)[f]
    sg = cell(m, "sig_trav", "DOM")                                # signature absente : non obligatoire pour le moment
    assert not sg["a_extraire"] and sg["info"] and not sg["open"] and not sg["bloq"]
    assert cell(m, "lu_approuve", "DOM")["text"] == "Présente"
    assert cell(m, "contrat", "CTS")["open"]                       # CT-2026-0002 ≠ n° contrat du DOM
    assert cell(m, "date_visa", "CTS")["text"] == "05/01/2026"
    # un autre dossier sans ces champs n'affiche plus « à extraire » mais « — » (champ connu, absent de la pièce)
    assert not cell(_mx(p)["DEMIR_ALI_CNT.pdf"], "sig_trav", "DOM")["a_extraire"]


def test_partie2_attend_partie1(env, monkeypatch):
    p, s = env
    from tlone import jobs, preflight
    monkeypatch.setattr(jobs.subprocess, "Popen", lambda *a, **k: type("P", (), {"pid": 1})())
    real = preflight.check_step
    monkeypatch.setattr(preflight, "check_step", lambda s_, step: [])
    jobs.launch(s, "dom_extract", 1, "testeur")
    monkeypatch.setattr(preflight, "check_step", real)
    ko = {r["contrôle"] for r in preflight.blocking(preflight.check_step(s, "dom_post"))}
    assert "Extraction DOM en cours" in ko


# ------------------------------------------------------------ worker GPU (modèle chargé une fois)
def test_worker_modele_charge_une_fois(env, tmp_path, monkeypatch):
    import dataclasses
    import sys as _sys
    import nbformat
    from tlone import jobs, preflight, worker
    p, s = env
    nb = nbformat.v4.new_notebook()
    nb.cells = [
        nbformat.v4.new_code_cell("MAX_PDFS = None\nfrom pathlib import Path\n"
                                  "INPUT_DIR = Path('/mnt/data/domiciliations_in')\n"
                                  "pdfs = sorted(INPUT_DIR.glob('*.pdf'))\n"
                                  "if MAX_PDFS is not None:\n    pdfs = pdfs[:int(MAX_PDFS)]"),
        nbformat.v4.new_code_cell("# simule AutoModel.from_pretrained(...)\nLOADS = globals().get('LOADS', 0) + 1\nCOUNT = 0"),
        nbformat.v4.new_code_cell("if not pdfs:\n    print('rien')\nelse:\n    COUNT += 1\n"
                                  "    (INPUT_DIR.parent / 'out.txt').write_text(f'{LOADS} {COUNT} {len(pdfs)}')"),
    ]
    fake = tmp_path / "fake_part1.ipynb"
    nbformat.write(nb, str(fake))
    s = dataclasses.replace(s, notebooks={**s.notebooks, "dom_extract": fake}, run_mode="worker")
    # sans worker : lancement bloqué
    assert any(r["contrôle"] == "Worker GPU" for r in preflight.blocking(preflight.check_step(s, "dom_extract")))
    w = worker.Worker(s, ["dom_extract"], pythons={"dom_extract": _sys.executable})
    w.beat()
    assert worker.worker_state(s)["state"] == "ATTENTE"
    monkeypatch.setattr(preflight, "check_step", lambda s_, step: [])
    r1 = jobs.launch(s, "dom_extract", 1, "testeur")
    w.loop(once=True)
    assert (p.data_root / "out.txt").read_text() == "1 1 1"
    r2 = jobs.launch(s, "dom_extract", None, "testeur")
    w.loop(once=True)
    assert (p.data_root / "out.txt").read_text() == "1 2 4"   # modèle chargé 1 fois, 2 traitements, 4 PDF
    runs = {r["run_id"]: r for r in jobs.list_runs(s)}
    assert runs[r1["run_id"]]["status"] == runs[r2["run_id"]]["status"] == "SUCCEEDED"
    w.stop()
    assert worker.worker_state(s) is None


def test_worker_reconnait_les_vrais_notebooks(env):
    """Sur les vrais notebooks : cellule de config, chargement du modèle et traitement bien repérés."""
    import sys as _sys
    from tlone import worker
    _, s = env
    for step, n_proc in (("dom_extract", 1), ("ov_extract", 2)):
        w = worker.WarmNotebook(s, step, _sys.executable)
        model = [i for i, c in enumerate(w.nb.cells) if c.cell_type == "code" and worker.MODEL_MARKER in c.source]
        assert model and w.cfg < model[0] < min(w.process), step     # config -> modèle -> traitement
        assert len(w.process) == n_proc and "GLOB_PDF_CASE_INSENSITIVE" in w.patches
        assert "MAX_PDFS" in w.nb.cells[w.cfg].source or step == "ov_extract"
        if step == "ov_extract":       # v27 : la boucle des PDF fait partie du traitement ; INPUT_DIR redirigé
            assert any("pdf_path in enumerate(pdfs" in w.nb.cells[i].source for i in w.process)
            assert "INPUT_DIR=transferts_in" in w.patches and "NO_PIP_INSTALL" in w.patches


def test_worker_signal_independant_du_fuseau(env):
    """L'App et le workspace peuvent avoir des fuseaux différents : l'âge du signal vient de l'epoch."""
    import json as _json
    import time as _time
    from tlone import worker
    _, s = env
    hb = worker.heartbeat_path(s)
    hb.parent.mkdir(parents=True, exist_ok=True)
    hb.write_text(_json.dumps({"ts": "2026-10-18T10:34:00+05:00", "epoch": _time.time(), "host": "ws",
                               "state": "ATTENTE", "current": None, "warm": {}}))
    assert worker.worker_state(s)["host"] == "ws"
    hb.write_text(_json.dumps({"ts": "x", "epoch": _time.time() - 600, "host": "ws", "state": "ATTENTE"}))
    assert worker.worker_state(s) is None
    assert "il y a 600 s" in worker.worker_absent_text(s)
    hb.unlink()
    assert worker.worker_absent_text(s).startswith("non démarré")


def test_domref_recontrole_apres_lecture_ctr_vide():
    """Partie 2 « n° illisible » (CTR vide au moment du run) : l'App recompare avec la valeur CTR courante."""
    from tlone import matrice as mx
    def cells(ctr, status="NUMERO_DOMICILIATION_NOT_NORMALIZABLE", cands=("271901|2019.3|40|002031|DZD",)):
        return {"CTR": {"value": ctr, "issues": []},
                "CARTHAGO": {"value": cands[0] if cands else None, "carthago_status": status, "candidats": list(cands),
                             "issues": [("bloq", "Carthago : n° de domiciliation illisible")]}}
    c = cells("27-19-01 | 2019.3 | 40 | 002031 | DZD")          # n° du contrat = Carthago
    mx._domref_recalcul(c)
    assert c["CARTHAGO"]["issues"] == [] and c["CARTHAGO"]["recalcul"]
    c = cells(None)                       # CTR sans n° : non bloquant, domiciliation unique du client retenue
    mx._domref_recalcul(c)
    assert [k for k, _ in c["CARTHAGO"]["issues"]] == ["info"] and c["CARTHAGO"]["value"] == "271901|2019.3|40|002031|DZD"
    c = cells(None, cands=("271901201934000111DZD", "271901202434000222DZD"))   # plusieurs : choisir
    mx._domref_recalcul(c)
    assert c["CARTHAGO"]["issues"][0][0] == "verif" and "2 domiciliations" in c["CARTHAGO"]["issues"][0][1]
    c = cells(None, cands=())
    mx._domref_recalcul(c)
    assert c["CARTHAGO"]["issues"][0][0] == "bloq"
    c = cells("271901201934000999DZD", status="NO_EXACT_DOMICILIATION_REFERENCE_MATCH")
    mx._domref_recalcul(c)
    assert c["CARTHAGO"]["issues"][0][0] == "bloq" and "Aucune domiciliation" in c["CARTHAGO"]["issues"][0][1]
    c = cells("x", status="CLIENT_NOT_FOUND")                   # autres statuts : inchangés
    mx._domref_recalcul(c)
    assert c["CARTHAGO"]["issues"][0][1].startswith("Carthago :")


def test_preflight_worker_utilise_l_environnement_du_worker(env, monkeypatch):
    """En mode worker, les versions contrôlées sont celles du workspace du worker, pas du conteneur de l'App."""
    import dataclasses
    import json as _json
    import time as _time
    from tlone import preflight, worker
    p, s = env
    s = dataclasses.replace(s, run_mode="worker")
    monkeypatch.setattr(preflight, "probe_python", lambda py, mods: {m: "5.20.1" for m in mods})  # conteneur App
    mods = preflight.STEP_MODULES["ov_extract"]
    hb = worker.heartbeat_path(s)
    hb.parent.mkdir(parents=True, exist_ok=True)
    hb.write_text(_json.dumps({"epoch": _time.time(), "ts": "x", "host": "ws", "state": "ATTENTE",
                               "steps": ["dom_extract", "dom_post"], "warm": {},
                               "env": {"ov_extract": {"python": "/opt/tf4/python",
                                                      "versions": {**{m: "ok" for m in mods},
                                                                   "transformers": "4.57.0"}}}}))
    rows = {r["contrôle"]: r for r in preflight.check_step(s, "ov_extract")}
    assert rows["transformers ov_extract (worker)"]["niveau"] == preflight.OK
    assert "transformers ov_extract" not in rows                      # pas de sondage local
    assert rows["Worker ov_extract"]["niveau"] == preflight.WARN      # modèle non préchargé : alerte seulement
    assert not preflight.blocking(list(rows.values())) or all(
        r["contrôle"] in ("PDF OV à traiter",) for r in preflight.blocking(list(rows.values())))


def test_worker_echec_chargement_n_arrete_pas_le_worker(env):
    """Un notebook qui échoue au préchargement (ex. ImportError transformers) n'arrête pas le worker."""
    import dataclasses
    import sys as _sys
    import nbformat
    from tlone import preflight, worker
    p, s = env
    nb = nbformat.v4.new_notebook()
    nb.cells = [nbformat.v4.new_code_cell("from transformers_inexistant import AutoModelForVision2Seq"),
                nbformat.v4.new_code_cell("pdfs = sorted([])"), nbformat.v4.new_code_cell("for pdf_path in pdfs: pass")]
    bad = p.data_root / "bad.ipynb"
    nbformat.write(nb, str(bad))
    s = dataclasses.replace(s, notebooks={**s.notebooks, "ov_extract": bad}, run_mode="worker")
    w = worker.Worker(s, ["ov_extract"], prewarm=True, pythons={"ov_extract": _sys.executable})
    w.loop(once=True)                                   # ne lève pas
    assert "ov_extract" in w.errors and "ov_extract" not in w.warm
    w.beat()
    rows = {r["contrôle"]: r for r in preflight.check_step(s, "ov_extract")}
    assert rows["Worker ov_extract"]["niveau"] == preflight.KO
    w.stop()


def test_patch_desactive_pip_install():
    out, applied = nbrunner.patch_source("%pip install -q pymupdf transformers\nif True:\n    !pip install x\nprint(1)",
                                         "ov_extract", "", None, None)
    assert "NO_PIP_INSTALL" in applied and "%pip" not in out.split("#")[0]
    compile(out, "cellule", "exec")


def _v27_excel(path, version="27", ov=None):
    """Classeur produit par la cellule « Export Excel » du vrai notebook v27 / v28 (code exécuté tel quel)."""
    import json as _json
    nb = _json.load(open(Path(__file__).resolve().parents[1] / "notebooks" / f"pipeline_ocr_vl{version}.ipynb", encoding="utf-8"))
    src = next("".join(c["source"]) for c in nb["cells"] if "def create_excel" in "".join(c["source"]))
    ns: dict = {}
    exec("from pathlib import Path\nfrom collections import defaultdict\nfrom openpyxl import Workbook\n"
         "from openpyxl.styles import Font, PatternFill, Alignment, Border, Side\n"
         "from openpyxl.utils import get_column_letter\n" + src, ns)
    row = {"fichier": "02700731010825900165-202603P1.pdf", "date_traitement": "2026-10-19 01:00:00",
           "temps_total_s": 12.0, "tokens_total": 1000, "pages_trouvees": "ANNEXE_I, ANNEXE_II, BULLETIN, OV",
           "pages_manquantes": None, "anomalies": None,
           "OV": {"monnaie": "DZD", "montant_chiffres": 312255.0, "montant_lettres": "trois cent douze mille deux cent cinquante-cinq dinars",
                  "periode": "MARS 2026 P1", "mois": "MARS", "annee": "2026", "tranche": "P1", "complement_ov": None,
                  "date_demande": "05/03/2026", "compte_donneur_ordre": "02700731010825900165",
                  "nature_paiement_autre_libelle": "EUR", "beneficiaire_nom": "KAYA AHMET",
                  "beneficiaire_compte": "TR33", "beneficiaire_adresse": "X",
                  "code_swift_banque_beneficiaire": "TEBUTRIS", "nom_banque_beneficiaire": "TEB"},
           "ANNEXE_I": {"nom_prenom_employe": "KAYA AHMET", "compte_bancaire_local": "02700731010825900165"},
           "ANNEXE_II": {"mois_transfert": "MARS 2026", "nom_prenom_travailleur": "KAYA AHMET",
                         "compte_bancaire_local": "02700731010825900165", "salaire_mensuel": 300000.0,
                         "part_transferable": 312255.0, "numero_domiciliation": "999"},
           "BULLETIN": {"nom_prenom_salarie": "KAYA AHMET", "periode_bulletin": "MARS 2026", "net_a_payer": 340000.0}}
    row["OV"].update(ov or {})
    ns["create_excel"](path, [row])


def test_lecture_classeur_v27_et_controles_retires(tmp_path):
    """v27 : libellés différents -> mêmes clés que v26 ; T16 (n° DOM) et T08 (net/salaire) ne sont plus contrôlés."""
    _v27_excel(tmp_path / "audit_transferts_20261019_0100.xlsx")
    a = transfers.read_audit_xlsx(tmp_path / "audit_transferts_20261019_0100.xlsx")
    for col in ("OV.compte_debiteur", "OV.swift", "OV.banque_nom", "OV.periode", "OV.mois", "ANNEXE_II.nom_prenom",
                "ANNEXE_II.compte_bancaire", "ANNEXE_II.numero_domiciliation", "BULLETIN.nom_prenom",
                "BULLETIN.mois_bulletin", "ANNEXE_I.nom_prenom"):
        assert col in a.columns, col
    s = load_settings()
    synth, rows = transfers.control_dossier(a.iloc[0].to_dict(), pd.DataFrame(), pd.DataFrame(columns=["PERIODE_ID", "DOM_ID", "FICHIER", "_MONTANT"]),
                                            pd.DataFrame(), pd.DataFrame(), s)
    codes = {r["CODE"]: r for r in rows}
    assert "T16" not in codes and "T08" not in codes
    assert codes["T07"]["RESULTAT"] == "OK"               # MARS 2026 partout, malgré « P1 »
    assert codes["T03"]["RESULTAT"] == "OK"


def test_lecture_case_70_periode_tranche_complement():
    from tlone.periode import lire
    cas = {"JUIN.2026 PART2": ("2026-06P2", None), "Juin 26": ("2026-06", None), "Sept 25 P2": ("2025-09P2", None),
           "Complément janvier 2026": ("2026-01", "COMPLEMENT"), "COMP JANV 26": ("2026-01", "COMPLEMENT"),
           "Salaire 05-2026 2eme partie": ("2026-05P2", None), "Fev.26 P 1": ("2026-02P1", None),
           "salaire mai 2026 15/06/2026": ("2026-05", None), "052026P1": ("2026-05P1", None),
           "RELIQUAT DEC 2025": ("2025-12", "RELIQUAT")}
    for txt, (per, comp) in cas.items():
        r = lire(txt)
        assert (r["periode"], r["complement"]) == (per, comp), (txt, r)
    assert lire("TR Août 2026 - 12 jrs")["statut"] == "A_PRECISER" and lire("Commission")["periode"] is None


def test_v28_case70_brute_et_complement(tmp_path):
    """v28 : la case 70 est recopiée telle quelle ; TL-One en déduit période / tranche / complément.
    Un complément sur une période déjà transférée est admis (T14 OK) si le cumul reste ≤ part de la période (T11)."""
    import json as _json
    nb = _json.load(open(Path(__file__).resolve().parents[1] / "notebooks" / "pipeline_ocr_vl28.ipynb", encoding="utf-8"))
    prompt = next("".join(c["source"]) for c in nb["cells"] if "PROMPT_UNIVERSEL" in "".join(c["source"]))
    assert "reference_paiement" in prompt and "observation du donneur d" in prompt and "complement_ov :" not in prompt
    _v27_excel(tmp_path / "audit_transferts_20261021_0100.xlsx", "28",
               {"reference_paiement": "Complément MARS.26", "periode": "Complément MARS.26", "mois": "AVRIL",
                "tranche": None, "montant_chiffres": 50000.0, "montant_lettres": "cinquante mille dinars"})
    a = transfers.read_audit_xlsx(tmp_path / "audit_transferts_20261021_0100.xlsx")
    d = a.iloc[0].to_dict()
    assert d["OV.reference_paiement"] == "Complément MARS.26"
    s = load_settings()
    pl = pd.DataFrame([{"PERIODE_ID": "pm", "DOM_ID": "d1", "SIEGE_RACINE": "07310108259", "PERIODE_TL": "2026-03",
                        "MOIS_BASE": "2026-03", "MONTANT_MAX_THEORIQUE": 400000.0, "NUMERO_DOMICILIATION": "999",
                        "DOM_COMPTE_LOCAL": "02700731010825900165", "NOM_TRAVAILLEUR": "KAYA AHMET",
                        "DATE_DEBUT_SEGMENT": "2026-03-01", "DATE_FIN_SEGMENT": "2026-03-31", "PARTIE": None}])
    reg = pd.DataFrame([{"PERIODE_ID": "pm", "DOM_ID": "d1", "FICHIER": "ATLAS:TL1", "_MONTANT": 300000.0}])
    sy, rows = transfers.control_dossier(d, pl, reg, pd.DataFrame(), pd.DataFrame(), s)
    c = {r["CODE"]: r for r in rows}
    assert sy["PERIODE_TL"] == "2026-03" and sy["COMPLEMENT"] == "COMPLEMENT"      # texte de la case, pas « AVRIL »
    assert "lecture du modèle différente" in c["T07"]["DETAIL"]
    assert c["T14"]["RESULTAT"] == "OK" and "Complément" in c["T14"]["CONTROLE"]
    assert c["T11"]["RESULTAT"] == "OK"                                           # 300 000 + 50 000 ≤ 400 000
    reg.loc[0, "_MONTANT"] = 380000.0
    _, rows = transfers.control_dossier(d, pl, reg, pd.DataFrame(), pd.DataFrame(), s)
    assert {r["CODE"]: r for r in rows}["T11"]["RESULTAT"] == "KO"                 # dépasse la part de la période
    _, rows = transfers.control_dossier(d, pl, reg.iloc[0:0], pd.DataFrame(), pd.DataFrame(), s)
    assert {r["CODE"]: r for r in rows}["T14"]["RESULTAT"] == "KO"                 # complément sans TL précédent


def test_execution_complement_pas_doublon():
    from tlone import execution as ex
    T = lambda ref, per, mt, st, comp=None: {"REF_TL": ref, "CLIENT": "07310100001", "NOM": "X", "PERIODE": per,  # noqa: E731
                                           "MOIS": per[:7], "TRANCHE": None, "COMPLEMENT": comp, "STATUT_PERIODE": "OK",
                                           "OBJET": per, "MONTANT_TRADE": mt, "STATUT_TRADE": st}
    trade = pd.DataFrame([T("TL1", "2026-01", 200000.0, "Exécuté"), T("TL2", "2026-01", 100000.0, "En cours", "COMPLEMENT"),
                          T("TL3", "2026-01", 100000.0, "En cours")])
    atlas = pd.DataFrame({"REF_TL": ["TL1"]})
    pl = pd.DataFrame([{"PERIODE_ID": "p", "SIEGE_RACINE": "07310100001", "PERIODE_TL": "2026-01", "MOIS_BASE": "2026-01",
                        "MONTANT_MAX_THEORIQUE": 250000.0}])
    r = ex.plan(trade, atlas, None, None, None, None, None, None, planning=pl).set_index("REF_TL")
    assert r.loc["TL3", "STATUT"] == "DOUBLON"                                   # sans mention de complément
    r = ex.plan(trade[trade["REF_TL"] != "TL3"], atlas, None, None, None, None, None, None, planning=pl).set_index("REF_TL")
    assert r.loc["TL2", "STATUT"] == "COMPLEMENT" and r.loc["TL2", "DOUBLON_AVEC"] == ""
    assert "> part de la période 250,000.00" in r.loc["TL2", "MOTIFS"]          # 200 000 + 100 000 > 250 000


def test_periode_annee_deux_chiffres():
    assert transfers.parse_periode("JANV 26 P2") == ("2026-01", "P2")
    assert transfers.parse_periode("MARS 2026") == ("2026-03", None)


# ------------------------------------------------------------------ vigilance
def _modele_vigilance(path):
    """Réplique du modèle « CHECK LIST TRANSFERT LIBRE » (mêmes libellés / lignes que le modèle BNP fourni)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Feuil1"
    ws["A2"] = "CHECK LIST TRANSFERT LIBRE"
    for r, lab in ((5, "Client:"), (6, "Réf Dossier:"), (7, "Circuit:"), (8, "Type de dossier:")):
        ws.cell(r, 1, lab)
    for c, lab in zip("ABCDE", ["Contrôle", "Type", "Valeur", "HIT / No HIT / KO", "COMMENTAIRES"]):
        ws[f"{c}10"] = lab
    ws["A11"] = "VIGILANCE"
    ws.merge_cells("A11:E11")
    rub = ["Banque du bénéficiaire", "Pays de la banque du bénéficiaire", "Banque intermédiaire",
           "Pays de la banque intermédiaire", "Bénéficiaire", "Pays du bénéficiaire", "Pays d’origine et provenance",
           "Nom du fabricant", "Pays du fabricant", "Compagnie maritime ou aérienne", "Navire", "Expéditeur sur le BL",
           "Agent signataire P/C compagnie maritime", "Transitaire", "Transporteur LVI", "Port ou aéroport d’embarquement",
           "Port ou aéroport de débarquement", "Lieu de prise en charge sur le BL", "Lieu établissement document transport",
           "Lieu de travail", "Divers individus sur les documents", "Divers entités sur les documents", "Divers lieux"]
    for i, lab in enumerate(rub):
        ws.cell(12 + i, 1, lab)
        ws.cell(12 + i, 2, "I" if "individus" in lab else "C")
    ws["A37"] = "Valideur :"
    wb.save(path)


def _retour(path_in, results: dict, individus=None) -> bytes:
    import io
    from openpyxl import load_workbook
    wb = load_workbook(path_in)
    ws = wb.worksheets[0]
    for row in range(12, 35):
        ws[f"D{row}"] = "OK ==> RAS" if ws[f"C{row}"].value else "Rubrique non renseignée"
    if ws["C32"].value:
        ws["D32"] = "\n".join(["OK ==> RAS"] * len(str(ws["C32"].value).split("\n")))
    for cell, val in results.items():
        ws[cell] = val
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_vigilance_regroupement_preparation_retours_decisions(tmp_path):
    import io
    import zipfile
    from openpyxl import load_workbook
    from tlone import vigilance as v
    tpl = tmp_path / "modele.xlsx"
    _modele_vigilance(tpl)
    assert v.check_template(tpl) == []
    ds = [dict(fichier="f1", ref_client="07310108547", client="KAYA AHMET", beneficiaire="AHMET KAYA", iban="TR33 0006",
               swift="TEBUTRIS", banque="TEB", pays_banque="TURQUIE", individus=["MEHMET KAYA", "AYSE KAYA"]),
          dict(fichier="f2", ref_client="07310108547", client="KAYA AHMET", beneficiaire="AHMET KAYA", iban="TR330006",
               swift="TEBUTRIS", banque="TEB", pays_banque="TURQUIE", individus=["AYSE KAYA", "SIGNATAIRE X"]),
          dict(fichier="f3", ref_client="07310262001", client="DEMIR ALI", beneficiaire="DEMIR FATMA", iban="TR99",
               swift="ZIRATRIS", banque="ZIRAAT", pays_banque="TURQUIE", individus=[]),
          dict(fichier="f4", ref_client="07310262001", client="DEMIR ALI", beneficiaire="DEMIR ALI", iban="TR11",
               swift="ZIRATRIS", banque="ZIRAAT", pays_banque="TURQUIE", individus=[])]
    units = v.build_units(ds)
    by = {tuple(u["fichiers"]): u for u in units}
    assert len(units) == 3                                     # KAYA : même bénéficiaire -> 1 fichier ; DEMIR : 2
    assert by[("f1", "f2")]["beneficiaire"] == ""              # bénéficiaire = client (ordre inversé) : omis
    assert by[("f1", "f2")]["individus"] == ["MEHMET KAYA", "AYSE KAYA", "SIGNATAIRE X"]
    assert by[("f3",)]["beneficiaire"] == "DEMIR FATMA"
    m = v.prepare(units, tpl, tmp_path / "lot", "TOSYALI IRON STEEL INDUSTRY ALGERIE")
    ws = load_workbook(tmp_path / "lot" / by[("f1", "f2")]["fichier"]).worksheets[0]
    assert (ws["C5"].value, ws["C6"].value, ws["C7"].value, ws["C8"].value) == ("KAYA AHMET", "07310 108547", "TL", "Création")
    assert ws["C12"].value == "TEB" and ws["C13"].value == "TURQUIE" and ws["C31"].value.startswith("TOSYALI")
    assert ws["C32"].value == "MEHMET KAYA\nAYSE KAYA\nSIGNATAIRE X" and ws["C16"].value is None
    assert zipfile.ZipFile(m["zip"]).namelist() == [u["fichier"] for u in units]

    jp = tmp_path / "journal.jsonl"
    v.append(jp, "u", "ENVOI", lot=m["lot"], unites=units)
    st = v.statuses(v.read_journal(jp))
    assert {f: x["statut"] for f, x in st.items()} == dict.fromkeys(["f1", "f2", "f3", "f4"], "EN_ATTENTE")
    ret = [("Recherche_Vigilance_Data_A.xlsx", _retour(tmp_path / "lot" / by[("f1", "f2")]["fichier"], {})),
           ("Recherche_Vigilance_Data_B.xlsx", _retour(tmp_path / "lot" / by[("f3",)]["fichier"], {"D16": "OK ==> HIT"})),
           ("Recherche_Vigilance_Data_C.xlsx", _retour(tmp_path / "lot" / by[("f4",)]["fichier"],
                                                       {"D12": "KO ==> La Plateforme Vigilance est indisponible"}))]
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as zz:
        for n, b in ret:
            zz.writestr(n, b)
    res = v.import_returns([("retours.zip", z.getvalue())], v.read_journal(jp), jp, "u", tmp_path / "res")
    assert set(res["synthese"]["Statut"]) == {"RAS", "HIT", "KO_A_REFAIRE"} and res["controles"].empty
    st = v.statuses(v.read_journal(jp))
    assert [st[f]["statut"] for f in ("f1", "f2", "f3", "f4")] == ["RAS", "RAS", "HIT", "KO_A_REFAIRE"]
    assert v.eligible(st) == {"f1", "f2"}
    dec = pd.DataFrame({"FICHIER": ["f1", "f2", "f3", "f4", "f5"], "DECISION": ["VALIDE"] * 5})
    assert v.candidates(dec, st) == ["f4", "f5"]               # KO à refaire + jamais envoyé
    with pytest.raises(ValueError):
        v.decide(jp, "u", ["f3"], "LEVEE", "  ")               # motif obligatoire
    v.decide(jp, "u", ["f3"], "LEVEE", "Homonyme : date de naissance différente")
    assert v.statuses(v.read_journal(jp))["f3"]["statut"] == "HIT_LEVE"
    assert v.eligible(v.statuses(v.read_journal(jp))) == {"f1", "f2", "f3"}


def test_vigilance_classement_et_controles():
    from tlone import vigilance as v
    assert v.classify("OK ==> RAS")[0] == "RAS" and v.classify("OK ==> HIT")[0] == "HIT"
    assert v.classify("KO ==> La Plateforme Vigilance est indisponible")[0] == "KO"
    assert v.classify("Rubrique non renseingée")[0] == "INFO" and v.classify("")[0] == "MANQUANT"
    assert v.classify("peut-être")[0] == "INCONNU"
    u = {"client": "A B", "individus": ["X", "Y"]}
    ret = {"client": "A B", "individus_row": 32, "rubriques": [
        {"ligne": 32, "rubrique": "Divers individus", "valeur": "X\nY", "resultat": "OK ==> RAS", "commentaire": ""}]}
    statut, _, ctl = v.analyse(u, ret)
    assert statut == "A_VERIFIER" and "Décalage" in ctl[0]     # 2 individus, 1 résultat : appariement à vérifier
    assert v.format_ref("07310108547") == "07310 108547" and v.pays_banque("TEBUTRIS", None) == "TURQUIE"
    assert v.safe("=HYPERLINK(1)") == "HYPERLINK(1)"


def test_registre_exige_vigilance():
    import tempfile
    sy = pd.DataFrame([{"FICHIER": f, "PERIODE_ID": f"P{f}", "MONTANT_OV": 1.0} for f in ("a", "b")])
    dec = pd.DataFrame({"FICHIER": ["a", "b"], "DECISION": ["VALIDE", "VALIDE"], "COMMENTAIRE": ["", ""]})
    with tempfile.TemporaryDirectory() as d:
        n = transfers.register(Path(d) / "reg.csv", sy, dec, "u", vigilance={"a": "RAS", "b": "HIT"})
        reg = transfers.load_registre(Path(d) / "reg.csv")
    assert n == 1 and list(reg["FICHIER"]) == ["a"] and list(reg["STATUT_VIGILANCE"]) == ["RAS"]


def test_vigilance_analyse_hit_dans_tlone(tmp_path):
    """HIT commenté dans l'app : non avéré -> TL libéré + fiche annotée ; avéré -> alerte confirmée."""
    from openpyxl import load_workbook
    from tlone import vigilance as v
    tpl = tmp_path / "modele.xlsx"
    _modele_vigilance(tpl)
    ds = [dict(fichier="f1", ref_client="07310108547", client="KAYA AHMET", beneficiaire="KAYA AHMET", iban="TR1",
               swift="TEBUTRIS", banque="TEB", pays_banque="TURQUIE", individus=["MEHMET KAYA", "AYSE KAYA"]),
          dict(fichier="f2", ref_client="07310262001", client="DEMIR ALI", beneficiaire="DEMIR ALI", iban="TR2",
               swift="ZIRATRIS", banque="ZIRAAT", pays_banque="TURQUIE", individus=[])]
    units = v.build_units(ds)
    v.prepare(units, tpl, tmp_path / "lot", "TOSYALI")
    jp = tmp_path / "j.jsonl"
    v.append(jp, "u", "ENVOI", lot="L1", unites=units)
    u1 = next(u for u in units if u["fichiers"] == ["f1"])
    u2 = next(u for u in units if u["fichiers"] == ["f2"])
    r1 = _retour(tmp_path / "lot" / u1["fichier"], {"D32": "OK ==> RAS\nOK ==> HIT"})
    r2 = _retour(tmp_path / "lot" / u2["fichier"], {"D12": "OK ==> HIT"})
    v.import_returns([(u1["fichier"], r1), (u2["fichier"], r2)], v.read_journal(jp), jp, "u", tmp_path / "res")
    j = v.read_journal(jp)
    ret1 = next(e for e in j if e["action"] == "RETOUR" and e["uid"] == u1["uid"])
    assert Path(ret1["chemin"]).exists() and ret1["points"][0][2] == "AYSE KAYA"   # HIT sur la mère (rang 2)
    pt = ret1["points"][0]
    with pytest.raises(ValueError):                                                  # commentaire obligatoire
        v.analyse_points(jp, "u", u1["uid"], ["f1"], [{"rubrique": pt[0], "rang": pt[1], "valeur": pt[2],
                                                       "verdict": "NON_AVERE", "commentaire": ""}], None, tmp_path)
    ev = v.analyse_points(jp, "u", u1["uid"], ["f1"],
                          [{"rubrique": pt[0], "rang": pt[1], "valeur": pt[2], "verdict": "NON_AVERE",
                            "commentaire": "Homonyme : date de naissance différente"}],
                          Path(ret1["chemin"]), tmp_path / "analyses")
    assert ev["decision"] == "LEVEE"
    ws = load_workbook(ev["fiche_annotee"]).worksheets[0]
    assert "NON AVÉRÉ — AYSE KAYA : Homonyme" in ws["E32"].value
    ret2 = next(e for e in v.read_journal(jp) if e["action"] == "RETOUR" and e["uid"] == u2["uid"])
    p2 = ret2["points"][0]
    ev2 = v.analyse_points(jp, "u", u2["uid"], ["f2"], [{"rubrique": p2[0], "rang": 1, "valeur": p2[2],
                                                         "verdict": "AVERE", "commentaire": "Banque sous sanctions"}],
                           Path(ret2["chemin"]), tmp_path / "analyses")
    assert ev2["decision"] == "CONFIRMEE"
    st = v.statuses(v.read_journal(jp))
    assert st["f1"]["statut"] == "HIT_LEVE" and st["f2"]["statut"] == "HIT_CONFIRME"
    assert v.eligible(st) == {"f1"}


def test_preflight_worker_autre_dossier(env):
    """Worker lancé depuis un autre dossier de version (ex. tlone_app_v2 alors que l'App tourne en v3) : bloquant."""
    import dataclasses
    import json as _json
    import time as _time
    from tlone import preflight, worker
    _, s = env
    s = dataclasses.replace(s, run_mode="worker")
    hb = worker.heartbeat_path(s)
    hb.parent.mkdir(parents=True, exist_ok=True)
    hb.write_text(_json.dumps({"epoch": _time.time(), "ts": "x", "host": "ws", "state": "ATTENTE", "warm": {},
                               "steps": ["ov_extract"], "app_dir": "/mnt/code/tlone_app_v2", "env": {}}))
    rows = {r["contrôle"]: r for r in preflight.check_step(s, "ov_extract")}
    assert rows["Dossier du worker"]["niveau"] == preflight.KO


# ------------------------------------------------------------------ historique Atlas × CarthagoTrade
@pytest.mark.parametrize("objet,attendu", [
    ("052026", "2026-05"), ("052026P1", "2026-05P1"), ("052026P2", "2026-05P2"), ("62026", "2026-06"),
    ("062026P1", "2026-06P1"), ("20257P2", "2025-07P2"), ("202605 partie 1", "2026-05P1"), ("TR Octobre 2021", "2021-10"),
    ("112025", "2025-11"), ("TR NOVEMBRE 2025", "2025-11"), ("072025P1", "2025-07P1"), ("202607", "2026-07"),
    ("TRANCHE 2 JUIN 2026", "2026-06P2")])
def test_historique_periode_objet(objet, attendu):
    from tlone import historique as h
    p = h.parse_periode(objet)
    assert p["periode"] == attendu and p["statut"] == "OK", p


def test_historique_periode_jours_et_illisible():
    from tlone import historique as h
    p = h.parse_periode("TR Octobre 2025-17 JRS")
    assert p["mois"] == "2025-10" and p["jours"] == 17 and p["statut"] == "A_PRECISER"
    assert h.parse_periode("salaire")["statut"] == "ILLISIBLE" and h.parse_periode(None)["statut"] == "ABSENT"


def _trade_xlsx(path, rows):
    cols = ["Référence", "Code client", "Nom", "Montant", "Statut", "Nom2", "Nom3", "Objet du transfert",
            "Compte à débiter (Siège+ Racine+ Ordinal)", "Date valeur"]
    pd.DataFrame(rows, columns=cols).to_excel(path, index=False)


def test_historique_rapprochement_atlas_trade(tmp_path):
    from tlone import historique as h
    d = tmp_path / "historique_in"
    d.mkdir()
    _trade_xlsx(d / "Transfert libre -2026-04-23T13-05-44.xlsx", [
        ["TL21000001", "07310100497", "A B", "100 000,00", "Exécuté", "Turquie", "Allemagne", "052026", "0731010049700132DZD", "02/06/2026"],
        ["TL21000002", "07310100497", "A B", "50 000,00", "Exécuté", "Turquie", "Allemagne", "052026P1", "", "03/06/2026"],
        ["TL21000003", "07310108309", "C D", "70 000,00", "Exécuté", "Inde", "Inde", "062026", "", "01/07/2026"],
        ["TL21000004", "07310108309", "C D", "70 000,00", "Rejeté", "Inde", "Inde", "072026", "", "01/08/2026"],
        ["TL21000005", "07310108309", "C D", "70 000,00", "Exécuté", "Inde", "Inde", "TR Août 2026 - 12 jrs", "", "01/09/2026"]])
    pd.DataFrame([
        ["07310100497", "02/06/2026", "TRF TL21000001 SALAIRE", "052026", "x", "02/06/2026", "-100000,00"],
        ["07310100497", "05/06/2026", "TRF TL21000001", "052026", "x", "05/06/2026", "-100000,00"],      # double
        ["07310100497", "03/06/2026", "TL 21000002", "052026P1", "x", "03/06/2026", "-50000,00"],
        ["07310108309", "02/08/2026", "TL21000004", "072026", "x", "02/08/2026", "-70000,00"],          # rejeté débité
        ["07310108309", "02/08/2026", "frais", "", "x", "02/08/2026", "-10"],
        ["07310108309", "02/09/2026", "TL21000009", "082026", "x", "02/09/2026", "-1"]],               # hors Trade
        columns=["id_tiers", "Date_operation_Atlas", "lib_complementaire_mvt", "Reference_denotage", "Reference_client",
                 "Date_Valeur_comptable", "Montant_mvt_comptable"]).to_excel(
        d / "Historique_mouvement_altas-2026-10-20T101010.xlsx", index=False)
    pd.DataFrame([["07310", "100497", "1", "DZD", "052026", "-105000", "20260601", "20260830"]],
                 columns=["SIEGE", "RACINE", "ORDINAL", "DEVISE", "REFERENCE DENOTAGE DU MOUVEMENT", "MT BLOCAGE DEV CPT",
                          "DATE DEBUT BLOCAGE", "DATE LIMITE BLOCAGE"]).to_excel(d / "Blocage-2026-10-20T101010.xlsx", index=False)
    pd.DataFrame([["0731010920500191DZD", "1 000,00", "900,00"]], columns=["compte_local", "Solde_Compta", "Solde_Dispo"]) \
        .to_excel(d / "comptes_solde-2026-10-20T101010.xlsx", index=False)
    f = h.latest(d)
    assert set(f) == {"trade", "atlas", "blocage", "soldes"}
    t, a, b, so = h.load_trade(f["trade"]), h.load_atlas(f["atlas"]), h.load_blocage(f["blocage"]), h.load_soldes(f["soldes"])
    assert list(t["NATIONALITE"])[0] == "Turquie" and list(t["PAYS_DESTINATION"])[0] == "Allemagne"   # Nom2 / Nom3
    assert b.iloc[0]["CLIENT"] == "07310100497" and str(b.iloc[0]["DATE_LIMITE"]) == "2026-08-30"
    assert so.iloc[0]["CLIENT"] == "07310109205" and so.iloc[0]["SOLDE_DISPO"] == 900.0
    an = h.reconcile(a, t, b)["anomalies"]
    codes = {(r.CODE, r.REF_TL) for r in an.itertuples()}
    assert ("H01", "TL21000001") in codes                        # même référence débitée deux fois
    assert ("H08", "TL21000004") in codes                        # rejeté mais débité
    assert ("H05", "TL21000009") in codes                        # débit Atlas sans TL Trade
    assert ("H09", "TL21000003") in codes                        # exécuté sans débit Atlas
    assert any(c == "H12" for c, _ in codes)                     # 2026-05 complet + 2026-05P1
    assert ("H10", "TL21000005") in codes                        # 12 jours sans tranche
    assert any(c == "H02" for c, _ in codes)                     # mouvement sans référence
    assert an.iloc[0]["GRAVITE"] == "CRITIQUE"
    per = h.periodes_par_client(t, a)
    assert set(per["PERIODE"]) == {"2026-05", "2026-05P1", "2026-07"}         # seul le débit Atlas fait foi


def test_historique_colonne_compte_decalee():
    from tlone import historique as h
    t = pd.DataFrame({"CLIENT": ["07310100001", "07310100002", "07310100003", "07310100004"],
                      "COMPTE_A_DEBITER": ["", "07310100001", "07310100002", "07310100003"]})
    assert h.colonne_compte_decalee(t) == 3


def test_execution_tl_en_attente_blocages_employeur():
    """TL saisis non débités : statut, blocage à faire (solde jamais débiteur), doublons, fichier employeur."""
    from datetime import date as _d
    from tlone import execution as ex
    T = lambda ref, cl, per, mt, st="En cours d'exécution": {  # noqa: E731
        "REF_TL": ref, "CLIENT": cl, "NOM": "N " + cl[-3:], "PERIODE": per, "MOIS": per[:7], "TRANCHE": per[7:] or None,
        "STATUT_PERIODE": "OK", "OBJET": per, "MONTANT_TRADE": mt, "STATUT_TRADE": st}
    trade = pd.DataFrame([T("TL000000001", "07310100001", "2026-06", 100000.0),          # bloqué -> à exécuter
                          T("TL000000002", "07310100002", "2026-06P1", 50000.0),         # à bloquer, solde OK
                          T("TL000000003", "07310100002", "2026-06P2", 50000.0),         # P2 : pas un doublon ; solde KO
                          T("TL000000004", "07310100003", "2026-07", 10.0),              # pas de dossier
                          T("TL000000005", "07310100004", "2026-05", 10.0),              # doublon avec 6 (exécuté)
                          T("TL000000006", "07310100004", "2026-05P1", 10.0, "Exécuté"),
                          T("TL000000007", "07310100005", "2026-05", 10.0, "Rejeté")])
    atlas = pd.DataFrame({"REF_TL": ["TL000000006"]})
    sy = pd.DataFrame([{"FICHIER": f"f{i}", "SIEGE_RACINE": cl, "PERIODE_TL": per, "MONTANT_OV": mt}
                       for i, (cl, per, mt) in enumerate([("07310100001", "2026-06", 100000.0),
                                                          ("07310100002", "2026-06P1", 50000.0),
                                                          ("07310100002", "2026-06P2", 50000.0),
                                                          ("07310100004", "2026-05", 10.0)])])
    ch = pd.DataFrame(columns=["FICHIER", "CODE", "RESULTAT", "SEVERITE"])
    dec = pd.DataFrame({"FICHIER": ["f0", "f1", "f2", "f3"], "DECISION": ["VALIDE"] * 4})
    vig = {"f0": "RAS", "f1": "RAS", "f2": "HIT_LEVE", "f3": "RAS"}
    bl = pd.DataFrame([{"CLIENT": "07310100001", "PERIODE": "2026-06", "MONTANT_BLOCAGE": 105000.0,
                        "DATE_LIMITE": _d(2099, 1, 1), "REFERENCE_DENOTAGE": "062026"}])
    so = pd.DataFrame([{"CLIENT": "07310100002", "ORDINAL": "001", "SOLDE_DISPO": 80000.0}])
    pl = ex.plan(trade, atlas, bl, so, sy, ch, dec, vig).set_index("REF_TL")
    assert "TL000000006" not in pl.index and "TL000000007" not in pl.index    # exécuté / rejeté exclus
    assert pl.loc["TL000000001", "STATUT"] == "A_EXECUTER"
    assert pl.loc["TL000000002", "STATUT"] == "A_BLOQUER" and pl.loc["TL000000002", "SOLDE_APRES_BLOCAGE"] == 25000.0
    assert pl.loc["TL000000003", "STATUT"] == "SOLDE_INSUFFISANT"             # 25 000 restant < 55 000
    assert pl.loc["TL000000004", "STATUT"] == "COMPLEMENT" and "non scanné" in pl.loc["TL000000004", "MOTIFS"]
    assert pl.loc["TL000000005", "STATUT"] == "DOUBLON" and pl.loc["TL000000005", "DOUBLON_AVEC"] == "TL000000006"
    fb = ex.fichier_blocages(pl.reset_index(), _d(2026, 10, 20))
    assert len(fb) == 1 and fb.iloc[0]["SIEGE"] == "07310" and fb.iloc[0]["RACINE"] == "100002"
    assert fb.iloc[0]["REFERENCE DENOTAGE DU MOUVEMENT"] == "062026P1" and fb.iloc[0]["MT BLOCAGE DEV CPT"] == 55000.0
    assert fb.iloc[0]["DATE LIMITE BLOCAGE"] == "20270118"                    # + 90 jours
    fe = ex.fichier_employeur(pl.reset_index())
    assert set(fe["Situation"]) >= {"Sera exécuté", "En attente de complément", "TL en double — à annuler"}
    # « Exécuté » dans Trade sans débit Atlas : ANOMALIE (ni exécuté, ni en cours), pas un doublon pour les autres
    trade2 = pd.concat([trade, pd.DataFrame([T("TL000000008", "07310100006", "2026-04", 10.0, "Exécuté"),
                                             T("TL000000009", "07310100006", "2026-04", 10.0)])], ignore_index=True)
    pl2 = ex.plan(trade2, atlas, bl, so, sy, ch, dec, vig).set_index("REF_TL")
    assert pl2.loc["TL000000008", "STATUT"] == "ANOMALIE"
    assert pl2.loc["TL000000009", "STATUT"] != "DOUBLON" and "TL000000008" in pl2.loc["TL000000009", "ALERTE"]
    # début des blocages = J+1 de l'extraction (nom du fichier des soldes)
    from pathlib import Path as _P
    d0, why = ex.debut_blocage({"soldes": _P("comptes_solde-2026-10-20T101010.xlsx")})
    assert d0 == _d(2026, 10, 21) and "20/10/2026" in why


def test_suivi_planning_depuis_debut_contrat_reliquats():
    """Planning complet depuis le début du CTS : exécutions Atlas (même anciennes) rattachées, reliquats visibles."""
    from datetime import date as _d
    from tlone import execution as ex
    from tlone import transfers
    pl = pd.DataFrame([{"PERIODE_ID": f"p{m}", "DOM_ID": "d1", "SIEGE_RACINE": "07310100001", "NOM_TRAVAILLEUR": "X",
                        "NUMERO_DOMICILIATION": "D1", "PERIODE_TL": per, "MOIS_BASE": per[:7],
                        "DATE_FIN_SEGMENT": fin, "MONTANT_MAX_THEORIQUE": 100.0}
                       for m, per, fin in [(1, "2025-04P2", "2025-04-30"), (2, "2025-05", "2025-05-31"),
                                           (3, "2025-06", "2025-06-30"), (4, "2025-07", "2025-07-31"),
                                           (5, "2026-12", "2026-12-31")]])
    T = lambda ref, per, st: {"REF_TL": ref, "CLIENT": "07310100001", "NOM": "X", "PERIODE": per, "MOIS": per[:7],  # noqa: E731
                              "TRANCHE": per[7:] or None, "OBJET": per, "MONTANT_TRADE": 100.0, "STATUT_TRADE": st}
    trade = pd.DataFrame([T("TL1", "2025-04", "Exécuté"),           # mois partiel P2 au Planning : rattaché par le mois
                          T("TL2", "2025-05", "Exécuté"),           # « Exécuté » sans débit Atlas -> anomalie
                          T("TL3", "2025-06", "En cours d'exécution")])
    atlas = pd.DataFrame({"REF_TL": ["TL1"], "MONTANT_ATLAS": [-100.0], "DATE_OPERATION": [_d(2025, 5, 3)]})
    hreg = ex.registre_atlas(pl, trade, atlas)
    assert list(hreg["PERIODE_ID"]) == ["p1"] and hreg.iloc[0]["FICHIER"] == "ATLAS:TL1"
    sp = ex.suivi_planning(pl, transfers.load_registre(_P_none()), trade, atlas, today=_d(2026, 10, 21)).set_index("PERIODE_TL")
    assert sp.loc["2025-04P2", "STATUT_PERIODE"] == "TRANSFERE" and sp.loc["2025-04P2", "TL_EXECUTES"] == "TL1"
    assert sp.loc["2025-05", "STATUT_PERIODE"] == "ANOMALIE_TRADE" and sp.loc["2025-05", "RELIQUAT"] == 100.0
    assert sp.loc["2025-06", "STATUT_PERIODE"] == "SAISI_EN_ATTENTE"
    assert sp.loc["2025-07", "STATUT_PERIODE"] == "NON_TRANSFERE" and sp.loc["2025-07", "RELIQUAT"] == 100.0
    assert sp.loc["2026-12", "STATUT_PERIODE"] == "A_VENIR" and sp.loc["2026-12", "RELIQUAT"] == 0
    # registre TL-One + Atlas : le même TL (même période, même montant) n'est pas compté deux fois
    reg = pd.DataFrame([{"FICHIER": "ov1.pdf", "PERIODE_ID": "p1", "MONTANT": "100", "DECISION": "VALIDE"}])
    assert len(ex.registre_complet(reg, hreg)) == 1


def _P_none():
    from pathlib import Path as _P
    return _P("/nonexistent/registre.csv")


def test_v27_limite_n_porte_sur_les_pdf_non_extraits(tmp_path):
    """« Limiter à N » sur v27 : les N premiers PDF non encore en JSON (sinon le nouveau PDF n'est jamais lu)."""
    import json as _j
    from tlone import nbrunner
    nb = _j.load(open(Path(__file__).resolve().parents[1] / "notebooks" / "pipeline_ocr_vl27.ipynb", encoding="utf-8"))
    src = next("".join(c["source"]) for c in nb["cells"] if "JSON_DIR   =" in "".join(c["source"]))
    out, ap = nbrunner.patch_source(src, "ov_extract", str(tmp_path), 1, None)
    assert any("non extraits" in a for a in ap)
    (tmp_path / "transferts_in").mkdir()
    (tmp_path / "transferts_out" / "json_dossiers").mkdir(parents=True)
    for n in ("a.pdf", "b.pdf"):
        (tmp_path / "transferts_in" / n).write_bytes(b"%PDF")
    (tmp_path / "transferts_out" / "json_dossiers" / "a.json").write_text("{}")
    line = next(l for l in out.splitlines() if l.startswith("pdfs = "))
    ns = {"INPUT_DIR": tmp_path / "transferts_in", "JSON_DIR": tmp_path / "transferts_out" / "json_dossiers"}
    exec(line, ns)
    assert [p.name for p in ns["pdfs"]] == ["b.pdf"]


def test_completude_lot_ov(env):
    from tlone import ui_ov
    p, _s = env
    p.tr_in.mkdir(parents=True, exist_ok=True)
    p.tr_out.mkdir(parents=True, exist_ok=True)
    for n in ("a.pdf", "b.pdf", "c.pdf"):
        (p.tr_in / n).write_bytes(b"%PDF")
    (p.tr_out / "pipeline.log").write_text("2026-10-21 10:00:00 — [   2/3] ❌ b.pdf — CUDA out of memory\n", encoding="utf-8")
    m = ui_ov.completude(p, pd.DataFrame({"fichier": ["a.pdf"]})).set_index("PDF")
    assert "CUDA out of memory" in m.loc["b.pdf", "Raison"] and "Pas encore extrait" in m.loc["c.pdf", "Raison"]


# ------------------------------------------------------------------ V2 : multi-employeur, base RH, BOCI, démo
def test_multi_employeur_registre_et_settings(tmp_path, monkeypatch):
    from tlone import employeurs
    from tlone.config import load_settings
    home = tmp_path / "pipeline"
    (home / "employeurs" / "TOSYALI").mkdir(parents=True)
    monkeypatch.setenv("TLONE_DATA_ROOT", str(home / "employeurs" / "TOSYALI"))
    s = load_settings()
    assert s.employeur == "TOSYALI" and s.home == home
    employeurs.ensure_registered(home, "TOSYALI", "TOSYALI IRON STEEL INDUSTRY ALGERIE")
    row = employeurs.create(home, "Sarl Aciéries de l'Ouest", "001631000123456", "testeur")
    assert row["CODE"] == "ACIERIES" and (home / "employeurs" / "ACIERIES" / "transferts_in").is_dir()
    with pytest.raises(ValueError):
        employeurs.create(home, "Autre", "001631000123456", "testeur")          # NIF déjà enregistré
    with pytest.raises(ValueError):
        employeurs.create(home, "Sans NIF", "", "testeur")
    s2 = load_settings("ACIERIES")
    assert s2.paths.data_root == home / "employeurs" / "ACIERIES" and s2.employeur_raison_sociale.startswith("SARL ACI")
    assert set(employeurs.search(employeurs.load(home), "1631")["CODE"]) == {"ACIERIES"}


def test_worker_file_commune_et_dossier_de_l_employeur(tmp_path, monkeypatch):
    """Un seul worker pour tous les employeurs : la demande porte le dossier de données de son employeur."""
    import dataclasses
    import sys as _sys
    import nbformat
    from tlone import jobs, preflight, worker
    from tlone.config import load_settings
    home = tmp_path / "h"
    for c in ("A", "B"):
        (home / "employeurs" / c / "domiciliations_in").mkdir(parents=True)
        (home / "employeurs" / c / "domiciliations_in" / f"{c}.pdf").write_bytes(b"%PDF")
    monkeypatch.setenv("TLONE_DATA_ROOT", str(home / "employeurs" / "A"))
    nb = nbformat.v4.new_notebook()
    nb.cells = [nbformat.v4.new_code_cell("from pathlib import Path\nINPUT_DIR = Path('/mnt/data/domiciliations_in')\n"
                                          "pdfs = sorted(INPUT_DIR.glob('*.pdf'))"),
                nbformat.v4.new_code_cell("# from_pretrained\nX = 1"),
                nbformat.v4.new_code_cell("if not pdfs:\n    pass\nelse:\n    (INPUT_DIR.parent / 'vu.txt').write_text(','.join(p.name for p in pdfs))")]
    fake = tmp_path / "nb.ipynb"
    nbformat.write(nb, str(fake))
    sa = dataclasses.replace(load_settings(), notebooks={**load_settings().notebooks, "dom_extract": fake}, run_mode="worker")
    sb = dataclasses.replace(load_settings("B"), notebooks=sa.notebooks, run_mode="worker")
    assert worker.queue_dir(sa) == worker.queue_dir(sb)
    w = worker.Worker(sa, ["dom_extract"], pythons={"dom_extract": _sys.executable})
    monkeypatch.setattr(preflight, "check_step", lambda s_, step: [])
    jobs.launch(sb, "dom_extract", None, "t")
    w.loop(once=True)
    assert (home / "employeurs" / "B" / "vu.txt").read_text() == "B.pdf"
    assert not (home / "employeurs" / "A" / "vu.txt").exists()
    w.stop()


def _rh_xlsx(path, rows):
    from tlone.ui_referentiels import RH_TEMPLATE
    pd.DataFrame([{**{c: None for c in RH_TEMPLATE}, **r} for r in rows]).to_excel(path, index=False)


def test_base_rh_lecture_controles_et_rapprochement(tmp_path):
    from tlone import rh
    base = {"Matricule SAP": "1030", "NOM ET PRENOM": "SOZER MEHMET", "Salaire Net": 759882.58,
            "Compte bancaire (Algérie)": "02700731010865400117", "Compte devise": "TR15 0006 7010 0000 0054 7919 27",
            "SWIFT/BIC": "YAPITRIS", "Devise du transfert": "USD", "MOIS DE PAIEMENT": "SEPTEMBRE-2026",
            "Nom du bénéficiaire": "SOZER", "Prénom du bénéficiaire": "MEHMET", "Fin de validité": "17/09/2026"}
    _rh_xlsx(tmp_path / "rh.xlsx", [
        dict(base, **{"Part transférable": 409070.12, "Part Payable": 21530.01, "Nombre de jours": 17, "PART": "PART 1"}),
        dict(base, **{"Part transférable": 312818.33, "Part Payable": 16464.12, "Nombre de jours": 13, "PART": "PART 2",
                      "Fin de validité": "17/09/2028"}),
        dict(base, **{"Matricule SAP": "1543", "NOM ET PRENOM": "KUCUK FATIH", "Compte bancaire (Algérie)": "02700731011083100127",
                      "Salaire Net": 456369.97, "Part transférable": 146428.87, "Part Payable": 20906.78, "Nombre de jours": 11,
                      "PART": "PART 1", "Compte devise": "TR70 0006 7010 0000 0023 0405 09"})])
    df = rh.read(tmp_path / "rh.xlsx")
    assert list(df["PERIODE"]) == ["2026-09P1", "2026-09P2", "2026-09P1"] and df["CLIENT"].iloc[0] == "07310108654"
    assert df["SWIFT"].iloc[0] == "YAPITRIS" and rh.swift11(df["SWIFT"].iloc[0]) == "YAPITRISXXX" and rh.iban_ok(df["IBAN"].iloc[0])
    q = rh.dq(df)
    assert set(q[q["MATRICULE"] == "1543"]["CODE"]) >= {"RH04"} and "RH04" not in set(q[q["MATRICULE"] == "1030"]["CODE"])
    r, how = rh.match(df, "07310108654", "2026-09", "P2")
    assert r["PART_TRANSFERABLE"] == 312818.33 and "part" in how
    r, how = rh.match(df, "07310108654", "2026-09", None)                  # OV du mois entier : P1 + P2
    assert r["PART_TRANSFERABLE"] == round(409070.12 + 312818.33, 2) and "additionnées" in how
    assert rh.iban_ok("TR57000320000000097105886") is False                # 25 caractères au lieu de 26


def test_atf_au_format_bnp_rapproche_par_matricule(tmp_path):
    from tlone import rh, transfers
    pd.DataFrame([{"SIRA": 1, "MATRICULE SAP": 1030, "Nom et Prenom": "SOZER MEHMET", "Adresse": "X", "NET IMPOSABLE": 1,
                   "IRG": 1, "SALAIRE NET": 759882.58, "MONTANT TRANSFERABLE": 721888.45, "BANQUE": "AGENCE",
                   "periode": pd.Timestamp(2026, 9, 1)}]).to_excel(tmp_path / "ATF_BNP_09_2026.xlsx", index=False)
    rhdf = pd.DataFrame({"MATRICULE": ["1030"], "RIB": ["02700731010865400117"]})
    a = transfers.load_referential(tmp_path / "ATF_BNP_09_2026.xlsx", transfers.ATF_COLUMNS, rhdf)
    assert a["_MOIS"].iloc[0] == "2026-09" and a["_COMPTE"].iloc[0] == "02700731010865400117"
    assert float(a["MONTANT_TRANSFERABLE"].iloc[0]) == 721888.45 and a["NOM_PRENOM"].iloc[0] == "SOZER MEHMET"
    (tmp_path / "vide.html").write_text("<html><body><iframe src='x'></iframe></body></html>")
    with pytest.raises(ValueError):
        transfers.load_referential(tmp_path / "vide.html", transfers.AUT_COLUMNS)
    (tmp_path / "aut.html").write_text("<table><tr><th>Nom et prénom</th><th>Montant autorisé</th><th>Mois</th></tr>"
                                       "<tr><td>SOZER MEHMET</td><td>721 888,45</td><td>09/2026</td></tr></table>")
    b = transfers.load_referential(tmp_path / "aut.html", transfers.AUT_COLUMNS)
    assert b["_MOIS"].iloc[0] == "2026-09" and b["NOM_PRENOM"].iloc[0] == "SOZER MEHMET"


def _ligne(**kw):
    l = dict(REF_TL="TL26012052", CLIENT="07310109737", PERIODE="2026-05P1", MONTANT=410625.17, DEVISE="USD",
             NOM_CLIENT="TOKA SERCAN", ADRESSE_CLIENT="BASE DE VIE 31015 BETHIOUA", DATE_DEMANDE="24/05/2026",
             PREDOM="2024|07310|000300", DOMICILIATION="271901202424000172DZD", DATE_DOMICILIATION="25/06/2024",
             BENEFICIAIRE="TOKA SERCAN", ADRESSE_BENEF="MADENLER MAH. 1039 SK. YATAGAN / MUGLA",
             IBAN="TR150006701000000054791927", SWIFT="TEBUTRISXXX", BANQUE="TEB/TURK EKONOMI BANK", PAYS="Turquie",
             SIEGE="07310", RACINE="109737", ORDINAL="001", CLE="32", COMPTE="0731010973700132DZD")
    l.update(kw)
    return l


def test_boci_generation_formats_et_valeurs_fixes():
    from datetime import date as _d
    from tlone import boci
    L = pd.DataFrame([_ligne(), _ligne(REF_TL="TL26015369", CLIENT="07310109137", PERIODE="2026-06P1", DEVISE="EUR",
                                       SWIFT="HDFCINBBXXX", IBAN="50100192206546", BANQUE="HDFC BANK LTD", PAYS="Inde",
                                       RACINE="109137", CLE="03", COMPTE="0731010913700103DZD")])
    out = boci.generer(L, _d(2026, 7, 20))
    t, d, r = out["trade"], out["dom"], out["reglement"]
    assert list(t.columns) == boci.TRADE_COLUMNS and list(d.columns) == boci.DOM_COLUMNS and list(r.columns) == boci.REG_COLUMNS
    assert t.loc[0, "Reference denotage"] == "2026-05P1" and t.loc[0, "Objet du transfert"] == "052026P1"
    assert t.loc[0, "Agence"] == "BNP/07310" and t.loc[0, "Nature du transfert"] == "Tranfert International"
    assert d.loc[0, "Description"] == "Salaire 2026-05P1" and d.loc[0, "Balance Payment Codification"] == "331"
    assert r.loc[0, "REF_DENO."] == "052026P1" and r.loc[0, "DATE VAL"] == "19/07/2026" and r.loc[0, "Value Date"] == "20/07/2026"
    assert r.loc[0, "CorrespondantSiege"] == "" and r.loc[1, "CorrespondantSiege"] == "09990"    # EUR seulement
    assert r.loc[1, "INTERMEDIAIRE FINAN"] == "CHASDEFXXXX"
    assert boci.verifier(out).empty                               # cohérents par construction
    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(boci.to_xlsx("trade", t))).active
    assert ws.title == "Trade_Consolidé" and ws["A2"].value == "07310109737" and ws["I2"].number_format == "#,##0.00"
    assert boci.file_names(26)["reglement"].startswith("Reglement_FINAL_BOCI_26TL_")


def test_boci_verifier_detecte_les_ecarts_entre_fichiers():
    from datetime import date as _d
    from tlone import boci
    out = boci.generer(pd.DataFrame([_ligne(), _ligne(REF_TL="TL2", CLIENT="07310109738", RACINE="109738", CLE="05",
                                                      COMPTE="0731010973800105DZD")]), _d(2026, 7, 20))
    out["reglement"].loc[0, "REF_DENO."] = "202605P1"                       # format AAAAMM
    out["dom"].loc[0, "Description"] = "Salaire 2026-05"                    # période différente
    out["reglement"].loc[0, "NUM CONTE BEN"] = "TR15000670100000005479192"   # IBAN tronqué
    out["trade"].loc[0, "Code CRDS"] = "MOCI"                               # valeur fixe
    out["reglement"].loc[1, "CorrespondantSiege"] = "09990"                 # correspondant sur un TL en USD
    an = boci.verifier(out)
    codes = set(an["CODE"])
    assert {"B05", "C05", "C09", "F02", "C18"} <= codes
    assert (an[an["CODE"] == "B05"]["ATTENDU"] == "BOCI").all()
    assert set(an[an["CODE"] == "C05"]["REF_TL"]) == {"TL26012052"}


def test_controles_ov_base_rh(tmp_path):
    """R01–R07 : l'OV est contrôlé par la base RH ; un IBAN différent ou un montant au-dessus de la part RH ressort."""
    from tlone import rh as _rh
    rhdf = pd.DataFrame([{"MATRICULE": "1", "NOM_PRENOM": "KAYA AHMET", "CLIENT": "07310108259", "MOIS": "2026-03",
                          "PART": None, "PERIODE": "2026-03", "PART_TRANSFERABLE": 312255.0, "IBAN": "TR330006100519786457841326",
                          "SWIFT": "TEBUTRISXXX", "DEVISE": "USD", "BENEFICIAIRE": "KAYA AHMET", "PERMIS_FIN": None}])
    s = load_settings()
    d = {"fichier": "x.pdf", "OV.montant_chiffres": 312255.0, "OV.periode": "MARS 2026", "OV.compte_debiteur": "02700731010825900165",
         "OV.beneficiaire_nom": "KAYA AHMET", "OV.beneficiaire_compte": "TR33 0006 1005 1978 6457 8413 26", "OV.swift": "TEBUTRIS"}
    _, rows = transfers.control_dossier(d, pd.DataFrame(), pd.DataFrame(columns=["PERIODE_ID", "DOM_ID", "FICHIER", "_MONTANT"]),
                                        pd.DataFrame(), pd.DataFrame(), s, rh=rhdf)
    c = {r["CODE"]: r for r in rows}
    assert c["R01"]["RESULTAT"] == c["R02"]["RESULTAT"] == c["R03"]["RESULTAT"] == c["R04"]["RESULTAT"] == "OK"
    d2 = dict(d, **{"OV.montant_chiffres": 330000.0, "OV.beneficiaire_compte": "TR330006100519786457841399"})
    _, rows = transfers.control_dossier(d2, pd.DataFrame(), pd.DataFrame(columns=["PERIODE_ID", "DOM_ID", "FICHIER", "_MONTANT"]),
                                        pd.DataFrame(), pd.DataFrame(), s, rh=rhdf)
    c = {r["CODE"]: r for r in rows}
    assert c["R02"]["RESULTAT"] == "KO" and c["R02"]["SEVERITE"] == "BLOQUANT" and c["R03"]["RESULTAT"] == "KO"


def test_ov_p2_non_rattache_a_la_p1_de_l_ancienne_dom():
    s = load_settings()
    pl = pd.DataFrame([{"PERIODE_ID": "p1", "DOM_ID": "d", "SIEGE_RACINE": "07310108259", "MOIS_BASE": "2026-09",
                        "PARTIE": "P1", "PERIODE_TL": "2026-09P1", "MONTANT_MAX_THEORIQUE": 100.0, "NUMERO_DOMICILIATION": "1",
                        "DOM_COMPTE_LOCAL": "02700731010825900165", "DATE_DEBUT_SEGMENT": "2026-09-01",
                        "DATE_FIN_SEGMENT": "2026-09-17", "NOM_TRAVAILLEUR": "KAYA AHMET"}])
    d = {"fichier": "07310-108259-202609P2.pdf", "OV.montant_chiffres": 50.0, "OV.periode": "SEPTEMBRE 2026 P2",
         "OV.compte_debiteur": "02700731010825900165"}
    sy, rows = transfers.control_dossier(d, pl, pd.DataFrame(columns=["PERIODE_ID", "DOM_ID", "FICHIER", "_MONTANT"]),
                                         pd.DataFrame(), pd.DataFrame(), s)
    assert sy["PERIODE_ID"] is None and {r["CODE"]: r for r in rows}["T04"]["RESULTAT"] == "KO"


def test_demo_bout_en_bout(tmp_path, monkeypatch):
    """Employeur de démonstration : TL prêts, fichiers BOCI sans écart, anomalies voulues présentes."""
    from datetime import date as _d
    from tlone import boci, cockpit, demo, execution, historique as h, planning, referentiels, vigilance as vg
    home = tmp_path / "home"
    monkeypatch.setenv("TLONE_DATA_ROOT", str(home / "employeurs" / "TOSYALI"))
    demo.build(home, today=_d(2026, 10, 21))
    s = load_settings("DEMO")
    P = s.paths
    A = transfers.read_audit_xlsx(transfers.latest_audit(P.tr_out)[0])
    PL = planning.load_planning(P.planning_csv)
    R = referentiels.load_all(P)
    sy, ch = transfers.control_each(A, PL, transfers.load_registre(P.registre_csv), R["atf"], R["aut"], s, None, R["rh"])
    st = dict(zip(sy["FICHIER"], sy["STATUT_CONTROLE"]))
    assert st["07310-990104-202609.pdf"] == "NON_CONFORME" and st["07310-990105-202609.pdf"] == "A_REVOIR"
    assert st["07310-990112-202609P2.pdf"] == "NON_CONFORME" and st["07310-990101-202609.pdf"] == "CONFORME"
    f = h.latest(P.historique_dir)
    d = {k: getattr(h, "load_" + k)(f[k]) for k in f}
    vig = {k: v["statut"] for k, v in vg.statuses(vg.read_journal(vg.journal_path(P))).items()}
    assert "HIT" in vig.values()
    pl = execution.plan(d["trade"], d["atlas"], d["blocage"], d["soldes"], sy, ch, transfers.load_decisions(P.decisions_csv),
                        vig, planning=PL)
    cnt = pl["STATUT"].value_counts()
    assert cnt.get("A_EXECUTER", 0) >= 4 and cnt.get("A_BLOQUER", 0) == 1 and cnt.get("SOLDE_INSUFFISANT", 0) == 1
    assert cnt.get("ANOMALIE", 0) == 1
    L = boci.preparer(pl[pl["STATUT"] == "A_EXECUTER"], d["trade"], sy, A, R["rh"], PL,
                      boci.carthago_dom_dates(P.carthago_xlsx), boci.load_params(P))
    assert (L["BLOQUANTS"] == "").all() and L["DATE_DOMICILIATION"].notna().all()
    assert boci.verifier(boci.generer(L, _d(2026, 10, 25)), None, PL).empty
    K = cockpit.indicateurs(P, s)
    assert K["prets"] == cnt["A_EXECUTER"] and K["salaries"] == 12 and K["nb_controles"] > 60


def test_boci_date_valeur_vide_et_name_client_rh():
    from tlone import boci
    out = boci.generer(pd.DataFrame([_ligne(NOM_CLIENT="YASIN KHAN YASIN KHAN", NOM_CLIENT_RH="YASIN KHAN")]))
    t, d, r = out["trade"].iloc[0], out["dom"].iloc[0], out["reglement"].iloc[0]
    assert t["Date valeur"] == d["Date valeur"] == r["Value Date"] == r["DATE VAL"] == r["DATE VAL1"] == ""
    assert r["Name CLIENT"] == "YASIN KHAN" and t["Nom"] == "YASIN KHAN YASIN KHAN"
    assert boci.verifier(out).empty


def test_vigilance_hit_deja_analyse_reconnu_le_mois_suivant(tmp_path):
    """Un HIT jugé non avéré est reconnu quand il ressort (même client / rubrique / valeur) : reconduction tracée.
    Une alerte sur une banque vaut pour tous les clients ; un nouveau point n'est jamais reconduit."""
    from tlone import vigilance as vg
    jp = tmp_path / "j.jsonl"
    u1 = {"uid": "A1", "ref_client": "07310108259", "client": "KAYA AHMET", "fichiers": ["sept.pdf"]}
    vg.append(jp, "x", "ENVOI", lot="L1", unites=[u1])
    vg.append(jp, "ana", "DECISION", uid="A1", fichiers=["sept.pdf"], decision="LEVEE", ref_client="07310108259",
              verdicts=[{"rubrique": "Bénéficiaire", "valeur": "KAYA AYSE", "verdict": "NON_AVERE", "commentaire": "épouse"},
                        {"rubrique": "Banque du bénéficiaire", "valeur": "HDFC BANK", "verdict": "NON_AVERE",
                         "commentaire": "banque correspondante"}])
    later = "2999-01-01T00:00:00"
    mem = vg.memoire_hits(vg.read_journal(jp))
    pts = [["Bénéficiaire", 1, "Kaya Aysé", "HIT", "HIT", "", ""]]
    r = vg.reconduire(mem, "07310 108259", pts, later)
    assert r and r[0]["verdict"] == "NON_AVERE" and "épouse" in r[0]["commentaire"] and r[0]["reconduit_de"]
    assert vg.reconduire(mem, "07310999999", pts, later) is None                     # autre client
    assert vg.reconduire(mem, "07310999999", [["Banque du bénéficiaire", 1, "HDFC BANK", "HIT", "HIT"]], later)  # global
    assert vg.reconduire(mem, "07310108259", pts + [["Divers individus sur les documents", 1, "X", "HIT", "HIT"]], later) is None
    ev = vg.analyse_points(jp, "ana2", "A1", ["oct.pdf"], r, None, tmp_path)
    assert ev["decision"] == "LEVEE" and ev["reconduit"] is True
    fp = vg.registre_faux_positifs(vg.read_journal(jp))
    assert int(fp[fp["Rubrique"] == "Bénéficiaire"]["Nombre d'analyses"].iloc[0]) == 2
    assert int(fp[fp["Rubrique"] == "Bénéficiaire"]["Dont reconduites"].iloc[0]) == 1


# ------------------------------------------------------------ v4 : autorisations du portail, régularisation, v29
_PORTAIL_TABLE = (
    "<html><body><div>Personnel étranger : Transfert des salaires des expatriés</div>"
    "<div>Code banque: 027 — Code agence: 00795</div>"
    "<p>NIF</p><p>000731010943515</p><p>RC</p><p>31 B 0000000-00/31</p><p>Raison sociale</p>"
    "<p>SPA TOSYALI IRONSTEEL INDUSTRY ALGERIE</p><p>Total de la demande</p><p>821 084,26 DZD</p><div>Articles (2)</div>"
    "<table><tr><th>N°</th><th>Désignation</th><th>Stock</th><th>En cours</th><th>Qté</th><th>Unité</th>"
    "<th>Prix unitaire</th><th>Total</th><th>Actions</th></tr>"
    "<tr><td>835</td><td>KAYA AHMET (AOUT)</td><td>—</td><td>—</td><td>1.0000</td><td>mois</td>"
    "<td>410 542,13 DZD</td><td>410 542,13 DZD</td><td><button>Domicilier</button></td></tr>"
    "<tr><td>836</td><td>ERHAN ALI (AOUT)</td><td>—</td><td>—</td><td>1.0000</td><td>mois</td>"
    "<td>410 542,13 DZD</td><td>410 542,13 DZD</td><td><button>Domicilier</button></td></tr></table></body></html>")


def test_autorisations_portail_tableau_texte_et_nif():
    from tlone import autorisations as au
    assert au.est_portail(_PORTAIL_TABLE)
    r = au.lire(_PORTAIL_TABLE, nom_fichier="Autorisation_Tosyali_2026_aout.html")
    e, L = r["entete"], r["lignes"]
    assert e["NIF"] == "000731010943515" and e["NB_ARTICLES"] == 2 and e["TOTAL_DEMANDE"] == pytest.approx(821084.26)
    assert e["CODE_BANQUE"] == "027" and e["CODE_AGENCE"] == "00795"
    assert list(L["MOIS"]) == ["2026-08", "2026-08"] and list(L["N"]) == [835, 836]
    assert L["MONTANT_AUTORISE"].sum() == pytest.approx(821084.26) and r["controles"] == []
    # même page rendue sans <table> (divs / spans) : lecture par le texte
    txt = _PORTAIL_TABLE.replace("<table>", "<div>").replace("</table>", "</div>").replace("<tr>", "<div>") \
        .replace("</tr>", "</div>").replace("<td>", "<span>").replace("</td>", "</span>") \
        .replace("<th>", "<span>").replace("</th>", "</span>")
    r2 = au.lire(txt, annee=2026)
    assert list(r2["lignes"]["NOM_PRENOM"]) == ["KAYA AHMET", "ERHAN ALI"] and r2["controles"] == []
    # année absente, total faux, NIF d'un autre employeur
    bad = au.lire(_PORTAIL_TABLE.replace("821 084,26", "900 000,00"))
    assert any("Année absente" in c for c in bad["controles"]) and any("total de la demande" in c for c in bad["controles"])
    _, ctl = au.vers_referentiel(r, nif_employeur="000016099999999")
    assert ctl and ctl[0].startswith("NIF de la demande")
    _, ctl = au.vers_referentiel(r, nif_employeur="000731010943515")
    assert ctl == []


def test_autorisations_portail_referentiel_json_et_compte_par_nom_rh(tmp_path):
    import json as _json
    f = tmp_path / "autorisations_aout.html"
    f.write_text(_PORTAIL_TABLE, encoding="utf-8")
    f.with_name(f.name + ".json").write_text(_json.dumps({"annee": 2026, "date_validation": "09/09/2026"}), encoding="utf-8")
    rh = pd.DataFrame([{"NOM_PRENOM": "KAYA AHMET", "RIB": "02700731010825900165", "MATRICULE": "1"}])
    df = transfers.load_referential(f, transfers.AUT_COLUMNS, rh)
    assert list(df["_MOIS"]) == ["2026-08", "2026-08"]
    k = df[df["NOM_PRENOM"] == "KAYA AHMET"].iloc[0]
    assert k["_COMPTE"] == "02700731010825900165" and str(k["DATE_VALIDATION"]).startswith(("09/09/2026", "2026-09-09"))
    assert df[df["NOM_PRENOM"] == "ERHAN ALI"]["COMPTE"].map(transfers.is_blank).all()


def test_regularisation_motifs_et_constats_sur_la_demo(tmp_path, monkeypatch):
    """Historique réutilisé : chaque période non transférée a un motif, chaque TL exécuté est contrôlé a posteriori."""
    from datetime import date as _d
    from tlone import demo, historique as h, referentiels, regularisation as rg, vigilance as vg
    home = tmp_path / "home"
    monkeypatch.setenv("TLONE_DATA_ROOT", str(home / "employeurs" / "TOSYALI"))
    demo.build(home, today=_d(2026, 10, 21))
    s = load_settings("DEMO")
    P = s.paths
    PL = planning.load_planning(P.planning_csv)
    R = referentiels.load_all(P)
    f = h.latest(P.historique_dir)
    d = {k: getattr(h, "load_" + k)(f[k]) for k in f}
    vig = {k: v["statut"] for k, v in vg.statuses(vg.read_journal(vg.journal_path(P))).items()}
    assert len(rg.audits_historiques(P)) >= 2                       # lot courant + extraction OV d'un mois passé
    df = rg.analyser(PL, d["trade"], d["atlas"], transfers.load_registre(P.registre_csv),
                     rg.index_ov(rg.audits_historiques(P)), R["atf"], R["aut"], R["rh"],
                     transfers.load_decisions(P.decisions_csv), today=_d(2026, 10, 21), vigilance=vig)
    by = {(n.split()[0], p): r for n, p, r in zip(df["NOM_TRAVAILLEUR"], df["PERIODE_TL"], df.to_dict("records"))}
    assert by[("SHARMA", "2026-05")]["CODE_MOTIF"] == "M03"         # OV rejeté (extraction d'un mois passé)
    assert by[("SHARMA", "2026-06")]["CODE_MOTIF"] == "M01"         # jamais scanné …
    assert "absent de l'ATF" in by[("SHARMA", "2026-06")]["MOTIFS"]  # … et absent de l'ATF du mois
    assert by[("TASKIN", "2026-07")]["CODE_MOTIF"] == "M10"         # « Exécuté » Trade sans débit Atlas
    assert by[("PATEL", "2026-09")]["CODE_MOTIF"] == "M12"          # HIT non levé
    assert by[("GUNES", "2026-09")]["CODE_MOTIF"] == "M04"          # OV au-dessus du plafond, non validé
    assert "supérieur à la part" in by[("ERDEM", "2026-04")]["CONSTATS"]   # double débit
    assert (df[df["ETAT"] == "TRANSFERE"]["CODE_MOTIF"] == "").all()
    sy = rg.synthese(df)
    assert set(sy["Motif"]) >= {rg.MOTIFS[c] for c in ("M01", "M03", "M04", "M10", "M12")}


def _nb_cells(name):
    import json as _json
    nb = _json.load(open(Path(__file__).resolve().parent.parent / "notebooks" / name, encoding="utf-8"))
    return {i: "".join(c["source"]) for i, c in enumerate(nb["cells"]) if c["cell_type"] == "code"}


def test_notebook_ov_v29_montants_champs_et_pages():
    import re as _re
    cells = _nb_cells("pipeline_ocr_vl29.ipynb")
    src = next(c for c in cells.values() if "def norm_montant" in c)
    g = {"re": _re}
    exec(src, g)
    nm = g["norm_montant"]
    for v, e in [("23.340.43", 23340.43), ("312:255:00", 312255.0), ("410 542,13", 410542.13), ("1.000", 1000.0),
                 ("1.234.567,89", 1234567.89), ("1,234,567.89", 1234567.89), ("312 255,00 DZD", 312255.0),
                 ("1,5", 1.5), (None, None), ("", None), (12, 12.0)]:
        assert nm(v) == e, v
    assert "nature_paiement_case" in g["CHAMPS_OV"] and "nombre_jours_absence" in g["CHAMPS_ANN2"]
    d = g["normalise_doc"]("ANNEXE_II", {"type": "ANNEXE_II", "nombre_jours_absence": "2", "salaire_mensuel": "23.340.43"})
    assert d["nombre_jours_absence"] == "2" and d["salaire_mensuel"] == 23340.43
    allsrc = "\n".join(cells.values())
    for lab in ("Pages par pièce", "Nature paiement (case)", "Nombre jours absence", "Date de naissance", "Résidence"):
        assert f"'{lab}'" in allsrc and any(k[1] == lab for k in transfers.V27_LABELS)
    assert "json_dossiers_v29" in allsrc and "data['page'] = i + 1" in allsrc


def test_notebook_ov_v30_beneficiaire_et_devise():
    """v30 : v29 inchangé ; nom du bénéficiaire après « Nom/Prénom », devise seule déduite de la case Autre."""
    import re as _re
    v29, v30 = _nb_cells("pipeline_ocr_vl29.ipynb"), _nb_cells("pipeline_ocr_vl30.ipynb")
    assert "json_dossiers_v29" in "\n".join(v29.values())                  # la version validée n'est pas modifiée
    allsrc = "\n".join(v30.values())
    assert "json_dossiers_v30" in allsrc and 'PIPELINE_VERSION = "OCR_OV_V30"' in allsrc
    assert 'APRÈS le libellé imprimé "Nom/Prénom"' in allsrc
    src = next(c for c in v30.values() if "def normalise_doc" in c)
    g = {"re": _re}
    exec(src, g)
    d = g["normalise_doc"]("OV", {"type": "OV", "beneficiaire_nom": "Nom/Prénom :          KAYA   AYSE",
                                  "nature_paiement_autre_libelle": "Transfert en USD",
                                  "beneficiaire_compte": "TR33 0006 1005 1978"})
    assert d["beneficiaire_nom"] == "KAYA AYSE" and d["devise_transfert"] == "USD"
    assert d["nature_paiement_autre_libelle"] == "Transfert en USD" and d["beneficiaire_compte"] == "TR33000610051978"
    assert g["devise_seule"]("transfert en euros") == "EUR" and g["devise_seule"]("Autre") is None
    for lab in ("Devise de transfert", "Nature paiement (autre libellé)"):
        assert f"'{lab}'" in allsrc and any(k[1] == lab for k in transfers.V27_LABELS)


def test_page_de_la_piece_et_r08_jours():
    from tlone.ui_ov import page_piece
    assert page_piece("OV:1, ANNEXE_I:2, ANNEXE_II:4+6, BULLETIN:5", "ANNEXE_II") == 4
    assert page_piece("OV:1", "BULLETIN") is None and page_piece(None, "OV") is None
    s = load_settings()
    rhdf = pd.DataFrame([{"MATRICULE": "1", "NOM_PRENOM": "KAYA AHMET", "CLIENT": "07310108259", "MOIS": "2026-03",
                          "PART": "P1", "PERIODE": "2026-03P1", "PART_TRANSFERABLE": 100.0, "IBAN": None, "SWIFT": None,
                          "DEVISE": None, "BENEFICIAIRE": None, "PERMIS_FIN": None, "NB_JOURS": 15, "RIB": None}])
    base = {"fichier": "x.pdf", "OV.montant_chiffres": 100.0, "OV.periode": "MARS 2026 P1",
            "OV.compte_debiteur": "02700731010825900165"}
    reg = pd.DataFrame(columns=["PERIODE_ID", "DOM_ID", "FICHIER", "_MONTANT"])
    # nombre de jours : plus comparé (décision du 24/10/2026)
    for jours in ("15", "17 jours"):
        _, rows = transfers.control_dossier(dict(base, **{"ANNEXE_II.nombre_jours": jours}), pd.DataFrame(), reg,
                                            pd.DataFrame(), pd.DataFrame(), s, rh=rhdf)
        assert "R08" not in {r["CODE"] for r in rows}


def test_ov_devise_iban_banque_prorata_et_date_domiciliation(tmp_path):
    """Devise lue dans « Transfert en USD », IBAN / banque comparés OV / Annexe II / RH, ATF ≈ part du contrat
    proratisée aux jours de la partie, date limite d'autorisation > date de domiciliation, part du Planning
    inchangée + part restante, OV hors du contrôle d'identité."""
    from datetime import date
    from tlone.common import parse_date
    assert parse_date("24052026") == date(2026, 5, 24) == parse_date("20260524")     # date de l'OV sans séparateurs
    assert transfers.devise_transfert("Transfert en USD") == "USD"
    assert transfers.devise_transfert("transfert en euros") == "EUR" and transfers.devise_transfert("EUR") == "EUR"
    assert transfers.devise_transfert("Autre") is None and transfers.devise_transfert(None) is None
    assert transfers.banque_similarite("ZIRAAT BANKASI", "T.C. ZIRAAT BANKASI A.S.") == 1.0
    assert transfers.banque_similarite("ZIRAAT BANKASI", "GARANTI BBVA") < 0.8
    s = load_settings()
    pl = pd.DataFrame([{"PERIODE_ID": "p1", "DOM_ID": "d1", "SIEGE_RACINE": "07310108259",
                        "DOM_COMPTE_LOCAL": "02700731010825900165", "MOIS_BASE": "2026-03", "PARTIE": "P1",
                        "PERIODE_TL": "2026-03P1", "NUMERO_DOMICILIATION": "07310 2026 000123 DZD",
                        "NOM_TRAVAILLEUR": "KAYA AHMET", "MONTANT_MAX_THEORIQUE": 51612.90,
                        "PART_TRANSFERABLE_CTS": 100000.0, "COEFFICIENT_PRORATA": 16 / 31,
                        "NB_JOURS_SEGMENT": 16, "NB_JOURS_MOIS": 31}])
    reg = pd.DataFrame([{"PERIODE_ID": "p1", "DOM_ID": "d1", "FICHIER": "deja.pdf", "_MONTANT": 10000.0,
                         "SIEGE_RACINE": "07310108259", "PERIODE_TL": "2026-03P1"}])
    atf = pd.DataFrame([{"NOM_PRENOM": "KAYA AHMET", "COMPTE": "02700731010825900165", "MOIS": "2026-03",
                         "MONTANT_TRANSFERABLE": 101000.0}])           # ≈ part mensuelle du contrat (+1 %)
    aut = pd.DataFrame([{"NOM_PRENOM": "KAYA AHMET", "COMPTE": "02700731010825900165", "MOIS": "2026-03",
                         "MONTANT_AUTORISE": 100000.0, "DATE_LIMITE": "10/03/2026"}])
    atf.to_excel(tmp_path / "atf.xlsx", index=False)
    aut.to_excel(tmp_path / "aut.xlsx", index=False)
    atf = transfers.load_referential(tmp_path / "atf.xlsx", transfers.ATF_COLUMNS)
    aut = transfers.load_referential(tmp_path / "aut.xlsx", transfers.AUT_COLUMNS)
    rhdf = pd.DataFrame([{"MATRICULE": "1", "NOM_PRENOM": "KAYA AHMET", "CLIENT": "07310108259", "MOIS": "2026-03",
                          "PART": "P1", "PERIODE": "2026-03P1", "PART_TRANSFERABLE": 41612.90,
                          "IBAN": "TR330006100519786457841326", "SWIFT": "TCZBTR2AXXX", "DEVISE": "EUR",
                          "BANQUE_ETR": "T.C. ZIRAAT BANKASI A.S.", "BENEFICIAIRE": "KAYA AYSE", "PERMIS_FIN": None,
                          "NB_JOURS": 16, "RIB": None}])
    d = {"fichier": "07310-108259-202603P1.pdf", "OV.montant_chiffres": 41612.90, "OV.periode": "MARS 2026 P1",
         "OV.compte_debiteur": "02700731010825900165", "OV.beneficiaire_nom": "KAYA AYSE",
         "OV.devise_transfert": "Transfert en USD", "OV.beneficiaire_compte": "TR33 0006 1005 1978 6457 8413 26",
         "OV.swift": "TCZBTR2A", "OV.banque_nom": "ZIRAAT BANKASI", "OV.date_demande": "03/03/2026",
         "ANNEXE_II.compte_devise": "TR330006100519786457841326", "ANNEXE_II.banque_etranger": "GARANTI BBVA",
         "ANNEXE_II.nom_prenom": "KAYA AHMET", "ANNEXE_II.part_transferable": 51612.90}
    sy, rows = transfers.control_dossier(d, pl, reg, atf, aut, s, today=date(2026, 3, 5), rh=rhdf,
                                         dom_dates={"073102026000123DZD": date(2026, 3, 12)})
    r = {x["CODE"]: x for x in rows}
    assert r["T12"]["RESULTAT"] == "OK" and "× 16/31" in r["T12"]["DETAIL"]      # 101 000 × 16/31 = 52 129,03
    assert sy["PLAFOND_ATF"] == round(101000 * 16 / 31, 2)
    assert sy["PLANNING_PART_PERIODE"] == 51612.90 and sy["PLANNING_PART_RESTANTE"] == 41612.90
    assert r["T13c"]["RESULTAT"] == "KO" and r["T13c"]["SEVERITE"] == "BLOQUANT"   # limite 10/03 < DOM 12/03
    assert r["R05"]["RESULTAT"] == "KO" and sy["DEVISE_TRANSFERT"] == "USD"        # USD sur l'OV / EUR en base RH
    assert r["R03"]["RESULTAT"] == "OK"                                           # espaces ignorés
    assert r["R09"]["RESULTAT"] == "KO" and "Annexe II" in r["R09"]["DETAIL"]
    assert "R08" not in r and "OV" not in r["T06"]["DETAIL"]
    from tlone import ov_matrice as ov
    m = ov.build_ov(d["fichier"], d, sy, pd.DataFrame(rows), pl, {}, rh=rhdf.iloc[0].to_dict())
    ligne = {x["ligne"].key: x for x in m["rows"]}
    assert "plafond" not in ligne and "OV" not in ligne["nom"]["ligne"].sources
    assert ligne["montant"]["cells"]["PLANNING"]["value"] == 51612.90               # part du contrat, inchangée
    assert ligne["restant"]["cells"]["PLANNING"]["value"] == 41612.90
    assert ligne["validite"]["cells"]["AUTORISATION"]["bloq"]
    assert ligne["devise"]["cells"]["RH"]["open"] and ligne["devise"]["retenu_txt"] == "USD"
    assert not ligne["benef_compte"]["cells"]["ANNEXE_II"]["open"] and ligne["banque"]["cells"]["ANNEXE_II"]["open"]


def test_vigilance_fiche_retournee_hors_modele_lue_sans_bloquer(tmp_path):
    """Fiche retournée sans les libellés d'en-tête (Client:, Réf Dossier:…) ou avec une ponctuation différente :
    lue quand même (rubriques), libellés absents signalés ; le modèle d'envoi reste contrôlé strictement."""
    import io
    from openpyxl import load_workbook
    from tlone import vigilance as v
    tpl = tmp_path / "modele.xlsx"
    _modele_vigilance(tpl)
    wb = load_workbook(tpl)
    ws = wb.worksheets[0]
    for c in ("A5", "A6", "A7", "A8"):
        ws[c] = None
    ws["A13"] = "Pays de la banque du bénéficiaire :"        # ponctuation différente : reconnu
    ws["A12"] = "Banque  du  BENEFICIAIRE"
    ws["D16"] = "OK ==> HIT"
    b = io.BytesIO()
    wb.save(b)
    ret = v.read_return(b.getvalue())
    assert ret["client"] == "" and ret["ref"] == ""
    assert set(ret["libelles_absents"]) == {"Client:", "Réf Dossier:", "Circuit:", "Type de dossier:"}
    assert any(rb["resultat"] == "OK ==> HIT" for rb in ret["rubriques"])
    with pytest.raises(ValueError):
        v.locate(load_workbook(io.BytesIO(b.getvalue())).worksheets[0])          # strict : modèle d'envoi
    with pytest.raises(ValueError):
        v.read_return(_xlsx_vide())


def _xlsx_vide() -> bytes:
    import io
    from openpyxl import Workbook
    b = io.BytesIO()
    Workbook().save(b)
    return b.getvalue()


def test_changement_d_employeur_depuis_une_page(tmp_path, monkeypatch):
    """« Ouvrir cet employeur » / démo : le sélecteur de la barre latérale n'est pas modifié après son affichage
    (StreamlitWidgetAlreadyInstantiatedError) ; il reprend le choix au prochain affichage."""
    from streamlit.testing.v1 import AppTest
    monkeypatch.setenv("TLONE_DATA_ROOT", str(tmp_path / "home" / "employeurs" / "TOSYALI"))
    monkeypatch.setenv("TLONE_RUN_MODE", "local")
    from tlone import employeurs
    employeurs.ensure_registered(tmp_path / "home", "TOSYALI", "TOSYALI IRON STEEL INDUSTRY ALGERIE")
    employeurs.create(tmp_path / "home", "AUTRE SPA", "000016012345678", "u")
    root = Path(__file__).resolve().parent.parent
    at = AppTest.from_file(str(root / "Accueil.py"), default_timeout=240)
    at.run()
    at.switch_page("pages/0_Employeurs.py")
    at.run()
    code = employeurs.load(tmp_path / "home").set_index("RAISON_SOCIALE").loc["AUTRE SPA", "CODE"]
    at.selectbox(key="emp_open").set_value(code)
    [x for x in at.button if x.label == "Ouvrir cet employeur"][0].click()
    at.run()
    assert not at.exception
    choisi = at.session_state["employeur"]
    assert choisi != "TOSYALI"
    import streamlit as _stl
    if tuple(int(x) for x in _stl.__version__.split(".")[:2]) < (1, 50):
        return   # AppTest < 1.50 rejoue les widgets de la page précédente : parcours vérifié en navigateur réel
    for pg in ("pages/2_Planning_TL.py", "pages/3_Transferts_OV.py", "Accueil.py"):   # reste choisi d'une page à l'autre
        at.switch_page(pg)
        at.run()
        assert not at.exception and at.session_state["employeur"] == choisi
        assert at.sidebar.selectbox[0].value == choisi
    at.sidebar.selectbox[0].set_value("TOSYALI")                                   # choix dans la barre latérale
    at.run()
    assert at.session_state["employeur"] == "TOSYALI"
    at.switch_page("pages/0_Employeurs.py")
    at.run()
    assert at.session_state["employeur"] == "TOSYALI" and at.sidebar.selectbox[0].value == "TOSYALI"


def test_worker_reel_vu_par_l_app_et_battement_en_tache_de_fond(tmp_path, monkeypatch):
    """run_worker.py lancé comme sur Domino (--data-root) : l'App le voit ; le battement continue pendant un long
    traitement ; un worker d'une version précédente est diagnostiqué."""
    import json as _json
    import subprocess
    import sys as _sys
    import time as _time
    from tlone import worker
    root = tmp_path / "home" / "employeurs" / "TOSYALI"
    root.mkdir(parents=True)
    monkeypatch.setenv("TLONE_DATA_ROOT", str(root))
    s = load_settings()
    assert worker.worker_state(s) is None and "non démarré" in worker.worker_absent_text(s)
    assert "--data-root " + str(root) in worker.commande_worker(s)
    app = Path(__file__).resolve().parent.parent
    env = {k: v for k, v in __import__("os").environ.items() if k != "TLONE_DATA_ROOT"}
    proc = subprocess.Popen([_sys.executable, str(app / "run_worker.py"), "--steps", "ov_extract",
                             "--data-root", str(root)], env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                            text=True)
    try:
        for _ in range(60):
            if worker.worker_state(s):
                break
            _time.sleep(0.25)
        w = worker.worker_state(s)
        assert w and w["state"] == "ATTENTE" and w["data_root"] == str(root) and w["app_dir"] == str(app)
    finally:
        proc.terminate()
        out = proc.communicate(timeout=20)[0]
    assert "signal de vie : " + str(worker.heartbeat_path(s)) in out
    # battement en tâche de fond pendant un traitement long (main thread occupé)
    monkeypatch.setattr(worker, "BEAT_EVERY_S", 0.05)
    wk = worker.Worker(s, ["ov_extract"])
    wk.start_beating()
    wk.state = "TRAITEMENT"
    t0 = _json.loads(worker.heartbeat_path(s).read_text())["epoch"]
    _time.sleep(0.3)
    d = _json.loads(worker.heartbeat_path(s).read_text())
    wk._stop.set()
    assert d["epoch"] > t0 and d["state"] == "TRAITEMENT"
    # ancien worker (v3) : signal dans runs/ de l'employeur, pas dans _commun/worker
    worker.heartbeat_path(s).unlink()
    s.paths.runs_dir.mkdir(parents=True, exist_ok=True)
    worker.legacy_heartbeat_path(s).write_text(_json.dumps({"epoch": _time.time(), "host": "h", "app_dir": "/mnt/code/tlone_app_v3"}))
    assert "version précédente" in worker.worker_absent_text(s) and "tlone_app_v3" in worker.worker_absent_text(s)


def test_planning_dates_de_reference_et_duree(env):
    """Décision du 23/10/2026 : le Planning va de la date de début à la date de fin de RÉFÉRENCE (matrice DOM) ;
    fin absente -> début + durée − 1 jour ; contrôle nombre de mois planifiés = durée de référence."""
    from tlone import ui
    p, _ = env
    dos = domvalid.load_sheet(p.validation_xlsx, "DOSSIERS_A_VALIDER")
    man = domvalid.load_sheet(p.validation_xlsx, "A_VALIDER_MANUELLEMENT")
    pa = pd.read_excel(p.param_aug_xlsx, dtype=object)
    base = planning.build_planning(dos, man, pa)
    kaya0 = base["Planning_TL"][base["Planning_TL"]["NOM_TRAVAILLEUR"] == "KAYA AHMET"]
    assert kaya0.iloc[0]["DATE_DEBUT_REFERENCE"] == "2026-01-03" and kaya0.iloc[0]["BASE_DATES"] == "REFERENCE"
    k = dos["FICHIER"] == "KAYA_AHMET_CNT.pdf"
    dos.loc[k, "CTS_DUREE_MOIS"] = "24"
    r = planning.build_planning(dos, man, pa)
    cd = r["CONTROLE_DUREE"].set_index("NOM_TRAVAILLEUR")
    assert cd.loc["KAYA AHMET", "DUREE_REFERENCE_MOIS"] == 24 and cd.loc["KAYA AHMET", "MOIS_PLANIFIES"] == 24
    assert cd.loc["KAYA AHMET", "CONTROLE_DUREE"] == "OK"
    # date de fin de référence absente : calculée avec la durée (12 mois)
    dos.loc[k, "CTS_DATE_FIN_VALIDITE_PERMIS"] = None
    dos.loc[k, "CTS_DUREE_MOIS"] = "12"
    aug = dos["FICHIER"] == "KAYA_AHMET_AUG_26.pdf"
    dos.loc[aug, "CTS_DATE_FIN_VALIDITE_PERMIS"] = "02/01/2027"
    r = planning.build_planning(dos, man, pa)
    kaya = r["Planning_TL"][r["Planning_TL"]["NOM_TRAVAILLEUR"] == "KAYA AHMET"]
    v0 = kaya[kaya["TYPE_VERSION"] == "NOUVEAU_CONTRAT"].iloc[0]
    assert v0["DATE_FIN_REFERENCE"] == "2027-01-02" and "DUREE" in v0["BASE_DATES"]
    assert r["CONTROLE_DUREE"].set_index("NOM_TRAVAILLEUR").loc["KAYA AHMET", "CONTROLE_DUREE"] == "OK"
    # dossier non validé dans la matrice : absent du Planning
    r = planning.build_planning(dos, man, pa, fichiers_valides={"KAYA_AHMET_CNT.pdf", "KAYA_AHMET_AUG_26.pdf"})
    assert set(r["Planning_TL"]["NOM_TRAVAILLEUR"]) == {"KAYA AHMET"}
    assert ui.HORIZONS == ("Toute la durée", "Période actuelle", "Intervalle")


def test_id_tiers_periode_et_bases_du_mois_entier(tmp_path):
    """Décision du 23/10/2026 : ID tiers « 07410 109109 » + période « 2026-05 » dans toutes les bases ; bases passées
    du mois entier (sans P1 / P2) ; ATF et autorisations du mois entier ; toutes agences (07310, 07410…)."""
    from tlone import rh, transfers
    from tlone.common import id_tiers, siege_racine
    assert id_tiers("07410 109109") == "07410109109" and siege_racine("02700741010910900155") == "07410109109"
    assert siege_racine("07410-109109-202605P2.pdf") == "07410109109"
    # base RH ancienne : ID tiers + période, sans RIB ni PART (mois entier)
    _rh_xlsx(tmp_path / "rh.xlsx", [
        {"ID TIERS": "07410 109109", "PERIODE": "2026-05", "NOM ET PRENOM": "ARSLAN EMRE", "Part transférable": 300000,
         "Salaire Net": 315789.47},
        {"ID TIERS": "07410 109110", "PERIODE": "2026-05P1", "NOM ET PRENOM": "BOZ ALI", "Part transférable": 100000,
         "Compte bancaire (Algérie)": "02700731010825900165"}])
    df = rh.read(tmp_path / "rh.xlsx")
    a = df.set_index("NOM_PRENOM")
    assert a.loc["ARSLAN EMRE", "CLIENT"] == "07410109109" and a.loc["ARSLAN EMRE", "MOIS"] == "2026-05"
    assert bool(a.loc["ARSLAN EMRE", "RH_MENSUELLE"]) and a.loc["BOZ ALI", "PART"] == "P1"
    q = rh.dq(df)
    assert "RH11" in set(q[q["NOM_PRENOM"] == "ARSLAN EMRE"]["CODE"])                 # RIB absent : à vérifier
    assert "RH10" in set(q[q["NOM_PRENOM"] == "BOZ ALI"]["CODE"])                     # ID tiers ≠ client du RIB
    r, how = rh.match(df, "07410109109", "2026-05", "P1")
    assert r is not None and "mois entier" in how
    # contrôle OV P1 face à une base RH du mois entier : part attendue = mois × quote-part du Planning
    s = load_settings()
    pl = pd.DataFrame([{"PERIODE_ID": p, "DOM_ID": "d", "SIEGE_RACINE": "07410109109", "MOIS_BASE": "2026-05",
                        "PARTIE": p[-2:], "PERIODE_TL": f"2026-05{p[-2:]}", "MONTANT_MAX_THEORIQUE": m,
                        "NUMERO_DOMICILIATION": "1", "DOM_COMPTE_LOCAL": "02700741010910900155",
                        "DATE_DEBUT_SEGMENT": "2026-05-01", "DATE_FIN_SEGMENT": "2026-05-31", "NOM_TRAVAILLEUR": "ARSLAN EMRE"}
                       for p, m in (("xP1", 170000.0), ("xP2", 130000.0))])
    reg = pd.DataFrame(columns=transfers.REGISTRE_COLUMNS)
    atf = pd.DataFrame([{"ID_TIERS": "07410 109109", "NOM_PRENOM": "ARSLAN EMRE", "MOIS": "2026-05",
                         "MONTANT_TRANSFERABLE": 300000}])
    atf.to_excel(tmp_path / "atf.xlsx", index=False)
    atf = transfers.load_referential(tmp_path / "atf.xlsx", transfers.ATF_COLUMNS, df)
    assert atf["_CLIENT"].iloc[0] == "07410109109"
    d = {"fichier": "07410-109109-202605P1.pdf", "OV.montant_chiffres": 170000.0, "OV.periode": "MAI 2026 P1",
         "OV.compte_debiteur": "02700741010910900155"}
    _, rows = transfers.control_dossier(d, pl, reg, atf, pd.DataFrame(), s, rh=df)
    c = {x["CODE"]: x for x in rows}
    assert c["R02"]["RESULTAT"] == "OK" and "quote-part P1" in c["R02"]["DETAIL"]
    assert c["T12"]["RESULTAT"] == "OK" and "ID TIERS" in c["T12"]["DETAIL"]
    # P2 alors que la P1 (170 000) est déjà validée : plafond ATF du mois = 300 000 − 170 000
    reg = pd.DataFrame([{"TL_ID": "t", "FICHIER": "07410-109109-202605P1.pdf", "PERIODE_ID": "xP1", "PERIODE_TL": "2026-05P1",
                         "SIEGE_RACINE": "07410109109", "MONTANT": "170000", "DECISION": "VALIDE"}]).reindex(columns=transfers.REGISTRE_COLUMNS)
    reg = transfers._registre_valide(reg)
    d2 = dict(d, fichier="07410-109109-202605P2.pdf", **{"OV.montant_chiffres": 140000.0, "OV.periode": "MAI 2026 P2"})
    _, rows = transfers.control_dossier(d2, pl, reg, atf, pd.DataFrame(), s, rh=df)
    t12 = {x["CODE"]: x for x in rows}["T12"]
    assert t12["RESULTAT"] == "KO" and "déjà sur l'autre partie" in t12["DETAIL"]
    pv = rh.planning_vs_rh(pl, df[df["NOM_PRENOM"] == "ARSLAN EMRE"])
    assert pv["STATUT"].iloc[0] == "OK" and pv["PART_PLANNING"].iloc[0] == 300000.0      # mois entier = P1 + P2


def test_matrice_dom_colonne_base_rh_arbitre():
    """La base RH arbitre une donnée d'identité en anomalie : pièce de référence vide -> valeur RH retenue."""
    from tlone import matrice as mx
    lg = mx.LIGNE_BY_KEY["naissance"]
    assert "RH" in lg.sources and "RH" in mx.COLS
    cells = {c: {"present": c in ("CTR", "RH"), "value": {"CTR": "02/01/1980", "RH": "02/01/1980"}.get(c)}
             for c in mx.COLS}
    assert mx._arbitre_rh(lg, cells, "CTS", None) == ("RH", "02/01/1980")
    cells["CTS"] = {"present": True, "value": "02/01/1981"}
    assert mx._arbitre_rh(lg, cells, "CTS", "02/01/1981") == ("RH", "02/01/1980")   # CTR confirme la base RH
    lgp = mx.LIGNE_BY_KEY["part"]
    assert mx._arbitre_rh(lgp, {"RH": {"present": True, "value": "1"}}, "CTS", None) == ("CTS", None)  # montant : au valideur


def _demo_P(tmp_path, monkeypatch):
    from datetime import date as _d
    from tlone import demo
    home = tmp_path / "home"
    monkeypatch.setenv("TLONE_DATA_ROOT", str(home / "employeurs" / "TOSYALI"))
    demo.build(home, today=_d(2026, 10, 21))
    s = load_settings("DEMO")
    return s, s.paths


def test_dom_decisions_en_lot(tmp_path, monkeypatch):
    """Sélection multiple DOM : seuls les dossiers prêts sont validés, après contrôle de l'échantillon ; rejet en lot."""
    from tlone import matrice
    _s, p = _demo_P(tmp_path, monkeypatch)

    def charger():
        lot = matrice.load_lot(p)
        return matrice.build_all(lot, matrice.read_journal(p.dom_journal), matrice.default_sources(p.sources_reference_xlsx))
    ms = charger()
    tous = [m["fichier"] for m in ms]
    rep = matrice.repartition(ms, tous)
    assert len(rep["prets"]) == 11 and len(rep["anomalies"]) == 2 and not rep["decides"]
    lot = matrice.etat_lot(ms)
    assert lot["echantillon"] == lot["reste"] and len(lot["reste"]) == 3 and set(lot["reste"]) <= set(rep["prets"])
    # tirage par empreinte : indépendant de l'ordre, quasi stable quand le lot grandit
    assert matrice.echantillon(["a", "b", "c", "d"]) == matrice.echantillon(["d", "c", "b", "a"])
    big = [f"F{i}.pdf" for i in range(200)]
    e1 = set(matrice.echantillon(big))
    assert len(e1) == 10 and len(e1 & set(matrice.echantillon(big + ["X.pdf"]))) >= 9
    # échantillon contrôlé un par un, puis validation en lot de toute la sélection
    for f in lot["reste"]:
        matrice.append_event(p.dom_journal, "v1", f, "VALIDE", n_corrections=0)
    ms = charger()
    assert matrice.etat_lot(ms)["reste"] == []
    n = matrice.valider_lot(p.dom_journal, "v1", ms, tous, matrice.etat_lot(ms)["echantillon"])
    assert n == 8
    by = {m["fichier"]: m for m in charger()}
    assert all(by[f]["decision"]["action"] == "VALIDE" for f in rep["prets"])
    assert all(by[f]["decision"] is None for f in rep["anomalies"])          # jamais validé en lot avec des cases rouges
    assert {by[f]["decision"].get("mode") for f in rep["prets"] if f not in lot["reste"]} == {"LOT"}
    assert matrice.etat_lot(list(by.values()))["reste"] == []               # l'échantillon reste celui contrôlé
    # rejet en lot : seulement les non décidés
    assert matrice.rejeter_lot(p.dom_journal, "v1", list(by.values()), tous, "Pièce manquante", None) == 2
    by = {m["fichier"]: m for m in charger()}
    assert all(by[f]["decision"]["action"] == "REJETE" for f in rep["anomalies"])


def test_ov_decisions_en_lot(tmp_path, monkeypatch):
    """Sélection multiple OV : seuls les conformes rattachés sont validés ; le rejet exige un motif."""
    from tlone import ov_matrice as ov, referentiels
    s, P = _demo_P(tmp_path, monkeypatch)
    A = transfers.read_audit_xlsx(transfers.latest_audit(P.tr_out)[0])
    R = referentiels.load_all(P)
    sy, ch = transfers.control_each(A, planning.load_planning(P.planning_csv), transfers.load_registre(P.registre_csv),
                                    R["atf"], R["aut"], s, None, R["rh"])
    df = ov.list_frame(sy, transfers.load_decisions(P.decisions_csv))
    tous = list(df["FICHIER"])
    rep = ov.repartition(df, tous)
    assert rep["conformes"] and rep["a_ouvrir"]
    assert "07310-990104-202609.pdf" in rep["a_ouvrir"] + rep["decides"]          # non conforme : jamais en lot
    assert set(rep["conformes"]) <= set(sy[sy["STATUT_CONTROLE"] == "CONFORME"]["FICHIER"])
    ed = ov.lignes_decision(sy, rep["conformes"], "VALIDE", "lot")
    assert transfers.check_decisions(ed) == [] and len(ed) == len(rep["conformes"])
    n = transfers.save_decisions(P.decisions_csv, ed, "v1")
    assert n == len(rep["conformes"])
    df2 = ov.list_frame(sy, transfers.load_decisions(P.decisions_csv))
    assert set(ov.repartition(df2, tous)["decides"]) >= set(rep["conformes"])
    rej = ov.lignes_decision(sy, rep["a_ouvrir"][:1], "REJETE", "")
    assert transfers.check_decisions(rej)                                          # rejet sans motif refusé


def test_tolerances_matrice_dom(env, tmp_path):
    """Tolérances du 24/10/2026 : « Agence » ignoré, employeur reconnu par mot-clé, nom ≥ 90 % avec identifiant
    client identique, n° de permis TTR facultatif (les écarts restent visibles en ⓘ, sans bloquer)."""
    p, _ = env
    from tlone import employeurs, matrice
    j = p.dom_journal
    f = "KAYA_AHMET_CNT.pdf"

    def cell(m, k, c):
        return next(x for x in m["rows"] if x["ligne"].key == k)["cells"][c]

    def build(mots=None):
        lot = matrice.load_lot(p)
        if mots is not None:
            lot.mots_cles = mots
        return {m["fichier"]: m for m in matrice.build_all(lot, matrice.read_journal(j), matrice.default_sources())}[f]
    # agence : « MOSTAGANEM » = « Agence Mostaganem » ; une autre agence reste une anomalie
    assert matrice.agence_norm("BNP Paribas El Djazaïr - Agence Mostaganem") == "MOSTAGANEM"
    assert matrice.same("agence", "Mostaganem", "AGENCE MOSTAGANEM")[0] is True
    assert matrice.same("agence", "Mostaganem", "Agence Arzew")[0] is False
    car_ag = cell(build(), "agence", "CARTHAGO")["value"]
    assert car_ag and "ORAN" in car_ag.upper()
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="agence", col="DOM", valeur="ORAN")
    c = cell(build(), "agence", "CARTHAGO")
    assert not c["open"] and c["info"]
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="agence", col="DOM", valeur="ARZEW")
    assert cell(build(), "agence", "CARTHAGO")["open"]
    # nom : 90 % de similarité + identifiant client identique DOM / Carthago -> tolérance ; en dessous -> anomalie
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="nom", col="CTR", valeur="KAYA AHMED")
    c = cell(build(), "nom", "CTR")
    assert not c["open"] and c["info"] and "identique" in c["issues"][-1][1]
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="nom", col="CTR", valeur="KAYA MEHMET")
    assert cell(build(), "nom", "CTR")["open"]
    # n° de permis du TTR : facultatif, un écart n'est pas une anomalie
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="permis", col="TTR", valeur="99-99999999 / 31-26-999999")
    c = cell(build(), "permis", "TTR")
    assert not c["open"] and c["info"] and "facultatif" in c["issues"][-1][1]
    # employeur : conforme dès que la pièce contient le mot-clé de l'employeur TL-One
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="employeur", col="CTS", valeur="SPA TOSYALI IRON STEEL INDUSTRY")
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="employeur", col="CTR", valeur="TOSYALI IRON STEEL INDUST")
    matrice.append_event(j, "v1", f, "CORRIGE", ligne="employeur", col="DOM", valeur="ACME SARL")
    m = build(["TOSYALI"])
    assert not cell(m, "employeur", "CTR")["open"] and not cell(m, "employeur", "CTS")["open"]
    assert cell(m, "employeur", "DOM")["open"]
    assert cell(build([]), "employeur", "CTR")["open"]        # sans mot-clé : comparaison entre pièces (inchangée)
    # mots-clés lus dans le registre des employeurs
    home = tmp_path / "reg"
    employeurs.ensure_registered(home, "TOSYALI", "TOSYALI IRON STEEL INDUSTRY ALGERIE")
    (home / "employeurs" / "TOSYALI").mkdir(parents=True, exist_ok=True)
    from tlone.config import Paths
    assert matrice.mots_cles_employeur(Paths(home / "employeurs" / "TOSYALI")) == ["TOSYALI"]
    assert matrice.mots_cles_employeur(p) == []


def test_worker_arret_commande_et_sigterm(tmp_path, monkeypatch):
    """Arrêt du worker sans Ctrl+C : « run_worker.py --arreter » depuis un autre terminal, ou kill / pkill (SIGTERM) :
    arrêt propre, signal de vie effacé (l'App le voit arrêté aussitôt)."""
    import signal as _sig
    import subprocess
    import sys as _sys
    import time as _time
    from tlone import worker
    root = tmp_path / "home" / "employeurs" / "TOSYALI"
    root.mkdir(parents=True)
    monkeypatch.setenv("TLONE_DATA_ROOT", str(root))
    s = load_settings()
    app = Path(__file__).resolve().parent.parent
    env = {k: v for k, v in __import__("os").environ.items() if k != "TLONE_DATA_ROOT"}

    def lancer():
        pr = subprocess.Popen([_sys.executable, str(app / "run_worker.py"), "--steps", "ov_extract",
                               "--data-root", str(root)], env=env, stdout=subprocess.PIPE,
                              stderr=subprocess.STDOUT, text=True)
        for _ in range(80):
            if worker.worker_state(s):
                return pr
            _time.sleep(0.25)
        pr.kill()
        raise AssertionError(pr.communicate()[0])
    assert worker.arreter(s, attente=1) == 0                 # rien à arrêter
    pr = lancer()
    r = subprocess.run([_sys.executable, str(app / "run_worker.py"), "--arreter", "--attente", "30",
                        "--data-root", str(root)], env=env, capture_output=True, text=True, timeout=60)
    out = pr.communicate(timeout=30)[0]
    assert r.returncode == 0 and "worker arrêté" in r.stdout, r.stdout
    assert "arrêt demandé (commande --arreter)" in out and "GPU libéré" in out and "--arreter" in out
    assert pr.returncode == 0 and not worker.heartbeat_path(s).exists() and not worker.stop_path(s).exists()
    pr = lancer()                                            # kill / pkill : même arrêt propre
    pr.send_signal(_sig.SIGTERM)
    out = pr.communicate(timeout=30)[0]
    assert "arrêt demandé" in out and "GPU libéré" in out and not worker.heartbeat_path(s).exists()


def test_carthago_type_inconnu_se_confirme(env):
    """Carthago « trouvée, type DOM/PREDOM inconnu » : à vérifier, se confirme (tracé) ; le dossier devient validable.
    En-tête : « Trouvé, type à vérifier » (et non « Non trouvé »)."""
    p, _ = env
    from tlone import matrice
    f = "KAYA_AHMET_CNT.pdf"
    st_ = "EXACT_DOMICILIATION_REFERENCE_MATCH_TYPE_UNKNOWN"

    def build():
        lot = matrice.load_lot(p)
        for df in (lot.dossiers, lot.carthago):
            if "CARTHAGO_MATCH_STATUS" in df.columns:
                df.loc[df["FICHIER"].astype(str) == f, "CARTHAGO_MATCH_STATUS"] = st_
        m = {x["fichier"]: x for x in matrice.build_all(lot, matrice.read_journal(p.dom_journal),
                                                         matrice.default_sources())}[f]
        return m, next(x for x in m["rows"] if x["ligne"].key == "domref")["cells"]["CARTHAGO"]
    m, car = build()
    assert car["open"] and not car["bloq"] and matrice.summary(m)["CARTHAGO"] == "Trouvé, type à vérifier"
    matrice.append_event(p.dom_journal, "v1", f, "CONFIRME", ligne="domref", col="CARTHAGO", valeur_vue=str(car["value"]))
    m, car = build()
    assert not car["open"] and car["edit"]["action"] == "CONFIRME"


def test_carthago_secours_compte_mal_lu(env):
    """Compte mal lu (0 lu 6) : la Partie 2 ne rapproche rien ; TL-One retrouve le client dans l'export CarthagoDom
    (nom du fichier / nom du salarié), signale le compte et propose le compte à clé juste ; après correction du
    compte, le rapprochement se fait par le compte."""
    p, _ = env
    from tlone import matrice
    f = "07310-108448-I.pdf"
    ok_compte = "02700731010844800130"
    assert matrice.proposition_compte("02760731016844800130", "07310108448") == ok_compte

    def build(journal):
        lot = matrice.load_lot(p)
        base = lot.dossiers.iloc[0].copy()
        base["FICHIER"] = f
        for k in ("DOM_COMPTE_LOCAL", "DOM_COMPTE_LOCAL_NORMALIZED"):
            base[k] = "02760731016844800130"
        base["CARTHAGO_MATCH_STATUS"] = "NO_EXACT_DOMICILIATION_REFERENCE_MATCH"
        base["CTS_NOM_PRENOM_TRAVAILLEUR"] = "ONAY CAN"
        import pandas as _pd
        lot.dossiers = _pd.concat([lot.dossiers, base.to_frame().T], ignore_index=True)
        lot.carthago_export = [{"client": "07310108448", "nom": "ONAY CAN", "dom": "27190120261400069DZD",
                                "agence": "Agence Mostaghanem"}]
        st = matrice.fold(journal).get(f)
        m = matrice.build(lot, f, st, matrice.default_sources())
        return m, {x["ligne"].key: x for x in m["rows"]}
    m, L = build([])
    assert matrice.summary(m)["CARTHAGO"] == "Trouvé, type à vérifier"
    assert L["domref"]["cells"]["CARTHAGO"]["value"] == "27190120261400069DZD"
    assert any(ok_compte in t for _k, t in L["compte"]["cells"]["DOM"]["issues"])
    j = [{"fichier": f, "action": "CORRIGE", "ligne": "compte", "col": "DOM", "valeur": ok_compte, "user": "v",
          "ts": "2026-10-24T23:00:00"}]
    m, L = build(j)
    assert not any("mal lu" in t for _k, t in L["compte"]["cells"]["DOM"]["issues"])
    assert L["domref"]["cells"]["CARTHAGO"]["value"] == "27190120261400069DZD"


def test_avancement_lecture(env):
    """Barre de progression : PDF déposés sans JSON = à lire ; JSON écrits depuis le lancement = lus."""
    import json as _json
    import time as _time
    from datetime import datetime as _dt
    from tlone import avancement as av
    _p0, _s0 = env
    s = load_settings(data_root=_p0.data_root / "vide")
    p = s.paths
    p.dom_in.mkdir(parents=True, exist_ok=True)
    p.raw_json_dir.mkdir(parents=True, exist_ok=True)
    for n in ("A_CNT", "B_CNT", "C_CNT", "D_CNT"):
        (p.dom_in / f"{n}.pdf").write_bytes(b"%PDF-1.4")
    (p.raw_json_dir / "A_CNT.json").write_text("{}")                    # déjà lu avant le lancement
    _time.sleep(1.1)
    run = {"step": "dom_extract", "launched_at": _dt.now().isoformat(timespec="seconds")}
    _time.sleep(1.1)
    assert [x.stem for x in av.a_lire(s, "dom_extract")] == ["B_CNT", "C_CNT", "D_CNT"]
    e = av.etat(s, run)
    assert (e["lus"], e["total"], e["en_cours"]) == (0, 3, "B_CNT")
    (p.raw_json_dir / "dossier_B_CNT_v15.json").write_text(_json.dumps({"source_file": "B_CNT.pdf"}))
    e = av.etat(s, run)
    assert (e["lus"], e["total"], e["en_cours"]) == (1, 3, "C_CNT") and e["eta_s"] is not None
    e = av.etat(s, dict(run, max_items=2))                                  # « Limiter à N »
    assert e["total"] == 2
    assert av.duree(30) == "moins d'une minute" and av.duree(300) == "environ 5 min"


def test_worker_enchaine_lecture_puis_controles(env, tmp_path, monkeypatch):
    """Vue essentielle : un seul bouton ; la lecture DOM réussie dépose d'elle-même la demande des contrôles."""
    import dataclasses
    import json as _json
    import sys as _sys
    import nbformat
    from tlone import jobs, preflight, worker
    p, s = env
    nb = nbformat.v4.new_notebook()
    nb.cells = [nbformat.v4.new_code_cell("from pathlib import Path\nINPUT_DIR = Path('/mnt/data/domiciliations_in')\n"
                                          "pdfs = sorted(INPUT_DIR.glob('*.pdf'))"),
                nbformat.v4.new_code_cell("# from_pretrained simulé\nX = 1"),
                nbformat.v4.new_code_cell("if not pdfs:\n    print('rien')")]
    fake = tmp_path / "fake.ipynb"
    nbformat.write(nb, str(fake))
    s = dataclasses.replace(s, notebooks={**s.notebooks, "dom_extract": fake}, run_mode="worker")
    monkeypatch.setattr(preflight, "check_step", lambda s_, step: [])
    w = worker.Worker(s, ["dom_extract"], pythons={"dom_extract": _sys.executable})
    r1 = jobs.launch(s, "dom_extract", None, "testeur", ensuite="dom_post")
    w.loop(once=True)
    q = [_json.loads(x.read_text()) for x in worker.queue_dir(s).glob("*.json")]
    assert [x["step"] for x in q] == ["dom_post"] and q[0]["user"] == "testeur"
    assert {r["run_id"]: r for r in jobs.list_runs(s)}[r1["run_id"]]["status"] == "SUCCEEDED"
    w.stop()


def test_worker_libre_apres_controles(env, monkeypatch):
    """Après une Partie 2 (exécutée à froid), le worker redevient « en attente » (voyant sans « lecture en cours »)."""
    import dataclasses
    from tlone import worker
    p, s = env
    s = dataclasses.replace(s, run_mode="worker")
    w = worker.Worker(s, ["dom_extract"])
    monkeypatch.setattr(worker.subprocess, "run", lambda *a, **k: None)
    w.handle({"run_id": "dom_post_x", "step": "dom_post"})
    st_ = worker.worker_state(s)
    assert w.current is None and st_["state"] == "ATTENTE" and st_["current"] is None
