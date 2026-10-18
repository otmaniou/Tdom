"""Module 1 — validation humaine du classeur Partie 2 (validation_domiciliations_V15_A.xlsx).

L'app remplace la saisie Excel de A_VALIDER_MANUELLEMENT, mais écrit dans LE MÊME
classeur et les MÊMES colonnes (VALEUR_VALIDEE / DECISION_MANUELLE / COMMENTAIRE) :
le notebook Planning_TL_V15A reste exécutable tel quel sur ce fichier.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd

from .common import backup, is_blank

MANUAL_SHEET = "A_VALIDER_MANUELLEMENT"
DOSSIERS_SHEET = "DOSSIERS_A_VALIDER"
EDITABLE = ["VALEUR_VALIDEE", "DECISION_MANUELLE", "COMMENTAIRE"]
DECISIONS = ["", "VALIDEE", "CORRIGEE", "REJETEE"]

CONTEXT_SHEETS = ["DQ_DASHBOARD", "CONTROLES_COHERENCE", "CARTHAGO_MATCH", "MOTIFS_CLIENT",
                  "ANOMALIES", "DOCUMENTS", "ERREURS", "PARAMETRES"]


def sheet_names(xlsx: Path) -> list[str]:
    return pd.ExcelFile(xlsx).sheet_names


def load_sheet(xlsx: Path, sheet: str) -> pd.DataFrame:
    try:
        return pd.read_excel(xlsx, sheet_name=sheet, dtype=object)
    except ValueError:
        return pd.DataFrame()


def load_manual(xlsx: Path) -> pd.DataFrame:
    df = load_sheet(xlsx, MANUAL_SHEET)
    for c in EDITABLE:
        if c not in df.columns:
            df[c] = None
    df = df.reset_index(drop=True)
    df.insert(0, "_ROW", range(2, len(df) + 2))  # n° de ligne Excel, clé d'écriture
    for c in EDITABLE:
        df[c] = df[c].map(lambda v: "" if is_blank(v) else str(v))
    return df


def check_decisions(edited: pd.DataFrame) -> list[str]:
    """Contrôles bloquants avant sauvegarde (mêmes exigences que Planning 2B)."""
    errs = []
    for _, r in edited.iterrows():
        d = (r.get("DECISION_MANUELLE") or "").strip().upper()
        if d and d not in DECISIONS:
            errs.append(f"Ligne {r['_ROW']} : décision inconnue « {d} »")
        if d == "CORRIGEE" and is_blank(r.get("VALEUR_VALIDEE")):
            errs.append(f"Ligne {r['_ROW']} ({r.get('CHAMP')}) : CORRIGEE sans VALEUR_VALIDEE")
        if d == "REJETEE" and is_blank(r.get("COMMENTAIRE")):
            errs.append(f"Ligne {r['_ROW']} ({r.get('CHAMP')}) : REJETEE sans motif en COMMENTAIRE")
    return errs


def save_decisions(xlsx: Path, edited: pd.DataFrame, expected_mtime: float | None = None) -> dict:
    """Écrit uniquement les 3 colonnes éditables des lignes modifiées, en place.

    `expected_mtime` protège contre l'écrasement si le fichier a été régénéré
    (nouveau run Partie 2) ou modifié par un autre valideur entre-temps.
    """
    from openpyxl import load_workbook

    if expected_mtime is not None and abs(xlsx.stat().st_mtime - expected_mtime) > 1e-6:
        raise RuntimeError("Le classeur a changé depuis son chargement (nouveau run Partie 2 "
                           "ou autre valideur). Rechargez la page avant d'enregistrer.")
    bkp = backup(xlsx)
    wb = load_workbook(xlsx)
    ws = wb[MANUAL_SHEET]
    headers = {c.value: c.column for c in ws[1]}
    for col in EDITABLE:
        if col not in headers:
            headers[col] = ws.max_column + 1
            ws.cell(row=1, column=headers[col], value=col)
    changes = []
    for _, r in edited.iterrows():
        row = int(r["_ROW"])
        # Vérifie que la ligne Excel correspond toujours au même dossier / champ.
        if "FICHIER" in headers and "CHAMP" in headers:
            f = ws.cell(row=row, column=headers["FICHIER"]).value
            c = ws.cell(row=row, column=headers["CHAMP"]).value
            if str(f) != str(r.get("FICHIER")) or str(c) != str(r.get("CHAMP")):
                raise RuntimeError(f"Ligne {row} désalignée ({f}/{c}) — rechargez le classeur.")
        for col in EDITABLE:
            new = r.get(col)
            new = None if is_blank(new) else str(new).strip()
            if col == "DECISION_MANUELLE" and new:
                new = new.upper()
            cell = ws.cell(row=row, column=headers[col])
            old = None if is_blank(cell.value) else str(cell.value)
            if old != new:
                cell.value = new
                changes.append({"row": row, "fichier": r.get("FICHIER"), "champ": r.get("CHAMP"),
                                "col": col, "old": old, "new": new})
    if changes:
        wb.save(xlsx)
    return {"changes": changes, "backup": str(bkp) if bkp else None}


def progress(manual: pd.DataFrame) -> pd.DataFrame:
    """Avancement de la validation par dossier."""
    if manual.empty:
        return pd.DataFrame(columns=["FICHIER", "LIGNES", "DECIDEES", "REJETEES", "RESTANT"])
    m = manual.copy()
    m["_DEC"] = m["DECISION_MANUELLE"].fillna("").astype(str).str.strip().str.upper()
    g = m.groupby("FICHIER", dropna=False)
    out = pd.DataFrame({
        "LIGNES": g.size(),
        "DECIDEES": g["_DEC"].apply(lambda s: (s != "").sum()),
        "REJETEES": g["_DEC"].apply(lambda s: (s == "REJETEE").sum()),
    }).reset_index()
    out["RESTANT"] = out["LIGNES"] - out["DECIDEES"]
    out["STATUT_VALIDATION"] = out.apply(
        lambda r: "REJETE" if r.REJETEES else ("TERMINE" if r.RESTANT == 0 else "EN_COURS"), axis=1)
    return out.sort_values(["RESTANT", "FICHIER"], ascending=[False, True])


def backups(xlsx: Path) -> list[Path]:
    d = xlsx.parent / "_backups"
    return sorted(d.glob(f"{xlsx.stem}_*{xlsx.suffix}"), key=lambda p: p.stat().st_mtime, reverse=True) \
        if d.exists() else []


def carry_over(previous: Path, current: Path) -> pd.DataFrame:
    """Prépare le report des décisions d'un ancien classeur vers le classeur régénéré.

    Une décision n'est reportée que si la ligne est identique (FICHIER, CHAMP, PAGE, RAW)
    et encore vide dans le nouveau classeur : si l'OCR a changé, on revalide.
    Retourne le DataFrame à passer à save_decisions (lignes concernées uniquement).
    """
    old, new = load_manual(previous), load_manual(current)
    if old.empty or new.empty:
        return new.iloc[0:0]
    key = [c for c in ("FICHIER", "CHAMP", "PAGE", "RAW") if c in old.columns and c in new.columns]
    o = old[old["DECISION_MANUELLE"] != ""].copy()
    for df in (o, new):
        df["_K"] = df[key].apply(lambda r: "|".join("" if is_blank(v) else str(v) for v in r), axis=1)
    o = o.drop_duplicates("_K", keep="last").set_index("_K")
    todo = new[(new["DECISION_MANUELLE"] == "") & new["_K"].isin(o.index)].copy()
    for c in EDITABLE:
        todo[c] = todo["_K"].map(o[c])
    return todo.drop(columns=["_K"])
