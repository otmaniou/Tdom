"""Module 3 — Contrôles des dossiers de transfert (OV / Annexe I / Annexe II / Bulletin).

Entrée  : classeur audit_transferts_*.xlsx produit par le notebook OV (v26 ou v27, extraction VLM)
          + Planning_TL (Module 2) + référentiels ATF / autorisations (optionnels)
          + registre des TL déjà validés.
Sortie  : une synthèse par dossier, le détail des contrôles, puis un registre.

Principe (hérité de la Partie 2) : une valeur OCR absente ou illisible n'est JAMAIS
un motif de rejet automatique -> REVIEW. Seul un écart objectivable sur des valeurs
lues (plafond dépassé, doublon, autorisation échue...) est BLOQUANT. La décision
finale reste humaine.
"""
from __future__ import annotations

import hashlib
import re
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd

from .common import (digits, is_blank, name_similarity, parse_amount,
                     parse_date, siege_racine)
from .montant_lettres import lettres_to_number

# Schéma du classeur v26 (cellule « Export Excel ») : (groupe, clé, libellé)
V26_SCHEMA = {
    "META": [("fichier", "Fichier"), ("date_traitement", "Date traitement"),
             ("pages_trouvees", "Pages trouvées"), ("pages_manquantes", "Pages manquantes"),
             ("anomalies", "Anomalies")],
    "OV": [("monnaie", "Monnaie"), ("montant_chiffres", "Montant (chiffres)"),
           ("montant_lettres", "Montant (lettres)"), ("periode", "Période"), ("tranche", "Tranche"),
           ("date_demande", "Date demande"), ("compte_debiteur", "Compte débiteur"),
           ("beneficiaire_nom", "Bénéficiaire nom"), ("beneficiaire_compte", "Bénéficiaire compte"),
           ("beneficiaire_adresse", "Bénéficiaire adresse"), ("swift", "SWIFT"),
           ("banque_nom", "Banque bénéficiaire"), ("status", "Statut"),
           ("validation_errors", "Erreurs validation")],
    "ANNEXE_I": [("nom_prenom", "Nom Prénom"), ("date_signature", "Date signature"),
                 ("status", "Statut"), ("validation_errors", "Erreurs validation")],
    "ANNEXE_II": [("mois_transfert", "Mois transfert"), ("nom_prenom", "Nom Prénom"),
                  ("compte_bancaire", "Compte bancaire"), ("salaire_mensuel", "Salaire mensuel"),
                  ("nombre_jours", "Nombre de jours"), ("part_transferable", "Part transférable"),
                  ("pays_destination", "Pays destination"), ("compte_devise", "Compte devise"),
                  ("numero_domiciliation", "N° domiciliation"),
                  ("date_domiciliation", "Date domiciliation"), ("status", "Statut"),
                  ("validation_errors", "Erreurs validation")],
    "BULLETIN": [("nom_prenom", "Nom Prénom"), ("matricule", "Matricule"),
                 ("mois_bulletin", "Mois bulletin"), ("salaire_base", "Salaire base"),
                 ("salaire_brut", "Salaire brut"), ("retenue_ss", "Retenue SS"),
                 ("retenue_irg", "Retenue IRG"), ("retenue_mutuelle", "Retenue mutuelle"),
                 ("net_a_payer", "Net à payer")],
}

# Classeur v27 (pipeline_ocr_vl27) : autres libellés / clés -> mêmes clés canoniques que v26,
# pour que les contrôles et la matrice restent inchangés. (groupe, libellé) -> clé canonique
V27_LABELS = {
    ("META", "Temps (s)"): "temps_total_s", ("META", "Tokens total"): "tokens_total",
    ("OV", "periode"): "periode", ("OV", "mois"): "mois", ("OV", "annee"): "annee",
    ("OV", "Reference paiement (case 70)"): "reference_paiement",
    ("OV", "complement"): "complement", ("OV", "compte Donneur ordre"): "compte_debiteur",
    ("OV", "Devise de transfert"): "devise_transfert", ("OV", "SWIFT"): "swift",
    ("OV", "Banque bénéficiaire"): "banque_nom",
    ("ANNEXE_II", "Nom Prénom annex2"): "nom_prenom", ("ANNEXE_II", "Compte bancaire"): "compte_bancaire",
    ("ANNEXE_II", "Nom_banque"): "banque_etranger", ("ANNEXE_II", "Numéro IBAN"): "compte_devise",
    ("ANNEXE_II", "Domiciliation import"): "numero_domiciliation",
    ("BULLETIN", "Nom Prénom Salarie"): "nom_prenom", ("BULLETIN", "Periode bulletin"): "mois_bulletin",
    ("ANNEXE_I", "Nom prenom client"): "nom_prenom", ("ANNEXE_I", "Compte du Client"): "compte_bancaire",
    ("ANNEXE_I", "Nom prenom Signataire"): "nom_prenom_signataire", ("ANNEXE_I", "Date de signature"): "date_signature",
    # v29 : page de chaque pièce, case nature du paiement, jours d'absence, date de naissance / résidence (Annexe I)
    ("META", "Pages par pièce"): "pages_pieces", ("OV", "Nature paiement (case)"): "nature_paiement_case",
    ("ANNEXE_II", "Nombre jours absence"): "nombre_jours_absence",
    ("ANNEXE_I", "Date de naissance"): "date_naissance", ("ANNEXE_I", "Résidence"): "residence",
    # v30 : « Devise de transfert » = la devise seule (USD…), texte de la case Autre dans une colonne à part
    ("OV", "Nature paiement (autre libellé)"): "nature_paiement_autre_libelle",
}

MOIS = {"JANVIER": 1, "FEVRIER": 2, "MARS": 3, "AVRIL": 4, "MAI": 5, "JUIN": 6, "JUILLET": 7,
        "AOUT": 8, "SEPTEMBRE": 9, "OCTOBRE": 10, "NOVEMBRE": 11, "DECEMBRE": 12,
        "JAN": 1, "FEV": 2, "AVR": 4, "JUIL": 7, "SEPT": 9, "OCT": 10, "NOV": 11, "DEC": 12,
        "JANV": 1, "FEVR": 2, "SEP": 9}

ATF_COLUMNS = ["NOM_PRENOM", "COMPTE", "MOIS", "SALAIRE_NET", "IRG", "MONTANT_TRANSFERABLE", "ATF_REFERENCE"]
AUT_COLUMNS = ["NOM_PRENOM", "COMPTE", "MOIS", "MONTANT_AUTORISE", "DATE_VALIDATION", "DATE_LIMITE",
               "AUTORISATION_REFERENCE"]
REQUIRED = {"ATF": ["NOM_PRENOM", "MOIS", "MONTANT_TRANSFERABLE"],
            "Autorisations": ["NOM_PRENOM", "MOIS", "MONTANT_AUTORISE"]}
REGISTRE_COLUMNS = ["TL_ID", "FICHIER", "DOM_ID", "PERIODE_ID", "PERIODE_TL", "SIEGE_RACINE", "NOM",
                    "NUMERO_DOMICILIATION", "MONTANT", "DEVISE", "PROVISION_ATLAS", "STATUT_CONTROLE",
                    "DECISION", "COMMENTAIRE", "VALIDE_PAR", "DATE_VALIDATION", "STATUT_EXECUTION",
                    "STATUT_VIGILANCE"]


# ------------------------------------------------------------------ lecture v26
def read_audit_xlsx(path: Path) -> pd.DataFrame:
    """Aplatis le classeur à double en-tête en colonnes « GROUPE.clé »."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Dossiers"] if "Dossiers" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return pd.DataFrame()
    groups, g = [], None
    for v in rows[0]:
        g = v if v not in (None, "") else g
        groups.append(g)
    label_to_key = {(grp, lab): key for grp, cols in V26_SCHEMA.items() for key, lab in cols}
    label_to_key.update(V27_LABELS)
    cols = []
    for grp, lab in zip(groups, rows[1]):
        key = label_to_key.get((grp, lab), re.sub(r"\W+", "_", str(lab or "")).lower())
        cols.append(key if grp == "META" else f"{grp}.{key}")
    df = pd.DataFrame([r for r in rows[2:] if any(x is not None for x in r)], columns=cols)
    df["AUDIT_SOURCE"] = path.name
    return df


def latest_audit(folder: Path) -> list[Path]:
    if not folder.exists():
        return []
    return sorted(folder.glob("audit_transferts_*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)


# ------------------------------------------------------------------ périodes
def parse_periode(v) -> tuple[str | None, str | None]:
    """'MARS 2026 PART 1' -> ('2026-03','P1') ; '03/2026' -> ('2026-03', None) ; règle unique : tlone.periode.lire."""
    from .periode import lire
    r = lire(v)
    return r["mois"], r["tranche"]


def periode_from_filename(name: str) -> tuple[str | None, str | None]:
    """Formats observés : ...-202603P1, ...-202606-OV-PART-2_NOM, ...-042026P2, ...-04 P1."""
    stem = Path(str(name)).stem.upper()
    part = None
    mp = re.search(r"(?:P|PART)[\s_-]*([12])(?![0-9])", stem)
    if mp:
        part = f"P{mp.group(1)}"
    m = re.search(r"(?<![0-9])(20\d{2})(0[1-9]|1[0-2])(?![0-9])", stem)      # AAAAMM
    if m:
        return f"{m.group(1)}-{m.group(2)}", part
    m = re.search(r"(?<![0-9])(0[1-9]|1[0-2])(20\d{2})(?![0-9])", stem)      # MMAAAA
    if m:
        return f"{m.group(2)}-{m.group(1)}", part
    return None, part


def periode_label(mois: str | None, part: str | None) -> str | None:
    if not mois:
        return None
    return mois + (part or "")


# ------------------------------------------------------------------ référentiels
# en-têtes des fichiers réels (ATF BNP mensuel, listes d'autorisations) -> colonnes canoniques
REF_ALIASES = {
    "NOM ET PRENOM": "NOM_PRENOM", "NOM PRENOM": "NOM_PRENOM", "NOM ET PRENOMS": "NOM_PRENOM",
    "NOM PRENOM TRAVAILLEUR": "NOM_PRENOM", "TRAVAILLEUR": "NOM_PRENOM", "BENEFICIAIRE": "NOM_PRENOM",
    "MATRICULE SAP": "MATRICULE", "MATRICULE": "MATRICULE", "PERIODE": "MOIS", "MOIS": "MOIS",
    "ID TIERS": "ID_TIERS", "IDTIERS": "ID_TIERS", "IDENTIFIANT TIERS": "ID_TIERS", "IDENTIFIANT CLIENT": "ID_TIERS",
    "MOIS DE TRANSFERT": "MOIS", "SALAIRE NET": "SALAIRE_NET", "MONTANT TRANSFERABLE": "MONTANT_TRANSFERABLE",
    "MONTANT TRANSFERT": "MONTANT_TRANSFERABLE", "IRG": "IRG", "NET IMPOSABLE": "NET_IMPOSABLE",
    "COMPTE": "COMPTE", "COMPTE BANCAIRE": "COMPTE", "RIB": "COMPTE", "MONTANT AUTORISE": "MONTANT_AUTORISE",
    "MONTANT": "MONTANT_AUTORISE", "DATE VALIDATION": "DATE_VALIDATION", "DATE DE VALIDATION": "DATE_VALIDATION",
    "DATE LIMITE": "DATE_LIMITE", "DATE D EXPIRATION": "DATE_LIMITE", "REFERENCE": "ATF_REFERENCE",
    "N AUTORISATION": "AUTORISATION_REFERENCE", "NUMERO AUTORISATION": "AUTORISATION_REFERENCE",
}


def read_any_table(path: Path) -> pd.DataFrame:
    """xlsx / xls / csv, ou HTML (liste d'autorisations enregistrée en page web) : premier tableau trouvé."""
    p = Path(path)
    if p.suffix.lower() in (".html", ".htm"):
        from . import autorisations
        raw = p.read_bytes()
        if autorisations.est_portail(raw):      # page du portail « Transfert des salaires des expatriés »
            res = autorisations.lire(raw, nom_fichier=p.name)
            if res["lignes"].empty:
                raise ValueError("Page du portail sans tableau d'articles : " + " ; ".join(res["controles"]))
            df = autorisations.vers_referentiel(res)[0]
            meta = p.with_name(p.name + ".json")          # date de validation saisie au dépôt (absente de la page)
            if meta.exists():
                import json as _json
                m = _json.loads(meta.read_text(encoding="utf-8"))
                df["DATE_VALIDATION"] = m.get("date_validation")
                if m.get("annee") and df["MOIS"].isna().any():
                    res = autorisations.lire(raw, annee=int(m["annee"]), nom_fichier=p.name)
                    df = autorisations.vers_referentiel(res)[0].assign(DATE_VALIDATION=m.get("date_validation"))
            return df
        from html.parser import HTMLParser

        class T(HTMLParser):
            def __init__(self):
                super().__init__()
                self.tables, self.row, self.cell, self.inc = [], None, None, False

            def handle_starttag(self, tag, a):
                if tag == "table":
                    self.tables.append([])
                elif tag == "tr" and self.tables:
                    self.row = []
                elif tag in ("td", "th") and self.row is not None:
                    self.cell, self.inc = [], True

            def handle_endtag(self, tag):
                if tag in ("td", "th") and self.row is not None and self.cell is not None:
                    self.row.append(" ".join("".join(self.cell).split()))
                    self.cell, self.inc = None, False
                elif tag == "tr" and self.row is not None and self.tables:
                    if any(self.row):
                        self.tables[-1].append(self.row)
                    self.row = None

            def handle_data(self, d):
                if self.inc and self.cell is not None:
                    self.cell.append(d)
        t = T()
        t.feed(p.read_text(encoding="utf-8", errors="replace"))
        tabs = [x for x in t.tables if len(x) >= 2]
        if not tabs:
            raise ValueError("Aucun tableau dans la page HTML (page enregistrée sans son contenu ?). "
                             "Enregistrer la liste au format Excel ou CSV.")
        tab = max(tabs, key=len)
        w = max(len(r) for r in tab)
        return pd.DataFrame([r + [None] * (w - len(r)) for r in tab[1:]], columns=tab[0] + [f"C{i}" for i in range(len(tab[0]), w)])
    from .common import cached_read
    if p.suffix.lower() == ".csv":
        return cached_read(p, "table", lambda q: pd.read_csv(q, dtype=object, sep=None, engine="python")).copy()
    return cached_read(p, "table", lambda q: pd.read_excel(q, dtype=object)).copy()


def _mois_ref(v):
    from datetime import date as _d, datetime as _dt
    if isinstance(v, (_dt, _d, pd.Timestamp)):
        return f"{v.year:04d}-{v.month:02d}"
    return parse_periode(v)[0]


def load_referential(path: Path, columns: list[str], rh: pd.DataFrame | None = None) -> pd.DataFrame:
    """Référentiel ATF / autorisations, au modèle TL-One ou au format reçu (en-têtes reconnus par alias).
    Lien au salarié, dans cet ordre : ID tiers (colonne ID_TIERS « 07410 109109 ») ; sinon, avec la base RH, le
    matricule SAP puis le nom exact. Le montant est celui du mois entier (ATF / autorisations non réparties P1 / P2)."""
    if not path or not Path(path).exists():
        return pd.DataFrame(columns=columns + ["_MOIS", "_COMPTE", "_CLIENT"])
    from .rh import hkey
    df = read_any_table(Path(path))
    ren = {}
    for c in df.columns:
        k = str(c).strip().upper()
        if k in columns:
            continue
        a = REF_ALIASES.get(hkey(c))
        if a and a not in df.columns and a not in ren.values():
            ren[c] = a
    df = df.rename(columns=ren)
    df.columns = [str(c).strip().upper() if str(c).strip().upper() in columns + ["MATRICULE", "ID_TIERS"] else c
                  for c in df.columns]
    for c in columns:
        if c not in df.columns:
            df[c] = None
    if "MATRICULE" in df.columns and rh is not None and not rh.empty:
        mat = {str(m): r for m, r in zip(rh["MATRICULE"], rh["RIB"]) if not is_blank(m)}
        df["COMPTE"] = [c if not is_blank(c) else mat.get(str(m).split(".")[0]) for c, m in zip(df["COMPTE"], df["MATRICULE"])]
    if rh is not None and not rh.empty and df["COMPTE"].map(is_blank).any():
        # sans compte ni matricule (autorisations du portail) : rapprochement par le nom exact de la base RH
        from .common import norm_name
        par_nom: dict = {}
        for n, r in zip(rh["NOM_PRENOM"], rh["RIB"]):
            par_nom.setdefault(norm_name(n), set()).add(r)
        df["COMPTE"] = [c if not is_blank(c) else (next(iter(par_nom.get(norm_name(n), []))) if len(par_nom.get(norm_name(n), [])) == 1 else None)
                        for c, n in zip(df["COMPTE"], df["NOM_PRENOM"])]
    from .common import id_tiers
    from .rh import rib_parts
    it = df["ID_TIERS"].map(lambda v: None if is_blank(v) else id_tiers(v)) if "ID_TIERS" in df.columns \
        else pd.Series([None] * len(df), index=df.index)
    if rh is not None and not rh.empty and {"CLIENT", "RIB"} <= set(rh.columns):   # compte retrouvé par l'ID tiers
        rib_cl = {c: r for c, r in zip(rh["CLIENT"], rh["RIB"]) if isinstance(c, str) and not is_blank(r)}
        df["COMPTE"] = [c if not is_blank(c) else rib_cl.get(t) for c, t in zip(df["COMPTE"], it)]
    df["_MOIS"] = df["MOIS"].map(_mois_ref)
    df["_COMPTE"] = df["COMPTE"].map(digits)
    df["_CLIENT"] = [t or (rib_parts(c) or {}).get("CLIENT") for t, c in zip(it, df["_COMPTE"])]
    return df


def template(columns: list[str]) -> pd.DataFrame:
    return pd.DataFrame(columns=columns)


def _match_ref(ref: pd.DataFrame, mois: str | None, compte: str, nom, seuil: float, client: str | None = None):
    if ref is None or ref.empty or not mois:
        return None, None
    cand = ref[ref["_MOIS"] == mois]
    if cand.empty:
        return None, None
    if client and "_CLIENT" in cand.columns:
        by_cl = cand[cand["_CLIENT"] == client]
        if len(by_cl) == 1:
            return by_cl.iloc[0], "ID TIERS"
        if len(by_cl) > 1:      # plusieurs lignes du mois (P1 / P2 listées à part) : plafond = somme du mois
            r = by_cl.iloc[0].copy()
            for k in ("MONTANT_TRANSFERABLE", "MONTANT_AUTORISE"):
                if k in by_cl.columns:
                    r[k] = round(sum(parse_amount(x) or 0 for x in by_cl[k]), 2)
            return r, f"ID TIERS ({len(by_cl)} lignes du mois additionnées)"
    if compte:
        by_compte = cand[cand["_COMPTE"] == compte]
        if not by_compte.empty:
            return by_compte.iloc[0], "COMPTE"
    best, best_s = None, 0.0
    for _, r in cand.iterrows():
        sim = name_similarity(nom, r.get("NOM_PRENOM")) or 0.0
        if sim > best_s:
            best, best_s = r, sim
    if best is not None and best_s >= seuil:
        return best, f"NOM ({best_s:.2f})"
    return None, None


# ------------------------------------------------------------------ registre
def load_registre(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame(columns=REGISTRE_COLUMNS)
    return pd.read_csv(path, dtype=str, encoding="utf-8-sig").reindex(columns=REGISTRE_COLUMNS)


def _registre_valide(reg: pd.DataFrame) -> pd.DataFrame:
    if reg.empty:
        return reg
    r = reg[reg["DECISION"].fillna("") == "VALIDE"].copy()
    r["_MONTANT"] = r["MONTANT"].map(parse_amount).fillna(0.0)
    return r


def tl_id(fichier: str, periode_id: str | None) -> str:
    return hashlib.sha256(f"{fichier}|{periode_id or ''}".encode()).hexdigest()[:20]


# ------------------------------------------------------------------ contrôles
class _Checks:
    def __init__(self, fichier: str):
        self.fichier, self.rows = fichier, []

    def add(self, code, libelle, ref, severite, resultat, detail=""):
        self.rows.append({"FICHIER": self.fichier, "CODE": code, "CONTROLE": libelle, "REFERENCE": ref,
                          "SEVERITE": severite, "RESULTAT": resultat, "DETAIL": detail})

    def cmp_cap(self, code, libelle, ref, montant, cap, tol, missing_msg, severite_missing="REVIEW"):
        if montant is None:
            self.add(code, libelle, ref, "REVIEW", "NON_CONTROLABLE", "Montant OV illisible")
        elif cap is None:
            self.add(code, libelle, ref, severite_missing, "NON_CONTROLABLE", missing_msg)
        elif montant <= cap + tol:
            self.add(code, libelle, ref, "BLOQUANT", "OK", f"{montant:,.2f} ≤ {cap:,.2f}")
        else:
            self.add(code, libelle, ref, "BLOQUANT", "KO", f"{montant:,.2f} > {cap:,.2f} (écart {montant - cap:,.2f})")


def _g(d: dict, k: str):
    v = d.get(k)
    return None if is_blank(v) else v


# Devise de transfert : l'OV la porte dans le libellé à côté de la case « Autre » de la zone nature du paiement
# (v29 : nature_paiement_autre_libelle, colonne « Devise de transfert »), ex. « Transfert en USD ». Seul le code de la
# devise est retenu (USD, EUR…) pour être comparé à la base RH.
DEVISES = ("USD", "EUR", "GBP", "CHF", "CAD", "TRY", "AED", "SAR", "CNY", "JPY", "INR")
_DEVISE_MOTS = {"DOLLARS": "USD", "DOLLAR": "USD", "EUROS": "EUR", "EURO": "EUR", "LIVRES": "GBP", "LIVRE": "GBP"}


def devise_transfert(v) -> str | None:
    """« Transfert en USD » -> USD ; « transfert en euros » -> EUR ; « € » -> EUR ; texte sans devise -> None."""
    if is_blank(v):
        return None
    from .common import strip_accents
    s = strip_accents(str(v)).upper()
    m = re.search(r"(?<![A-Z])(" + "|".join(DEVISES) + r")(?![A-Z])", s)
    if m:
        return m.group(1)
    for mot, code in _DEVISE_MOTS.items():
        if re.search(rf"\b{mot}\b", s):
            return code
    if "€" in s:
        return "EUR"
    if "$" in s:
        return "USD"
    return None


_BANQUE_VIDES = {"BANK", "BANKASI", "BANKASY", "BANQUE", "BANCO", "BANCA", "BANKA", "THE", "OF", "DE", "DU", "LA", "LE",
                 "AS", "SA", "SAS", "AG", "PLC", "LTD", "LIMITED", "TAS", "TC", "INC", "CO", "NV", "SPA", "GROUP"}


def banque_similarite(a, b) -> float | None:
    """Similarité de deux noms de banque, sans les mots génériques (BANK, BANKASI, A.Ş., T.C.…) :
    « ZIRAAT BANKASI » / « T.C. ZIRAAT BANKASI A.S. » -> 1.0. None si une valeur manque."""
    from difflib import SequenceMatcher

    from .common import strip_accents

    def toks(v):
        s = re.sub(r"[^A-Z0-9 ]+", " ", strip_accents(str(v)).upper().replace(".", ""))
        return [t for t in s.split() if t not in _BANQUE_VIDES]
    if is_blank(a) or is_blank(b):
        return None
    ta, tb = toks(a), toks(b)
    if not ta or not tb:
        return None
    inclus = len(set(ta) & set(tb)) / min(len(set(ta)), len(set(tb)))
    return round(max(inclus, SequenceMatcher(None, " ".join(ta), " ".join(tb)).ratio()), 3)


def date_domiciliation(prow, dom_dates: dict | None):
    """Date de la domiciliation (fichier CarthagoDom : n° DOM -> date ; à défaut colonne du Planning)."""
    if prow is None:
        return None
    if dom_dates:
        k = digits(prow.get("NUMERO_DOMICILIATION"))
        if k and (k + "DZD") in dom_dates:
            return dom_dates[k + "DZD"]
    return parse_date(prow.get("DATE_DOMICILIATION")) if "DATE_DOMICILIATION" in prow else None


def control_dossier(d: dict, planning: pd.DataFrame, reg_ok: pd.DataFrame, atf: pd.DataFrame,
                    aut: pd.DataFrame, s, today: date | None = None, rh: pd.DataFrame | None = None,
                    dom_dates: dict | None = None) -> tuple[dict, list[dict]]:
    today = today or date.today()
    fichier = str(d.get("fichier"))
    ck = _Checks(fichier)
    tol = s.tolerance_montant_dzd

    # ---------- T01 pièces
    manq = [x.strip() for x in str(_g(d, "pages_manquantes") or "").split(",") if x.strip()]
    bloq = [x for x in manq if x in ("OV", "ANNEXE_II", "BULLETIN")]
    ck.add("T01", "Pièces du dossier présentes", "REG-11 / REG-12 / AS-04",
           "BLOQUANT" if bloq else ("REVIEW" if manq else "BLOQUANT"),
           "KO" if manq else "OK", ("Manquantes : " + ", ".join(manq)) if manq else "OV, Annexe I, Annexe II, Bulletin")
    if _g(d, "anomalies") is not None:
        ck.add("T01b", "Pages en doublon", "AS-04", "REVIEW", "KO", str(d.get("anomalies")))

    # ---------- T02 mentions obligatoires OV
    oblig = {"OV.montant_chiffres": "montant", "OV.montant_lettres": "montant en lettres",
             "OV.periode": "période", "OV.compte_debiteur": "compte débiteur",
             "OV.beneficiaire_nom": "bénéficiaire", "OV.beneficiaire_compte": "compte bénéficiaire",
             "OV.swift": "SWIFT", "OV.banque_nom": "banque"}
    absents = [lab for k, lab in oblig.items() if _g(d, k) is None]
    ck.add("T02", "Mentions obligatoires OV lues", "CTL-09", "REVIEW", "KO" if absents else "OK",
           ("Non lues (OCR ou absentes) : " + ", ".join(absents)) if absents else "Toutes lues")

    montant = parse_amount(_g(d, "OV.montant_chiffres"))
    devise = (str(_g(d, "OV.monnaie") or "")).upper() or None

    # ---------- T03 chiffres / lettres
    ml = lettres_to_number(_g(d, "OV.montant_lettres"))
    if montant is None or _g(d, "OV.montant_lettres") is None:
        ck.add("T03", "Montant chiffres = lettres", "CTL-09", "REVIEW", "NON_CONTROLABLE", "Une des deux mentions non lue")
    elif ml is None:
        ck.add("T03", "Montant chiffres = lettres", "CTL-09", "REVIEW", "NON_CONTROLABLE",
               f"Lettres non interprétables : « {_g(d, 'OV.montant_lettres')} »")
    else:
        ok = abs(ml - montant) <= tol
        ck.add("T03", "Montant chiffres = lettres", "CTL-09", "BLOQUANT", "OK" if ok else "KO",
               f"chiffres {montant:,.2f} / lettres {ml:,.2f}")

    # ---------- T17 devise
    if not devise:
        ck.add("T17", "Devise de l'OV", "AS-04", "REVIEW", "NON_CONTROLABLE", "non lue")
    elif devise in ("DZD", "DA"):
        ck.add("T17", "Devise de l'OV", "AS-04", "REVIEW", "OK", devise)
    else:
        ck.add("T17", "Devise de l'OV", "AS-04", "REVIEW", "KO", f"{devise} — plafonds exprimés en DZD, conversion à vérifier")

    # ---------- Périodes
    # Case 70 « Référence du paiement et observation du donneur d'ordre » : texte brut (v28 reference_paiement,
    # v26/v27 periode) ; période, tranche et complément sont déduits ici par une règle unique (tlone.periode).
    from .periode import lire
    z70 = _g(d, "OV.reference_paiement") or _g(d, "OV.periode")
    lu = lire(z70)
    p_ov = (lu["mois"], lu["tranche"])
    p_mod = (parse_periode(f"{_g(d, 'OV.mois') or ''} {_g(d, 'OV.annee') or ''}")[0] if _g(d, "OV.mois") else None,
             parse_periode("P " + str(_g(d, "OV.tranche")))[1] if _g(d, "OV.tranche") else None)
    if not p_ov[0] and p_mod[0]:   # v27 : mois / année extraits séparément par le modèle
        p_ov = (p_mod[0], p_ov[1])
    if not p_ov[1] and p_mod[1]:
        p_ov = (p_ov[0], p_mod[1])
    complement = lu["complement"] or (None if is_blank(_g(d, "OV.complement")) else "COMPLEMENT")
    ecart_modele = [x for x, a, b in (("mois", lu["mois"], p_mod[0]), ("tranche", lu["tranche"], p_mod[1]))
                    if a and b and a != b]
    p_a2 = parse_periode(_g(d, "ANNEXE_II.mois_transfert"))
    p_bu = parse_periode(_g(d, "BULLETIN.mois_bulletin"))
    p_fn = periode_from_filename(fichier)
    mois = p_ov[0] or p_a2[0] or p_fn[0]
    part = p_ov[1] or p_a2[1] or p_fn[1]
    vus = {"OV": p_ov[0], "Annexe II": p_a2[0], "Bulletin": p_bu[0], "Nom fichier": p_fn[0]}
    distincts = {v for v in vus.values() if v}
    ck.add("T07", "Période cohérente entre pièces", "CTL-11 / CTL-12", "REVIEW",
           "OK" if len(distincts) == 1 else ("NON_CONTROLABLE" if not distincts else "KO"),
           " | ".join(f"{k}={v or '?'}" for k, v in vus.items()) + (f" | partie={part}" if part else "")
           + (f" | case 70 « {z70} »" if not is_blank(z70) else "")
           + (f" | lecture du modèle différente ({', '.join(ecart_modele)}) : texte de la case retenu" if ecart_modele else "")
           + (f" | {lu['regle']}" if lu["statut"] == "A_PRECISER" else ""))

    # ---------- Comptes / siège
    c_ov = digits(_g(d, "OV.compte_debiteur"))
    c_a2 = digits(_g(d, "ANNEXE_II.compte_bancaire"))
    siege = siege_racine(c_ov) or siege_racine(c_a2) or siege_racine(Path(fichier).stem)

    # ---------- T04 rattachement Planning
    prow, plan_dom = None, pd.DataFrame()
    if planning is not None and not planning.empty:
        pl = planning
        cand = pl[pl["SIEGE_RACINE"].astype(str) == str(siege)] if siege else pl.iloc[0:0]
        if cand.empty and (c_ov or c_a2):
            cand = pl[pl["DOM_COMPTE_LOCAL"].map(digits).isin([x for x in (c_ov, c_a2) if x])]
        plan_dom = cand
        if not cand.empty and mois:
            m = cand[cand["MOIS_BASE"] == mois]
            parties = m["PARTIE"].fillna("").astype(str)
            if part and (parties == part).any():
                m = m[parties == part]
            elif part and len(m) and (parties != "").all():
                # OV d'une partie (ex. P2) que la DOM ne couvre pas (DOM jusqu'au 17 : seule P1 planifiée)
                m = m.iloc[0:0]
            elif len(m) > 1 and (parties == "").any():
                m = m[parties == ""]
            if len(m) == 1:
                prow = m.iloc[0]
            elif len(m) > 1:
                ck.add("T04", "Rattachement au Planning TL", "AS-04 / AS-07", "REVIEW", "KO",
                       f"Période {mois} fractionnée ({', '.join(m['PERIODE_TL'])}) : préciser la partie sur l'OV")
    if prow is not None:
        ck.add("T04", "Rattachement au Planning TL", "AS-04", "BLOQUANT", "OK",
               f"{prow['PERIODE_TL']} — DOM {prow['NUMERO_DOMICILIATION']}")
    elif not any(r["CODE"] == "T04" for r in ck.rows):
        if planning is None or planning.empty:
            detail = "Planning TL non généré"
        elif plan_dom.empty:
            detail = f"Aucune domiciliation planifiée pour siège {siege or '?'} / compte {c_ov or c_a2 or '?'}"
        else:
            detail = (f"DOM trouvée mais période {periode_label(mois, part) or '?'} hors fenêtre "
                      f"({plan_dom['PERIODE_TL'].iloc[0]} → {plan_dom['PERIODE_TL'].iloc[-1]})")
        ck.add("T04", "Rattachement au Planning TL", "AS-04", "BLOQUANT", "KO", detail)

    # ---------- T05 comptes
    c_dom = digits(prow["DOM_COMPTE_LOCAL"]) if prow is not None else ""
    vals = {"OV": c_ov, "Annexe II": c_a2, "DOM": c_dom}
    present = {k: v for k, v in vals.items() if v}
    if len(present) < 2:
        ck.add("T05", "Compte local identique OV / Annexe II / DOM", "CTL-17", "REVIEW", "NON_CONTROLABLE",
               " | ".join(f"{k}={v or '?'}" for k, v in vals.items()))
    else:
        ok = len(set(present.values())) == 1
        ck.add("T05", "Compte local identique OV / Annexe II / DOM", "CTL-17", "REVIEW",
               "OK" if ok else "KO", " | ".join(f"{k}={v or '?'}" for k, v in vals.items()))

    # T16 (n° domiciliation Annexe II = DOM) : retiré — non contrôlé (décision métier du 19/10/2026).

    # ---------- T06 identité
    # L'OV ne porte pas le nom du donneur d'ordre (le salarié), seulement celui du bénéficiaire : il n'entre pas
    # dans le contrôle d'identité (le bénéficiaire est contrôlé contre la base RH, R06).
    ref_nom = prow["NOM_TRAVAILLEUR"] if prow is not None else _g(d, "ANNEXE_II.nom_prenom")
    noms = {"Annexe I": _g(d, "ANNEXE_I.nom_prenom"),
            "Annexe II": _g(d, "ANNEXE_II.nom_prenom"), "Bulletin": _g(d, "BULLETIN.nom_prenom")}
    sims = {k: name_similarity(ref_nom, v) for k, v in noms.items()}
    faibles = [f"{k} ({v:.2f})" for k, v in sims.items() if v is not None and v < s.seuil_similarite_nom]
    if not ref_nom:
        ck.add("T06", "Identité travailleur concordante", "REG-01", "REVIEW", "NON_CONTROLABLE", "Nom de référence absent")
    else:
        ck.add("T06", "Identité travailleur concordante", "REG-01", "REVIEW", "KO" if faibles else "OK",
               f"Réf. {'DOM/CTS' if prow is not None else 'Annexe II'} : {ref_nom}"
               + (" | écarts : " + ", ".join(faibles) if faibles else ""))

    # T08 (net bulletin = salaire Annexe II) : retiré — l'Annexe II donne un salaire hors primes, le
    # bulletin un net avec primes : un écart est normal (décision métier du 19/10/2026).
    net = parse_amount(_g(d, "BULLETIN.net_a_payer"))
    sal_a2 = parse_amount(_g(d, "ANNEXE_II.salaire_mensuel"))

    # ---------- Plafonds REG-05 / CTL-15
    caps: dict[str, float] = {}
    cap_a2 = parse_amount(_g(d, "ANNEXE_II.part_transferable"))
    ck.cmp_cap("T09", "OV ≤ part transférable Annexe II", "REG-05 / CTL-11", montant, cap_a2, tol,
               "Part transférable Annexe II non lue")
    if cap_a2 is not None:
        caps["Annexe II"] = cap_a2

    assiette = net if net is not None else sal_a2
    cap95 = round(assiette * s.plafond_pct_net, 2) if assiette is not None else None
    ck.cmp_cap("T10", f"OV ≤ {s.plafond_pct_net:.0%} du net perçu", "REG-05 (assiette à confirmer, Q-01)",
               montant, cap95, tol, "Net non lu (bulletin et Annexe II)")
    if cap95 is not None:
        caps[f"{s.plafond_pct_net:.0%} net"] = cap95

    deja_periode = 0.0
    cap_plan = None
    if prow is not None:
        if not reg_ok.empty:
            deja_periode = float(reg_ok.loc[(reg_ok["PERIODE_ID"] == prow["PERIODE_ID"]) & (reg_ok["FICHIER"] != fichier),
                                            "_MONTANT"].sum())
        cap_plan = round(float(prow["MONTANT_MAX_THEORIQUE"]) - deja_periode, 2)
    ck.cmp_cap("T11", "OV ≤ part DOM de la période (CTS × prorata − déjà transféré)", "REG-04 / REG-05",
               montant, cap_plan, tol, "Période non rattachée au Planning", "BLOQUANT")
    if cap_plan is not None:
        caps["DOM période"] = cap_plan

    nom_ref = ref_nom or _g(d, "ANNEXE_I.nom_prenom") or _g(d, "BULLETIN.nom_prenom")
    compte_ref = c_dom or c_ov or c_a2
    # ATF et autorisations sont données pour le MOIS ENTIER (non réparties P1 / P2). Pour une partie :
    # - si le montant ATF / autorisé est la part mensuelle du contrat (± 2 %), le plafond de la partie est ce montant
    #   × jours de la partie / jours du mois (coefficient du Planning TL) ;
    # - sinon, le montant du mois moins ce qui est déjà transféré / validé sur l'autre partie du même mois.
    part_cts = parse_amount(prow.get("PART_TRANSFERABLE_CTS")) if prow is not None else None
    coef = None
    if prow is not None and part and not is_blank(prow.get("COEFFICIENT_PRORATA")):
        coef = float(prow["COEFFICIENT_PRORATA"])
    ddom, aut_lim = date_domiciliation(prow, dom_dates), None
    deja_mois = 0.0
    if part and mois and not reg_ok.empty and siege and {"SIEGE_RACINE", "PERIODE_TL"} <= set(reg_ok.columns):
        autres = reg_ok[(reg_ok["SIEGE_RACINE"].astype(str) == str(siege))
                        & reg_ok["PERIODE_TL"].astype(str).str.startswith(mois)
                        & (reg_ok["FICHIER"] != fichier)
                        & (reg_ok["PERIODE_TL"].astype(str) != f"{mois}{part}")]
        deja_mois = float(autres["_MONTANT"].sum())

    def _cap_mois(v, lib):
        """(plafond de la période, explication)."""
        if v is None:
            return None, ""
        if coef is not None and part_cts and abs(v - part_cts) <= 0.02 * part_cts:
            nbj, nbm = prow.get("NB_JOURS_SEGMENT"), prow.get("NB_JOURS_MOIS")
            return round(v * coef, 2), (f" | {lib} {v:,.2f} ≈ part mensuelle du contrat {part_cts:,.2f} (± 2 %) : "
                                        f"× {nbj}/{nbm} jours de la partie {part} (Planning TL)")
        return round(v - deja_mois, 2), (f" | {lib} du mois entier − {deja_mois:,.2f} déjà sur l'autre partie"
                                         if deja_mois else "")
    r_atf, how = _match_ref(atf, mois, compte_ref, nom_ref, s.seuil_similarite_nom, siege)
    if atf is None or atf.empty:
        ck.add("T12", "OV ≤ montant transférable ATF", "REG-05 / CTL-13", "INFO", "NON_CONTROLABLE", "Référentiel ATF non chargé")
    elif r_atf is None:
        ck.add("T12", "OV ≤ montant transférable ATF", "REG-05 / CTL-13", "BLOQUANT", "KO",
               f"Travailleur absent de l'ATF {mois or ''}")
    else:
        cap_atf, why = _cap_mois(parse_amount(r_atf.get("MONTANT_TRANSFERABLE")), "ATF")
        ck.cmp_cap("T12", "OV ≤ montant transférable ATF", "REG-05 / CTL-13", montant, cap_atf, tol,
                   "Montant ATF vide")
        ck.rows[-1]["DETAIL"] += f" | match {how}" + why
        if cap_atf is not None:
            caps["ATF"] = cap_atf

    r_aut, how = _match_ref(aut, mois, compte_ref, nom_ref, s.seuil_similarite_nom, siege)
    if aut is None or aut.empty:
        ck.add("T13", "Autorisation ministérielle : plafond et validité", "REG-05 / REG-07 / CTL-14", "INFO",
               "NON_CONTROLABLE", "Référentiel autorisations non chargé")
    elif r_aut is None:
        ck.add("T13", "Autorisation ministérielle : plafond et validité", "REG-05 / CTL-14", "BLOQUANT", "KO",
               f"Travailleur absent des autorisations {mois or ''}")
    else:
        cap_aut, why = _cap_mois(parse_amount(r_aut.get("MONTANT_AUTORISE")), "Autorisation")
        ck.cmp_cap("T13", "OV ≤ montant autorisé", "REG-05 / CTL-14", montant, cap_aut, tol, "Montant autorisé vide")
        ck.rows[-1]["DETAIL"] += f" | match {how}" + why
        if cap_aut is not None:
            caps["Autorisation"] = cap_aut
        lim = parse_date(r_aut.get("DATE_LIMITE"))
        dval = parse_date(r_aut.get("DATE_VALIDATION"))
        if lim is None and dval is not None:
            lim = dval + timedelta(days=s.autorisation_validite_jours)
        aut_lim = lim
        dref = parse_date(_g(d, "OV.date_demande")) or today
        if lim is None:
            ck.add("T13b", "Autorisation non échue", "REG-07 (point de départ à confirmer, Q-03)", "REVIEW",
                   "NON_CONTROLABLE", "Ni DATE_LIMITE ni DATE_VALIDATION")
        else:
            ck.add("T13b", "Autorisation non échue", "REG-07 (point de départ à confirmer, Q-03)", "BLOQUANT",
                   "OK" if dref <= lim else "KO", f"date réf. {dref:%d/%m/%Y} / limite {lim:%d/%m/%Y}")
        # T13c : l'autorisation doit courir au-delà de la date de la domiciliation (fichier CarthagoDom)
        if lim is None:
            pass
        elif ddom is None:
            ck.add("T13c", "Date limite d'autorisation > date de domiciliation", "REG-07", "INFO", "NON_CONTROLABLE",
                   "Date de domiciliation absente du fichier CarthagoDom" if prow is not None
                   else "Période non rattachée au Planning")
        else:
            ck.add("T13c", "Date limite d'autorisation > date de domiciliation", "REG-07", "BLOQUANT",
                   "OK" if lim > ddom else "KO", f"limite {lim:%d/%m/%Y} / domiciliation {ddom:%d/%m/%Y}")

    # ---------- T14 doublon / T15 cumul DOM
    if prow is not None:
        dup = reg_ok[(reg_ok["PERIODE_ID"] == prow["PERIODE_ID"]) & (reg_ok["FICHIER"] != fichier)] \
            if not reg_ok.empty else pd.DataFrame(columns=["FICHIER"])
        deja = ", ".join(dup["FICHIER"]) + f" ({deja_periode:,.2f})"
        if complement and not dup.empty:
            # complément d'une période déjà transférée : admis si le cumul reste dans la part de la période (T11)
            ck.add("T14", "Complément d'une période déjà transférée", "AS-06 / CTL-19", "REVIEW", "OK",
                   f"Complément ({complement}) de {deja} ; reste disponible sur la période "
                   f"{cap_plan:,.2f} — plafond contrôlé en T11")
        elif complement:
            ck.add("T14", "Complément d'une période déjà transférée", "AS-06 / CTL-19", "REVIEW", "KO",
                   f"Case 70 « {z70} » annonce un complément mais aucun TL exécuté ou validé sur la période")
        elif not dup.empty:
            ck.add("T14", "Pas de TL déjà validé sur la période", "AS-06 / CTL-19", "REVIEW", "KO",
                   f"Déjà transféré / validé : {deja} — si c'est un complément, l'indiquer dans la case 70")
        elif not reg_ok.empty:
            ck.add("T14", "Pas de TL déjà validé sur la période", "AS-06 / CTL-19", "BLOQUANT", "OK", "Aucun")
    plafond_dom = cumul_avant = solde_dom = None
    if prow is not None:
        plafond_dom = round(float(planning.loc[planning["DOM_ID"] == prow["DOM_ID"], "MONTANT_MAX_THEORIQUE"].sum()), 2)
        cumul_avant = float(reg_ok.loc[(reg_ok["DOM_ID"] == prow["DOM_ID"]) & (reg_ok["FICHIER"] != fichier),
                                       "_MONTANT"].sum()) if not reg_ok.empty else 0.0
        solde_dom = round(plafond_dom - cumul_avant - (montant or 0), 2)
        ck.add("T15", "Cumul des prélèvements ≤ montant domicilié", "REG-08 / CTL-19 (Q-02)", "BLOQUANT",
               "OK" if solde_dom >= -tol else "KO",
               f"plafond DOM {plafond_dom:,.2f} − cumul {cumul_avant:,.2f} − OV {montant or 0:,.2f} = {solde_dom:,.2f}")

    # ---------- synthèse
    # ---------- R01–R08 : base RH de l'employeur (référence des documents imprimés ; l'OV reste la base du TL)
    plan_mois = plan_dom[plan_dom["MOIS_BASE"] == mois] if not plan_dom.empty and mois else pd.DataFrame()
    r_rh = _controles_rh(ck, d, rh, siege, mois, part, montant, devise, nom_ref, tol, prow, plan_mois)
    _controles_beneficiaire(ck, d, r_rh)
    dv_ov = devise_transfert(_g(d, "OV.devise_transfert"))

    ko_b = [r for r in ck.rows if r["SEVERITE"] == "BLOQUANT" and r["RESULTAT"] == "KO"]
    ko_r = [r for r in ck.rows if r["SEVERITE"] == "REVIEW" and r["RESULTAT"] in ("KO", "NON_CONTROLABLE")]
    statut = "NON_CONFORME" if ko_b else ("A_REVOIR" if ko_r else "CONFORME")
    plafond = min(caps.values()) if caps else None
    synth = {
        "FICHIER": fichier, "AUDIT_SOURCE": d.get("AUDIT_SOURCE"), "STATUT_CONTROLE": statut,
        "NB_BLOQUANT": len(ko_b), "NB_REVIEW": len(ko_r),
        "MOTIFS": " | ".join(f"{r['CODE']} {r['CONTROLE']}" for r in ko_b + ko_r),
        "NOM": nom_ref, "SIEGE_RACINE": siege, "COMPTE_LOCAL": compte_ref or None,
        "PERIODE_TL": prow["PERIODE_TL"] if prow is not None else periode_label(mois, part),
        "PERIODE_ID": prow["PERIODE_ID"] if prow is not None else None,
        "DOM_ID": prow["DOM_ID"] if prow is not None else None,
        "NUMERO_DOMICILIATION": prow["NUMERO_DOMICILIATION"] if prow is not None else _g(d, "ANNEXE_II.numero_domiciliation"),
        "REFERENCE_PAIEMENT": z70, "COMPLEMENT": complement, "LECTURE_PERIODE": lu["regle"] or lu["statut"],
        "MONTANT_OV": montant, "DEVISE": devise, "DEVISE_TRANSFERT": dv_ov,
        # Planning TL : la part de la période reste celle du contrat (inchangée) ; ce qui reste après les TL déjà
        # transférés / validés est une donnée à part (plafond du contrôle T11)
        "PLANNING_PART_PERIODE": (round(float(prow["MONTANT_MAX_THEORIQUE"]), 2) if prow is not None else None),
        "PLANNING_PART_RESTANTE": cap_plan,
        "DATE_DOMICILIATION": ddom.isoformat() if ddom else None,
        "AUT_DATE_LIMITE": aut_lim.isoformat() if aut_lim else None,
        "RH_MATRICULE": (r_rh or {}).get("MATRICULE"), "RH_PART_TRANSFERABLE": (r_rh or {}).get("_PART_ATTENDUE", (r_rh or {}).get("PART_TRANSFERABLE")),
        "RH_ADRESSE_BENEF": (r_rh or {}).get("ADRESSE_BENEF"), "RH_DEVISE": (r_rh or {}).get("DEVISE"),
        **{f"PLAFOND_{k.upper().replace(' ', '_').replace('%', 'PCT')}": v for k, v in caps.items()},
        "PLAFOND_RETENU": plafond, "PLAFOND_LIMITANT": min(caps, key=caps.get) if caps else None,
        "MARGE": round(plafond - montant, 2) if (plafond is not None and montant is not None) else None,
        "PROVISION_ATLAS": round(montant + s.provision_majoration_dzd, 2) if montant is not None else None,
        "PLAFOND_DOM": plafond_dom, "CUMUL_DOM_AVANT": cumul_avant, "SOLDE_DOM_APRES": solde_dom,
        "BENEFICIAIRE": _g(d, "OV.beneficiaire_nom"), "BENEFICIAIRE_COMPTE": _g(d, "OV.beneficiaire_compte"),
        "SWIFT": _g(d, "OV.swift"), "BANQUE": _g(d, "OV.banque_nom"), "DATE_DEMANDE": _g(d, "OV.date_demande"),
    }
    return synth, ck.rows


def _controles_rh(ck, d, rh, client, mois, part, montant, devise, nom, tol, prow=None, plan_mois=None) -> dict | None:
    from . import rh as _rh
    if rh is None or rh.empty:
        ck.add("R01", "Salarié présent dans la base RH du mois", "Base RH employeur", "INFO", "NON_CONTROLABLE",
               "Base RH non chargée")
        return None
    r, how = _rh.match(rh, client, mois, part, nom)
    if r is None:
        ck.add("R01", "Salarié présent dans la base RH du mois", "Base RH employeur", "REVIEW", "KO",
               f"Absent de la base RH {mois or '?'}" + (f" ({how})" if how else ""))
        return None
    ck.add("R01", "Salarié présent dans la base RH du mois", "Base RH employeur", "REVIEW", "OK",
           f"matricule {r.get('MATRICULE')} · {r.get('NOM_PRENOM')} · {r.get('PERIODE')} (rapprochement {how})")
    prt = r.get("PART_TRANSFERABLE")
    note = ""
    if prt is not None and part and r.get("RH_MENSUELLE"):
        # base RH « mois entier » (bases passées) : part attendue = montant du mois × quote-part de la période
        tot = float(sum(parse_amount(x) or 0 for x in plan_mois["MONTANT_MAX_THEORIQUE"])) \
            if plan_mois is not None and not plan_mois.empty else 0.0
        mt = parse_amount(prow["MONTANT_MAX_THEORIQUE"]) if prow is not None else None
        if tot and mt is not None:
            prt = round(prt * mt / tot, 2)
            note = f" (base RH du mois entier, quote-part {part} = {mt / tot:.0%})"
        else:
            prt, note = None, " (base RH du mois entier : quote-part de la période inconnue)"
    r = dict(r, _PART_ATTENDUE=prt, _NOTE_RH=note.strip(" ()") or None)
    if montant is None or prt is None:
        ck.add("R02", "Montant OV = part transférable RH", "Base RH employeur", "REVIEW", "NON_CONTROLABLE",
               "Montant OV ou part RH absent")
    elif abs(montant - prt) <= max(tol, 1.0 if note else tol):
        ck.add("R02", "Montant OV = part transférable RH", "Base RH employeur", "REVIEW", "OK", f"{montant:,.2f}{note}")
    else:
        ck.add("R02", "Montant OV = part transférable RH", "Base RH employeur",
               "BLOQUANT" if montant > prt + tol and not note else "REVIEW", "KO",
               f"OV {montant:,.2f} / RH {prt:,.2f} (écart {montant - prt:,.2f}){note}")
    # R03 (IBAN), R04 (SWIFT), R09 (banque) : _controles_beneficiaire, aussi sans base RH (OV / Annexe II)
    raw = _g(d, "OV.devise_transfert")
    dv, dv_rh = devise_transfert(raw), devise_transfert(r.get("DEVISE"))
    ck.add("R05", "Devise de transfert OV = base RH", "Base RH employeur", "INFO" if not dv or not dv_rh else "REVIEW",
           "NON_CONTROLABLE" if not dv or not dv_rh else ("OK" if dv == dv_rh else "KO"),
           f"OV {dv or '?'}" + (f" (« {raw} »)" if raw and str(raw).strip().upper() != dv else "")
           + f" / RH {dv_rh or '?'}")
    sim = name_similarity(_g(d, "OV.beneficiaire_nom"), r.get("BENEFICIAIRE"))
    ck.add("R06", "Bénéficiaire OV = base RH", "Base RH employeur", "REVIEW",
           "NON_CONTROLABLE" if sim is None else ("OK" if sim >= 0.8 else "KO"),
           f"OV « {_g(d, 'OV.beneficiaire_nom') or '?'} » / RH « {r.get('BENEFICIAIRE') or '?'} »"
           + ("" if sim is None else f" ({sim:.0%})")
           + (" — bénéficiaire ≠ salarié" if r.get("BENEFICIAIRE") and (name_similarity(r.get("BENEFICIAIRE"), r.get("NOM_PRENOM")) or 0) < 0.8 else ""))
    fin = r.get("PERMIS_FIN")
    if fin and mois:
        y, m = map(int, mois.split("-"))
        deb = date(y, m, 16 if part == "P2" else 1)
        ck.add("R07", "Permis de travail (RH) valide sur la période", "Base RH employeur", "REVIEW",
               "OK" if fin >= deb else "KO", f"fin de validité {fin:%d/%m/%Y}")
    # R08 (nombre de jours Annexe II = base RH) : retiré — non comparé (décision métier du 24/10/2026)
    return r


def _controles_beneficiaire(ck, d, r: dict | None) -> None:
    """Coordonnées du bénéficiaire : IBAN (OV / Annexe II / base RH, comparés sans espaces), SWIFT (OV / base RH),
    nom de la banque (OV / Annexe II / base RH)."""
    from . import rh as _rh
    r = r or {}
    ib = {"OV": _rh.iban_clean(_g(d, "OV.beneficiaire_compte")),
          "Annexe II": _rh.iban_clean(_g(d, "ANNEXE_II.compte_devise")),
          "Base RH": _rh.iban_clean(r.get("IBAN"))}
    vus = {k: v for k, v in ib.items() if v}
    ck.add("R03", "IBAN du bénéficiaire identique OV / Annexe II / base RH (sans espaces)", "Base RH employeur",
           "INFO" if len(vus) < 2 else "REVIEW", "NON_CONTROLABLE" if len(vus) < 2 else ("OK" if len(set(vus.values())) == 1 else "KO"),
           " | ".join(f"{k} {v or '?'}" for k, v in ib.items()))
    sw_ov, sw_rh = _rh.swift11(_g(d, "OV.swift") or "")[:8], _rh.swift11(r.get("SWIFT") or "")[:8]
    ck.add("R04", "Code SWIFT OV = base RH", "Base RH employeur", "INFO" if not sw_ov or not sw_rh else "REVIEW",
           "NON_CONTROLABLE" if not sw_ov or not sw_rh else ("OK" if sw_ov == sw_rh else "KO"),
           f"OV {sw_ov or '?'} / RH {sw_rh or '?'}")
    bq = {"OV": _g(d, "OV.banque_nom"), "Annexe II": _g(d, "ANNEXE_II.banque_etranger"), "Base RH": r.get("BANQUE_ETR")}
    vus = {k: v for k, v in bq.items() if not is_blank(v)}
    ref = next(iter(vus.values()), None)
    sims = {k: banque_similarite(ref, v) for k, v in vus.items()}
    faibles = [f"{k} ({v:.0%})" for k, v in sims.items() if v is not None and v < 0.8]
    ck.add("R09", "Banque du bénéficiaire identique OV / Annexe II / base RH", "Base RH employeur",
           "INFO" if len(vus) < 2 else "REVIEW",
           "NON_CONTROLABLE" if len(vus) < 2 else ("KO" if faibles else "OK"),
           " | ".join(f"{k} « {v or '?'} »" for k, v in bq.items()) + (" | écarts : " + ", ".join(faibles) if faibles else ""))


def _jours(v) -> int | None:
    """« 15 », « 15 jours », 15.0 -> 15 ; texte sans nombre -> None."""
    if v is None or (isinstance(v, float) and v != v):
        return None
    m = re.search(r"\d{1,2}", str(v))
    return int(m.group()) if m else None


def run_controls(audit: pd.DataFrame, planning: pd.DataFrame, registre: pd.DataFrame,
                 atf: pd.DataFrame, aut: pd.DataFrame, s, today: date | None = None, rh: pd.DataFrame | None = None,
                 dom_dates: dict | None = None):
    sy, ch = control_each(audit, planning, registre, atf, aut, s, today, rh, dom_dates)
    return mark_lot_duplicates(sy, ch)


def control_each(audit: pd.DataFrame, planning: pd.DataFrame, registre: pd.DataFrame,
                 atf: pd.DataFrame, aut: pd.DataFrame, s, today: date | None = None, rh: pd.DataFrame | None = None,
                 dom_dates: dict | None = None):
    """Contrôles dossier par dossier (sans les doublons internes au lot). dom_dates : n° DOM -> date de
    domiciliation (fichier CarthagoDom, voir dates_domiciliation)."""
    reg_ok = _registre_valide(registre)
    synth, checks = [], []
    for _, r in audit.iterrows():
        a, b = control_dossier(r.to_dict(), planning, reg_ok, atf, aut, s, today, rh, dom_dates)
        synth.append(a)
        checks.extend(b)
    return pd.DataFrame(synth), pd.DataFrame(checks)


def dates_domiciliation(path) -> dict:
    """{n° DOM compacté (chiffres + DZD) : date} depuis l'export CarthagoDom (colonne « Date domiciliation »)."""
    from .boci import carthago_dom_dates
    df = carthago_dom_dates(path)
    return dict(zip(df["_DOM"], df["_DATE"])) if not df.empty else {}


def mark_lot_duplicates(sy: pd.DataFrame, ch: pd.DataFrame):
    """T14b : deux dossiers du même lot sur la même période (vectorisé : lots de ~1000 dossiers)."""
    sy, ch = sy.copy(), ch.copy()
    if sy.empty:
        return sy, ch
    m = sy["PERIODE_ID"].notna() & sy.duplicated("PERIODE_ID", keep=False)
    if not m.any():
        return sy, ch
    groups = sy[m].groupby("PERIODE_ID")["FICHIER"].apply(list).to_dict()
    new = []
    for f, pid in zip(sy.loc[m, "FICHIER"], sy.loc[m, "PERIODE_ID"]):
        autres = [x for x in groups[pid] if x != f]
        detail = "Même période que " + ", ".join(autres[:5]) + (f" (+{len(autres) - 5})" if len(autres) > 5 else "")
        new.append({"FICHIER": f, "CODE": "T14b", "CONTROLE": "Doublon dans le lot", "REFERENCE": "CTL-19",
                    "SEVERITE": "BLOQUANT", "RESULTAT": "KO", "DETAIL": detail})
    ch = pd.concat([ch, pd.DataFrame(new)], ignore_index=True)
    sy.loc[m, "STATUT_CONTROLE"] = "NON_CONFORME"
    sy.loc[m, "NB_BLOQUANT"] = sy.loc[m, "NB_BLOQUANT"] + 1
    sy.loc[m, "MOTIFS"] = sy.loc[m, "MOTIFS"].map(lambda x: (x + " | " if x else "") + "T14b Doublon dans le lot")
    return sy, ch


# ------------------------------------------------------------------ décisions / registre
DECISION_VALUES = ["", "VALIDE", "REJETE", "EN_ATTENTE"]


def load_decisions(path: Path) -> pd.DataFrame:
    cols = ["FICHIER", "AUDIT_SOURCE", "DECISION", "COMMENTAIRE", "USER", "TS"]
    if not path.exists():
        return pd.DataFrame(columns=cols)
    df = pd.read_csv(path, dtype=str, encoding="utf-8-sig").reindex(columns=cols)
    return df.drop_duplicates("FICHIER", keep="last")


def save_decisions(path: Path, edited: pd.DataFrame, user: str) -> int:
    old = load_decisions(path)
    old_map = {r.FICHIER: (r.DECISION or "", r.COMMENTAIRE or "") for r in old.fillna("").itertuples()}
    new_rows = []
    for _, r in edited.iterrows():
        dec, com = (r.get("DECISION") or "").strip(), (r.get("COMMENTAIRE") or "").strip()
        if old_map.get(r["FICHIER"], ("", "")) != (dec, com) and (dec or com or r["FICHIER"] in old_map):
            new_rows.append({"FICHIER": r["FICHIER"], "AUDIT_SOURCE": r.get("AUDIT_SOURCE"), "DECISION": dec,
                             "COMMENTAIRE": com, "USER": user, "TS": datetime.now().isoformat(timespec="seconds")})
    if new_rows:
        path.parent.mkdir(parents=True, exist_ok=True)
        hist = pd.read_csv(path, dtype=str, encoding="utf-8-sig") if path.exists() else pd.DataFrame()
        pd.concat([hist, pd.DataFrame(new_rows)], ignore_index=True).to_csv(path, index=False, encoding="utf-8-sig")
    return len(new_rows)


def check_decisions(edited: pd.DataFrame) -> list[str]:
    errs = []
    for _, r in edited.iterrows():
        dec = (r.get("DECISION") or "").strip()
        if dec == "VALIDE" and r.get("STATUT_CONTROLE") == "NON_CONFORME" and not (r.get("COMMENTAIRE") or "").strip():
            errs.append(f"{r['FICHIER']} : validation d'un dossier NON_CONFORME sans justification")
        if dec == "VALIDE" and not r.get("PERIODE_ID"):
            errs.append(f"{r['FICHIER']} : impossible de valider un TL non rattaché au Planning")
        if dec == "REJETE" and not (r.get("COMMENTAIRE") or "").strip():
            errs.append(f"{r['FICHIER']} : rejet sans motif")
    return errs


def register(path: Path, synth: pd.DataFrame, decisions: pd.DataFrame, user: str,
             vigilance: dict | None = None) -> int:
    """Ajoute au registre les TL validés (idempotent par TL_ID).
    vigilance = {fichier: statut} : si fourni, seuls les TL au statut RAS ou HIT_LEVE sont inscrits."""
    reg = load_registre(path)
    known = set(reg["TL_ID"].dropna())
    dec = decisions.set_index("FICHIER") if not decisions.empty else pd.DataFrame()
    rows = []
    for _, r in synth.iterrows():
        if r["FICHIER"] not in dec.index or dec.loc[r["FICHIER"], "DECISION"] != "VALIDE":
            continue
        if vigilance is not None and vigilance.get(r["FICHIER"]) not in ("RAS", "HIT_LEVE"):
            continue
        tid = tl_id(r["FICHIER"], r.get("PERIODE_ID"))
        if tid in known:
            continue
        rows.append({"TL_ID": tid, "FICHIER": r["FICHIER"], "DOM_ID": r.get("DOM_ID"),
                     "PERIODE_ID": r.get("PERIODE_ID"), "PERIODE_TL": r.get("PERIODE_TL"),
                     "SIEGE_RACINE": r.get("SIEGE_RACINE"), "NOM": r.get("NOM"),
                     "NUMERO_DOMICILIATION": r.get("NUMERO_DOMICILIATION"), "MONTANT": r.get("MONTANT_OV"),
                     "DEVISE": r.get("DEVISE"), "PROVISION_ATLAS": r.get("PROVISION_ATLAS"),
                     "STATUT_CONTROLE": r.get("STATUT_CONTROLE"), "DECISION": "VALIDE",
                     "COMMENTAIRE": dec.loc[r["FICHIER"], "COMMENTAIRE"], "VALIDE_PAR": user,
                     "DATE_VALIDATION": datetime.now().isoformat(timespec="seconds"),
                     "STATUT_EXECUTION": "A_EXECUTER",
                     "STATUT_VIGILANCE": (vigilance or {}).get(r["FICHIER"])})
    if rows:
        path.parent.mkdir(parents=True, exist_ok=True)
        pd.concat([reg, pd.DataFrame(rows)], ignore_index=True).reindex(columns=REGISTRE_COLUMNS) \
            .to_csv(path, index=False, encoding="utf-8-sig")
    return len(rows)


def consumption(planning: pd.DataFrame, registre: pd.DataFrame, today: date | None = None,
                tol: float = 1.0) -> pd.DataFrame:
    """Planning enrichi : montant transféré, solde, statut par période (base AS-07)."""
    today = today or date.today()
    if planning is None or planning.empty:
        return pd.DataFrame()
    p = planning.copy()
    reg = _registre_valide(registre)
    sums = reg.groupby("PERIODE_ID")["_MONTANT"].sum() if not reg.empty else pd.Series(dtype=float)
    last = (reg.dropna(subset=["DATE_VALIDATION"]).sort_values("DATE_VALIDATION").drop_duplicates("PERIODE_ID", keep="last")
            .set_index("PERIODE_ID")["DATE_VALIDATION"] if not reg.empty else pd.Series(dtype=str))
    p["MONTANT_TRANSFERE"] = p["PERIODE_ID"].map(sums).fillna(0.0).round(2)
    p["SOLDE_RESTANT"] = (p["MONTANT_MAX_THEORIQUE"].astype(float) - p["MONTANT_TRANSFERE"]).round(2)
    p["DATE_EXECUTION_TL"] = p["PERIODE_ID"].map(last)

    def statut(r):
        if r["MONTANT_TRANSFERE"] > 0 and r["SOLDE_RESTANT"] <= tol:
            return "TRANSFERE"
        if r["MONTANT_TRANSFERE"] > 0:
            return "PARTIEL"
        fin = parse_date(r["DATE_FIN_SEGMENT"])
        return "NON_TRANSFERE" if fin and fin < today else "A_VENIR"

    p["STATUT_PERIODE"] = p.apply(statut, axis=1)
    return p
