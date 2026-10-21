"""Employeur de démonstration : jeu de données entièrement FICTIF, cohérent de bout en bout.

Sert la démonstration en direct (aucun GPU nécessaire) : dossiers DOM lus, Planning TL, base RH / ATF /
autorisations du mois, lot d'OV lus et contrôlés (avec des anomalies voulues), vigilance (un HIT à analyser),
historique CarthagoTrade / Atlas / blocages / soldes, puis TL prêts pour les fichiers BOCI.
Noms, comptes, IBAN et montants sont inventés (clés RIB / IBAN calculées pour être valides).

    from tlone import demo; demo.build(home)          # crée / recrée employeurs/DEMO
"""
from __future__ import annotations

import calendar
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd

from . import employeurs
from .config import Paths

CODE = "DEMO"
RAISON = "SIDERURGIE DEMO SPA (DONNEES FICTIVES)"
NIF = "000016099999999"
ADRESSE_CLIENT = "BASE DE VIE SITE DEMO 31015 ORAN"
BANQUES_TR = {"YAPITRIS": ("00067", "YAPI VE KREDI BANKASI"), "TEBUTRIS": ("00032", "TEB/TURK EKONOMI BANK"),
              "TCZBTR2A": ("00010", "TURKIYE ZIRAAT BANKASI")}
BANQUES_IN = {"HDFCINBBXXX": "HDFC BANK LTD", "PUNBINBBBFB": "PUNJAB NATIONAL BANK"}

# (nom, prénom, nationalité, BIC, salaire net, mois avant M du début du contrat, jour de début, renouvellement,
#  bénéficiaire tiers)
SALARIES = [
    ("AKSOY", "MURAT", "TR", "YAPITRIS", 612400.00, 9, 1, None, None),
    ("DEMIRCI", "KEMAL", "TR", "TEBUTRIS", 548900.00, 8, 1, None, None),
    ("ERDEM", "SELIM", "TR", "YAPITRIS", 701250.00, 10, 1, None, None),
    ("GUNES", "ONUR", "TR", "TCZBTR2A", 455300.00, 7, 1, None, None),
    ("KAPLAN", "EMRE", "TR", "TEBUTRIS", 389600.00, 9, 1, None, None),
    ("OZTURK", "BARIS", "TR", "YAPITRIS", 520000.00, 8, 1, None, None),
    ("POLAT", "CEM", "TR", "TEBUTRIS", 466750.00, 6, 1, None, None),
    ("PATEL", "RAVI KUMAR", "IN", "HDFCINBBXXX", 498300.00, 7, 1, None, None),
    ("SAHIN", "VOLKAN", "TR", "YAPITRIS", 575800.00, 9, 1, None, "SAHIN AYSE"),
    ("SHARMA", "ARJUN", "IN", "PUNBINBBBFB", 432900.00, 10, 1, None, None),
    ("TASKIN", "HAKAN", "TR", "TEBUTRIS", 633100.00, 11, 18, "DOMICILIEE", None),
    ("YAVUZ", "ORHAN", "TR", "YAPITRIS", 541200.00, 11, 18, "NON_DOMICILIEE", None),
]
UNITES = ["", "UN", "DEUX", "TROIS", "QUATRE", "CINQ", "SIX", "SEPT", "HUIT", "NEUF", "DIX", "ONZE", "DOUZE", "TREIZE",
          "QUATORZE", "QUINZE", "SEIZE"]
DIZAINES = {2: "VINGT", 3: "TRENTE", 4: "QUARANTE", 5: "CINQUANTE", 6: "SOIXANTE"}


# ----------------------------------------------------------------------------- utilitaires
def _cent(n: int) -> str:
    if n < 17:
        return UNITES[n]
    if n < 20:
        return "DIX " + UNITES[n - 10]
    d, u = divmod(n, 10)
    if d in (7, 9):
        base = "SOIXANTE" if d == 7 else "QUATRE VINGT"
        reste = 10 + u
        return base + (" ET " if reste == 11 and d == 7 else " ") + _cent(reste)
    base = "QUATRE VINGT" if d == 8 else DIZAINES[d]
    if u == 0:
        return base
    return base + (" ET " if u == 1 and d != 8 else " ") + UNITES[u]


def _mille(n: int) -> str:
    c, r = divmod(n, 100)
    out = ("CENT" if c == 1 else f"{UNITES[c]} CENT") if c else ""
    if r:
        out = (out + " " if out else "") + _cent(r)
    return out


def en_lettres(x: float) -> str:
    """Montant en lettres (format des OV) : « TROIS CENT DOUZE MILLE … DINARS ET 25 CTS »."""
    n = int(round(x * 100))
    e, c = divmod(n, 100)
    m, r = divmod(e, 1000)
    mil, m = divmod(m, 1000)
    parts = []
    if mil:
        parts.append(("UN MILLION" if mil == 1 else f"{_mille(mil)} MILLIONS"))
    if m:
        parts.append("MILLE" if m == 1 else f"{_mille(m)} MILLE")
    if r:
        parts.append(_mille(r))
    s = " ".join(parts) + " DINARS ALGERIENS"
    return s + (f" ET {c} CTS" if c else "")


def _rib(client: str) -> str:
    b = "0270" + client + "001"
    cle = 97 - (int(b[3:]) * 100) % 97
    return b + f"{cle:02d}"


def _iban_tr(bank5: str, acct16: str) -> str:
    bban = bank5 + "0" + acct16
    num = "".join(str(int(ch, 36)) for ch in bban + "TR00")
    return f"TR{98 - int(num) % 97:02d}{bban}"


def _add_months(d: date, n: int) -> date:
    y, m = divmod(d.month - 1 + n, 12)
    y += d.year
    m += 1
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


def demo_month(today: date) -> date:
    """Dernier mois échu de 30 jours (prorata RH « /30 » = prorata calendaire du Planning)."""
    d = date(today.year, today.month, 1)
    for _ in range(12):
        d = _add_months(d, -1)
        if calendar.monthrange(d.year, d.month)[1] == 30:
            return d
    return _add_months(date(today.year, today.month, 1), -1)


MOIS_FR = ["", "JANVIER", "FEVRIER", "MARS", "AVRIL", "MAI", "JUIN", "JUILLET", "AOUT", "SEPTEMBRE", "OCTOBRE",
           "NOVEMBRE", "DECEMBRE"]


def _workers(M: date) -> list[dict]:
    out = []
    for i, (nom, pre, nat, bic, net, mois_avant, jour, renouv, tiers) in enumerate(SALARIES, start=1):
        client = f"07310{990100 + i:06d}"
        rib = _rib(client)
        if nat == "TR":
            bank5, bname = BANQUES_TR[bic]
            iban = _iban_tr(bank5, f"{9000000000000000 + i * 7919:016d}")
            swift, devise, pays = bic, "USD", "TURQUIE"
        else:
            bname = BANQUES_IN[bic]
            iban = f"{50100000000000 + i * 104729}"
            swift, devise, pays = bic, "EUR", "INDE"
        start = _add_months(M, -mois_avant).replace(day=jour)
        if renouv:
            end = M.replace(day=17)
            new_start, new_end = M.replace(day=18), _add_months(M.replace(day=17), 24)
        else:
            end = _add_months(start, 24) - timedelta(days=1)
            new_start = new_end = None
        out.append(dict(i=i, nom=nom, prenom=pre, nom_prenom=f"{nom} {pre}", nat=nat, client=client, rib=rib,
                        racine=client[5:], iban=iban, swift=swift, banque=bname, devise=devise, pays=pays, net=net,
                        part=round(0.95 * net, 2), start=start, end=end, renouv=renouv, new_start=new_start,
                        new_end=new_end, benef=tiers or f"{nom} {pre}", matricule=str(2000 + i),
                        pere=f"{['ALI', 'MEHMET', 'HASAN', 'IBRAHIM', 'OSMAN', 'MUSTAFA'][i % 6]}",
                        mere=f"{nom} {['FATMA', 'AYSE', 'EMINE', 'HATICE', 'ZEYNEP'][i % 5]}",
                        permis=f"2{i % 5}-000{10400 + i * 37} / 31-2{4 + i % 2}-00{1500 + i}",
                        permis_new=f"2{i % 5}-000{10400 + i * 37} / 31-26-00{2600 + i}",
                        dom=f"27190120{start.year % 100:02d}34{i:06d}DZD",
                        dom_new=f"27190120{(new_start.year % 100) if new_start else 0:02d}35{i:06d}DZD",
                        predom=f"{start.year}|07310|{i:06d}", date_dom=start - timedelta(days=12),
                        adresse=f"{['CUMHURIYET', 'ATATURK', 'INONU', 'FATIH', 'YILDIZ', 'BARBAROS'][i % 6]} MAH. "
                                f"{100 + i * 3} SK. NO:{i} DAIRE {i + 2} "
                                f"{['MERKEZ/HATAY', 'DORTYOL/HATAY', 'KORFEZ/KOCAELI', 'MENEMEN/IZMIR'][i % 4]}"
                        if nat == "TR" else f"H NO {i}-{40 + i} MG ROAD WARD {i} PIN:5600{10 + i:02d} KARNATAKA INDE"))
    return out


# ----------------------------------------------------------------------------- construction
def build(home: Path, user: str = "demo", today: date | None = None, reset: bool = True) -> Path:
    home = Path(home)
    today = today or date.today()
    root = employeurs.employeurs_dir(home) / CODE
    if reset and root.exists():
        shutil.rmtree(root)
    employeurs.create(home, RAISON, NIF, user, mots_cles="DEMO", agence="07310", demo=True, code=CODE)
    P = Paths(root)
    P.ensure()
    M = demo_month(today)
    W = _workers(M)
    stamp = (datetime.combine(today, datetime.min.time()) - timedelta(hours=6))
    _dom(P, W, M)
    PL = _planning(P)
    _referentiels(P, W, M, today)
    _historique(P, W, M, PL, today, stamp)
    _ov(P, W, M, today)
    _decisions_vigilance(P, home, user, today)
    return root


def _dom(P: Paths, W: list[dict], M: date) -> None:
    """Classeur de la Partie 2 (dossiers DOM lus et normalisés), détail des champs, export CarthagoDom, PDF."""
    dossiers, champs, carthago = [], [], []
    for w in W:
        versions = [("CNT", w["start"], w["end"], w["dom"])]
        if w["renouv"] == "DOMICILIEE":
            versions.append(("CNT2", w["new_start"], w["new_end"], w["dom_new"]))
        for suffix, deb, fin, dom in versions:
            f = f"{w['nom']}_{w['prenom'].split()[0]}_{suffix}.pdf"
            d = dict(FICHIER=f, TYPE_DOSSIER="NOUVEAU_CONTRAT", DOSSIER_STATUS="REVIEW",
                     CTS_DATE_DEBUT_VALIDITE_PERMIS=f"{deb:%d/%m/%Y}", CTS_DATE_FIN_VALIDITE_PERMIS=f"{fin:%d/%m/%Y}",
                     CTS_SALAIRE_NET=f"{w['net']:.2f}", CTS_SALAIRE_NET_ANCIEN=None,
                     CTS_PART_TRANSFERABLE=f"{w['part']:.2f}", CTS_PART_PAYABLE_DZD=f"{w['net'] - w['part']:.2f}",
                     CTS_NOM_PRENOM_TRAVAILLEUR=w["nom_prenom"], CTS_NUMERO_PERMIS_TRAVAIL=(w["permis"] if suffix == "CNT" else w["permis_new"]),
                     CTS_REFERENCE_DOCUMENT=f"CT-{deb.year}-{w['i']:04d}", CTS_DATE_DOCUMENT=f"{deb - timedelta(days=20):%d/%m/%Y}",
                     DOM_COMPTE_LOCAL_NORMALIZED=w["rib"], SIEGE_RACINE=w["client"], NUMERO_DOMICILIATION_PROPOSE=dom,
                     DQ_SCORE=92 - (w["i"] % 4) * 6, AUGMENTATION_YEAR=None,
                     CARTHAGO_MATCH_STATUS="EXACT_DOMICILIATION_REFERENCE_MATCH", MODULO97_OK=True,
                     CTS_NOM_PERE=w["pere"], CTS_NOM_MERE=w["mere"])
            dossiers.append(d)
            champs += _champs(d, w, deb, fin, dom)
            carthago.append({"Agence": "AGENCE MOSTAGANEM", "Client": w["client"], "Devise": "DZD",
                             "Nom complet": w["nom_prenom"], "Name": w["nom_prenom"], "Numéro de domiciliation": dom,
                             "Date domiciliation": f"{deb - timedelta(days=12):%d/%m/%Y}"})
            _pdf_dom(P.dom_in / f, d, w)
    with pd.ExcelWriter(P.validation_xlsx, engine="openpyxl") as xw:
        pd.DataFrame([{"FICHIER": d["FICHIER"], "TYPE_DOSSIER": d["TYPE_DOSSIER"], "DQ_SCORE": d["DQ_SCORE"]}
                      for d in dossiers]).to_excel(xw, sheet_name="DQ_DASHBOARD", index=False)
        pd.DataFrame(columns=["FICHIER", "TYPE_DOSSIER", "PAGE", "TYPE_DOCUMENT", "CHAMP", "RAW", "NORMALIZED",
                              "CONFIDENCE_QUALITE", "NIVEAU_CONFIDENCE", "VALEUR_PROPOSEE", "DECISION_MANUELLE",
                              "CONTEXTE_VALIDATION", "SOURCE_PROPOSITION", "VALEUR_VALIDEE", "COMMENTAIRE",
                              "DOSSIER_STATUS"]).to_excel(xw, sheet_name="A_VALIDER_MANUELLEMENT", index=False)
        pd.DataFrame(dossiers).to_excel(xw, sheet_name="DOSSIERS_A_VALIDER", index=False)
        pd.DataFrame([{"FICHIER": d["FICHIER"], "CONTROLE": "IDENTITE_CTS_CTR", "STATUT": "OK"} for d in dossiers]) \
            .to_excel(xw, sheet_name="CONTROLES_COHERENCE", index=False)
    pd.DataFrame(champs).to_csv(P.fields_csv, index=False, encoding="utf-8-sig")
    pd.DataFrame(carthago).to_excel(P.carthago_xlsx, index=False)
    pd.DataFrame([{"FICHIER": d["FICHIER"], "SIEGE_RACINE": d["SIEGE_RACINE"], "CARTHAGO_CLIENT": d["SIEGE_RACINE"],
                   "CARTHAGO_NOM_COMPLET": d["CTS_NOM_PRENOM_TRAVAILLEUR"],
                   "CARTHAGO_NUM_DOM_RAW": d["NUMERO_DOMICILIATION_PROPOSE"], "REFERENCE_EXACT_MATCH": True,
                   "CARTHAGO_TYPE": "DOM"} for d in dossiers]).to_csv(P.carthago_match_csv, index=False, encoding="utf-8-sig")
    pd.DataFrame({"ANNEE_AUGMENTATION": [M.year], "MOIS_AUGMENTATION": [4]}).to_excel(P.param_aug_xlsx, index=False)


def _champs(d: dict, w: dict, deb: date, fin: date, dom: str) -> list[dict]:
    from .matrice import months_between
    mois = months_between(deb, fin)
    vals = {
        "DOM_NOM_RAISON_SOCIAL_CLIENT": w["nom_prenom"], "DOM_COMPTE_LOCAL": w["rib"], "DOM_NUMERO_CONTRAT": d["CTS_REFERENCE_DOCUMENT"],
        "DOM_AGENCE_DOMICILIATAIRE": "AGENCE MOSTAGANEM", "DOM_DUREE_CONTRAT_MOIS": str(mois),
        "DOM_DATE_DEBUT_CONTRAT": f"{deb:%d/%m/%Y}", "DOM_DATE_FIN_CONTRAT": f"{fin:%d/%m/%Y}",
        "DOM_SALAIRE_NET_MENSUEL": f"{w['net']:.2f}", "DOM_PART_TRANSFERABLE": f"{w['part']:.2f}", "DOM_TAUX_TRANSFERABLE": "95",
        "DOM_MONTANT_TOTAL_DOMICILIE": f"{w['part'] * mois:.2f}", "DOM_DATE_SIGNATURE": f"{deb - timedelta(days=15):%d/%m/%Y}",
        "CTR_REFERENCE_DOCUMENT": d["CTS_REFERENCE_DOCUMENT"], "CTR_DUREE_MOIS": str(mois),
        "CTR_DATE_DEBUT_CONTRAT": f"{deb:%d/%m/%Y}", "CTR_NOM_PRENOM_TRAVAILLEUR": w["nom_prenom"],
        "CTR_NUMERO_PERMIS_TRAVAIL": d["CTS_NUMERO_PERMIS_TRAVAIL"], "CTR_DATE_DEBUT_VALIDITE_PERMIS": f"{deb:%d/%m/%Y}",
        "CTR_DATE_FIN_VALIDITE_PERMIS": f"{fin:%d/%m/%Y}", "CTR_SALAIRE_NET": f"{w['net']:.2f}",
        "CTR_REFERENCE_DOMICILIATION": dom, "CTR_SIGNATURE_TRAVAILLEUR_PRESENTE": "True",
        "CTR_SIGNATURE_EMPLOYEUR_PRESENTE": "True", "CTR_CACHET_EMPLOYEUR_PRESENT": "True",
        "CTS_REFERENCE_DOCUMENT": d["CTS_REFERENCE_DOCUMENT"], "CTS_DUREE_MOIS": str(mois),
        "CTS_DATE_DEBUT_CONTRAT": f"{deb:%d/%m/%Y}", "CTS_NOM_PRENOM_TRAVAILLEUR": w["nom_prenom"],
        "CTS_NUMERO_PERMIS_TRAVAIL": d["CTS_NUMERO_PERMIS_TRAVAIL"], "CTS_DATE_DEBUT_VALIDITE_PERMIS": f"{deb:%d/%m/%Y}",
        "CTS_DATE_FIN_VALIDITE_PERMIS": f"{fin:%d/%m/%Y}", "CTS_SALAIRE_NET": f"{w['net']:.2f}",
        "CTS_SALAIRE_NET_ANCIEN": None, "CTS_PART_TRANSFERABLE": f"{w['part']:.2f}",
        "CTS_SIGNATURE_TRAVAILLEUR_PRESENTE": "True", "CTS_SIGNATURE_EMPLOYEUR_PRESENTE": "True",
        "CTS_CACHET_EMPLOYEUR_PRESENT": "True", "CTS_VISA_INSPECTION_TRAVAIL_PRESENT": "True",
        "TTR_NUMERO_PERMIS": d["CTS_NUMERO_PERMIS_TRAVAIL"], "TTR_NOM": w["nom"], "TTR_PRENOM": w["prenom"],
        "TTR_DATE_DEBUT": f"{deb:%d/%m/%Y}", "TTR_DATE_FIN": f"{fin:%d/%m/%Y}",
    }
    anomalies = {}
    if w["i"] == 3:     # date du titre de travail différente du CTS : à arbitrer dans la matrice
        anomalies["TTR_DATE_FIN"] = (f"{fin - timedelta(days=365):%d/%m/%Y}", 93, "OK", None)
    if w["i"] == 6:     # part lue avec une erreur d'OCR
        anomalies["DOM_PART_TRANSFERABLE"] = (f"{w['part']:.2f}".replace("0", "O", 1), 58, "REVIEW", "MONTANT_ILLISIBLE")
    page = {"DOM": 1, "CTR": 2, "CTS": 3, "TTR": 4}
    doc = {"DOM": "ENGAGEMENT_DOMICILIATION", "CTR": "CONTRAT_TRAVAIL", "CTS": "CONTRAT_SPECIFIQUE", "TTR": "TITRE_TRAVAIL"}
    rows = []
    for champ, v in vals.items():
        pre = champ.split("_")[0]
        norm, conf, dq, why = v, 95, "OK" if v is not None else "MISSING", None
        if champ in anomalies:
            norm, conf, dq, why = anomalies[champ]
        rows.append({"FICHIER": d["FICHIER"], "TYPE_DOSSIER": d["TYPE_DOSSIER"], "PAGE": page[pre], "TYPE_DOCUMENT": doc[pre],
                     "CHAMP": champ, "RAW": norm, "NORMALIZED": norm, "NORMALIZATION_STATUS": "AUTO_OK" if norm else "MISSING",
                     "NORMALIZATION_RULE": "DEMO", "FIELD_CONFIDENCE": conf,
                     "FIELD_CONFIDENCE_LEVEL": "HIGH" if conf >= 90 else "MEDIUM" if conf >= 70 else "LOW",
                     "DQ_STATUS_FIELD": dq, "DQ_REASON": why, "DQ_IMPORTANCE": "STANDARD"})
    return rows


def _fitz():
    try:
        import pymupdf as fitz
    except ImportError:
        try:
            import fitz  # type: ignore
        except ImportError:
            return None
    return fitz


def _pdf(path: Path, pages: list[tuple[str, list[str]]]) -> None:
    fitz = _fitz()
    if fitz is None:
        return
    doc = fitz.open()
    for title, lines in pages:
        pg = doc.new_page()
        pg.insert_text((50, 50), "SPECIMEN DE DEMONSTRATION - document fictif", fontsize=9, color=(0.7, 0, 0))
        pg.insert_text((50, 90), title, fontsize=16)
        y = 130
        for ln in lines:
            pg.insert_text((50, y), ln, fontsize=11)
            y += 22
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _pdf_dom(path: Path, d: dict, w: dict) -> None:
    _pdf(path, [
        ("ENGAGEMENT DE DOMICILIATION", [f"Client : {w['nom_prenom']}", f"Compte : {w['rib']}",
                                         f"Contrat : {d['CTS_REFERENCE_DOCUMENT']}",
                                         f"Du {d['CTS_DATE_DEBUT_VALIDITE_PERMIS']} au {d['CTS_DATE_FIN_VALIDITE_PERMIS']}",
                                         f"Part transferable mensuelle : {d['CTS_PART_TRANSFERABLE']} DZD"]),
        ("CONTRAT DE TRAVAIL", [f"Travailleur : {w['nom_prenom']}", f"Permis : {d['CTS_NUMERO_PERMIS_TRAVAIL']}",
                                f"Salaire net : {d['CTS_SALAIRE_NET']} DZD"]),
        ("CONTRAT DE TRAVAIL SPECIFIQUE", [f"Travailleur : {w['nom_prenom']}", f"Salaire net : {d['CTS_SALAIRE_NET']}",
                                           f"Part transferable : {d['CTS_PART_TRANSFERABLE']}",
                                           f"Validite : {d['CTS_DATE_DEBUT_VALIDITE_PERMIS']} - {d['CTS_DATE_FIN_VALIDITE_PERMIS']}"]),
        ("TITRE DE TRAVAIL", [f"Nom : {w['nom']}", f"Prenom : {w['prenom']}", f"Permis n : {d['CTS_NUMERO_PERMIS_TRAVAIL']}"]),
    ])


def _planning(P: Paths) -> pd.DataFrame:
    from . import planning
    dos = pd.read_excel(P.validation_xlsx, sheet_name="DOSSIERS_A_VALIDER", dtype=str)
    man = pd.read_excel(P.validation_xlsx, sheet_name="A_VALIDER_MANUELLEMENT", dtype=str)
    aug = pd.read_excel(P.param_aug_xlsx)
    res = planning.build_planning(dos, man, aug)
    planning.export(res, P.planning_xlsx, P.planning_csv)
    return planning.load_planning(P.planning_csv)


def _rh_rows(W: list[dict], M: date) -> list[dict]:
    rows = []
    for w in W:
        parts = [(None, 30, w["start"], w["end"], w["permis"])]
        if w["renouv"]:
            parts = [("PART 1", 17, w["start"], w["end"], w["permis"]),
                     ("PART 2", 13, w["new_start"], w["new_end"], w["permis_new"])]
        for part, jours, deb, fin, permis in parts:
            pt = round(0.95 * w["net"] * jours / 30, 2)
            rows.append({
                "Matricule SAP": w["matricule"], "Matricule Kelio": f"{w['nat']}-{7000 + w['i']}", "NOM ET PRENOM": w["nom_prenom"],
                "Nom": w["nom"], "Prénom": w["prenom"], "DATE DE NAISSANCE": f"{10 + w['i']:02d}/0{1 + w['i'] % 9}/198{w['i'] % 10}",
                "Lieu de naissance": "ISTANBUL" if w["nat"] == "TR" else "BANGALORE", "NOM DU PERE": w["pere"],
                "NOM DE LA MERE": w["mere"], "Date d'entrée en Algérie": f"{w['start'] - timedelta(days=30):%d/%m/%Y}",
                "Fonction": "TECHNICIEN", "Qualification": "CERTIFICAT D'EXPERIENCE", "Permit travail n°": permis.replace(" ", ""),
                "DATE DE DELIVRANCE": f"{deb - timedelta(days=20):%d/%m/%Y}", "Date d'émission": f"{deb:%d/%m/%Y}",
                "Fin de validité": f"{fin:%d/%m/%Y}", "DUREE DU PERMIS DE TRAVAIL (Mois)": 24,
                "Date d'attestation": f"{deb:%d/%m/%Y}", "Salaire Poste": round(w["net"] * 1.62, 2),
                "Retenue SS": round(w["net"] * 0.146, 2), "IRG": -round(w["net"] * 0.49, 2), "Retenue Mutuelle": -1500,
                "Salaire Net": w["net"], "Part transférable": pt, "Part Payable": round(w["net"] * jours / 30 - pt, 2),
                "Compte bancaire (Algérie)": w["rib"], "Nom de banque (Algérie)": "BNP PARIBAS",
                "Agence de bank (Algérie)": "AGENCE MOSTAGANEM",
                "Compte devise": " ".join(w["iban"][k:k + 4] for k in range(0, len(w["iban"]), 4)) if w["nat"] == "TR" else w["iban"],
                "Nom de banque (Etrange)": w["banque"][:19], "Pays de la banque (Etrange)": w["pays"],
                "Nom du bénéficiaire": w["benef"].split()[0], "Prénom du bénéficiaire": " ".join(w["benef"].split()[1:]),
                "Devise du transfert": w["devise"], "MOIS DE PAIEMENT": f"{MOIS_FR[M.month]}-{M.year}",
                "SWIFT/BIC": w["swift"][:8] if w["nat"] == "TR" else w["swift"], "Adresse": w["adresse"],
                "Contrat N°": str(600 + w["i"]), "Date d'établissement": f"{_add_months(M, -1):%d/%m/%Y}",
                "Nombre de jours": jours, "PART": part})
    return rows


def _referentiels(P: Paths, W: list[dict], M: date, today: date) -> None:
    from . import referentiels
    rh = pd.DataFrame(_rh_rows(W, M))
    d = referentiels.folder(P, "rh")
    d.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(d / f"base_RH_{M:%Y%m}.xlsx", engine="openpyxl") as xw:
        rh.to_excel(xw, sheet_name=f"BNP_{M:%m_%Y}_Consolidé", index=False)
    atf = pd.DataFrame([{"SIRA": k, "MATRICULE SAP": int(w["matricule"]), "Nom et Prenom": w["nom_prenom"],
                         "Adresse": w["adresse"][:22], "NET IMPOSABLE": round(w["net"] * 1.47, 2),
                         "IRG": round(w["net"] * 0.47, 2), "SALAIRE NET": w["net"], "MONTANT TRANSFERABLE": w["part"],
                         "BANQUE": "AGENCE MOSTAGANEM", "periode": datetime(M.year, M.month, 1)}
                        for k, w in enumerate(W, start=1)])
    d = referentiels.folder(P, "atf")
    d.mkdir(parents=True, exist_ok=True)
    atf.to_excel(d / f"ATF_BNP_{M:%m_%Y}.xlsx", index=False, sheet_name=f"ATF_BNP_{M:%m_%Y}")
    # autorisations : page du portail « Transfert des salaires des expatriés » (HTML), comme le fichier réel
    d = referentiels.folder(P, "aut")
    d.mkdir(parents=True, exist_ok=True)
    fmt = lambda x: f"{x:,.2f}".replace(",", " ").replace(".", ",")  # noqa: E731
    total = sum(w["part"] for w in W)
    lignes = "".join(f"<tr><td>{k}</td><td>{w['nom_prenom']} ({MOIS_FR[M.month]})</td><td>—</td><td>—</td><td>1.0000</td>"
                     f"<td>mois</td><td>{fmt(w['part'])} DZD</td><td>{fmt(w['part'])} DZD</td><td><button>Domicilier</button></td></tr>"
                     for k, w in enumerate(W, start=1))
    html = (f"<html><body><div>Recherche/Personnel étranger : Transfert des salaires des expatriés</div>"
            f"<div>Code banque: 027 — Code agence: 00795</div><h1>Personnel étranger : Transfert des salaires des expatriés</h1>"
            f"<p>NIF</p><p>{NIF}</p><p>RC</p><p>31 B 0000000-00/31</p><p>Raison sociale</p><p>{RAISON}</p>"
            f"<p>Total de la demande</p><p>{fmt(total)} DZD</p><div>Articles ({len(W)})</div>"
            f"<table><tr><th>N°</th><th>Désignation</th><th>Stock</th><th>En cours</th><th>Qté</th><th>Unité</th>"
            f"<th>Prix unitaire</th><th>Total</th><th>Actions</th></tr>{lignes}</table></body></html>")
    f = d / f"{M.year}_demo {MOIS_FR[M.month].lower()}.html"
    f.write_text(html, encoding="utf-8")
    import json as _json
    f.with_name(f.name + ".json").write_text(_json.dumps({"annee": M.year, "date_validation":
                                                          f"{today - timedelta(days=12):%d/%m/%Y}"}), encoding="utf-8")


def _periodes(PL: pd.DataFrame, client: str) -> pd.DataFrame:
    return PL[PL["SIEGE_RACINE"].astype(str) == client].sort_values("DATE_DEBUT_SEGMENT")


def _historique(P: Paths, W: list[dict], M: date, PL: pd.DataFrame, today: date, stamp: datetime) -> None:
    """Exports CarthagoTrade / Atlas / blocages / soldes : passé exécuté (avec anomalies voulues) + TL du mois."""
    trade, atlas = [], []
    seq = [26010000]

    def ref():
        seq[0] += 37
        return f"TL{seq[0]}"
    mois_m = f"{M:%Y-%m}"
    for w in W:
        for r in _periodes(PL, w["client"]).itertuples():
            per = r.PERIODE_TL
            if per[:7] >= mois_m:
                continue
            if w["i"] == 10 and per[:7] in (f"{_add_months(M, -4):%Y-%m}", f"{_add_months(M, -3):%Y-%m}"):
                continue                                             # reliquat : 2 mois jamais transférés
            mt = float(r.MONTANT_MAX_THEORIQUE)
            d_dem = date.fromisoformat(r.DATE_FIN_SEGMENT) + timedelta(days=6)
            d_val = d_dem + timedelta(days=12)
            t = ref()
            statut = "Exécuté"
            objet = f"{per[5:7]}{per[:4]}{per[7:]}"
            trade.append(_trade_row(w, t, mt, objet, statut, d_dem, d_val))
            anomalie_trade = w["i"] == 11 and per[:7] == f"{_add_months(M, -2):%Y-%m}"
            if not anomalie_trade:                                   # « Exécuté » Trade sans débit Atlas
                atlas.append(_atlas_row(w, t, mt, objet, d_val))
                if w["i"] == 3 and per[:7] == f"{_add_months(M, -5):%Y-%m}":
                    atlas.append(_atlas_row(w, t, mt, objet, d_val + timedelta(days=3)))   # débité deux fois
        if w["i"] == 2:     # TL rejeté puis ressaisi
            trade.append(_trade_row(w, ref(), w["part"], f"{_add_months(M, -6):%m%Y}", "Rejeté",
                                    _add_months(M, -5), None))
    pd.DataFrame(trade, columns=TRADE_EXPORT_COLUMNS).to_excel(
        P.historique_dir / f"Transfert libre -{stamp:%Y-%m-%dT%H-%M-%S}.xlsx", index=False)
    pd.DataFrame(atlas).to_excel(P.historique_dir / f"Historique_mouvement_altas-{stamp:%Y-%m-%dT%H%M%S}.xlsx", index=False)


TRADE_EXPORT_COLUMNS = ["Référence", "Code client", "Nom", "Code CRDS", "Montant", "Devise ordre de virement",
                        "Devise cible", "Date valeur", "Agence", "Date demande", "Statut", "Etape en cours",
                        "Utilisateur de destination", "Pli reçu niveau siège", "Nature du transfert", "Nature du compte",
                        "Nom/Raison sociale bénéficiaire", "Nom2", "Nom banque bénéficiaire", "Nom3", "Objet du transfert",
                        "Décision conformité", "Décision Trade", "Référence prédom", "N° domiciliation",
                        "Type de transfert", "Compte à débiter (Siège+ Racine+ Ordinal)", "Adresse",
                        "Adresse bénéficiaire", "Compte bénéficiaire", "Code Swift", "Client Vip"]


def _trade_row(w, ref, mt, objet, statut, d_dem, d_val, dom=None) -> dict:
    return {"Référence": ref, "Code client": w["client"], "Nom": w["nom_prenom"], "Code CRDS": "BOCI",
            "Montant": round(mt, 2), "Devise ordre de virement": "DZD", "Devise cible": w["devise"],
            "Date valeur": f"{d_val:%d/%m/%Y}" if d_val else None, "Agence": "BNP07310",
            "Date demande": f"{d_dem:%d/%m/%Y}" if d_dem else None, "Statut": statut,
            "Etape en cours": "06.Dispatcher BOCI" if statut.startswith("En cours") else None,
            "Utilisateur de destination": None, "Pli reçu niveau siège": "Non",
            "Nature du transfert": "Transfert International", "Nature du compte": "Dinar",
            "Nom/Raison sociale bénéficiaire": w["benef"], "Nom2": "Turquie" if w["nat"] == "TR" else "Inde",
            "Nom banque bénéficiaire": w["banque"], "Nom3": "Turquie" if w["nat"] == "TR" else "Inde",
            "Objet du transfert": objet, "Décision conformité": None, "Décision Trade": None,
            "Référence prédom": w["predom"], "N° domiciliation": dom or w["dom"], "Type de transfert": "Salaire",
            "Compte à débiter (Siège+ Racine+ Ordinal)": f"{w['rib'][4:]}DZD", "Adresse": ADRESSE_CLIENT,
            "Adresse bénéficiaire": w["adresse"][:80], "Compte bénéficiaire": w["iban"], "Code Swift": w["swift"][:8],
            "Client Vip": "Non"}


def _atlas_row(w, ref, mt, objet, d_val) -> dict:
    return {"id_tiers": w["client"], "Date_operation_Atlas": f"{d_val:%d/%m/%Y}",
            "lib_complementaire_mvt": f"TRF {ref} SALAIRE {objet}", "Reference_denotage": objet,
            "Reference_client": w["client"], "Date_Valeur_comptable": f"{d_val:%d/%m/%Y}",
            "Montant_mvt_comptable": -round(mt, 2)}


def _ov(P: Paths, W: list[dict], M: date, today: date) -> None:
    """Lot d'OV du mois M « lus par l'IA » (classeur au format du notebook) + PDF spécimens."""
    from openpyxl import Workbook
    from .transfers import V26_SCHEMA
    extra = {"OV": [("devise_transfert", "Devise de transfert"), ("reference_paiement", "Reference paiement (case 70)"),
                    ("beneficiaire_adresse", "Bénéficiaire adresse")]}
    schema = {g: list(cs) + [x for x in extra.get(g, []) if x[0] not in {k for k, _ in cs}] for g, cs in V26_SCHEMA.items()}
    labels = {g: dict(cs) for g, cs in schema.items()}
    rows = []
    for w in W:
        parts = [(None, 30)] if not w["renouv"] else [("P1", 17), ("P2", 13)]
        for part, jours in parts:
            mt = round(0.95 * w["net"] * jours / 30, 2)
            iban = w["iban"]
            pages = "ANNEXE_I, ANNEXE_II, BULLETIN, OV"
            manquantes = None
            if w["i"] == 4:
                mt = round(mt + 25000, 2)                       # OV au-dessus de la part transférable
            if w["i"] == 5:
                iban = iban[:-6] + str((int(iban[-6]) + 3) % 10) + iban[-5:]   # IBAN mal recopié (1 chiffre)
            if w["i"] == 6:
                pages, manquantes = "ANNEXE_II, BULLETIN, OV", "ANNEXE_I"
            case70 = {None: f"{MOIS_FR[M.month].title()} {M.year}", "P1": f"{MOIS_FR[M.month][:4].title()}.{M.year % 100} PART1",
                      "P2": f"{MOIS_FR[M.month]} {M.year} P2"}[part]
            f = f"07310-{w['racine']}-{M:%Y%m}{part or ''}.pdf"
            a2_part = round(0.95 * w["net"] * jours / 30, 2)
            rows.append({
                "META": dict(fichier=f, date_traitement=f"{today:%Y-%m-%d} 09:12:00", pages_trouvees=pages,
                             pages_manquantes=manquantes),
                "OV": dict(monnaie="DZD", montant_chiffres=mt, montant_lettres=en_lettres(mt), periode=case70,
                           reference_paiement=case70, tranche=part, date_demande=f"{_add_months(M, 1).replace(day=3):%d/%m/%Y}",
                           compte_debiteur=w["rib"], beneficiaire_nom=w["benef"], beneficiaire_compte=iban,
                           beneficiaire_adresse=w["adresse"], swift=w["swift"][:8], banque_nom=w["banque"],
                           devise_transfert=w["devise"], status="OK"),
                "ANNEXE_I": dict(nom_prenom=w["nom_prenom"], date_signature=f"{_add_months(M, 1).replace(day=2):%d/%m/%Y}",
                                 status="OK") if w["i"] != 6 else {},
                "ANNEXE_II": dict(mois_transfert=f"{MOIS_FR[M.month]} {M.year}", nom_prenom=w["nom_prenom"],
                                  compte_bancaire=w["rib"], salaire_mensuel=round(w["net"] * 0.93, 2), nombre_jours=jours,
                                  part_transferable=a2_part, pays_destination=w["pays"], compte_devise=w["iban"],
                                  numero_domiciliation=(w["dom_new"] if part == "P2" else w["dom"]), status="OK"),
                "BULLETIN": dict(nom_prenom=w["nom_prenom"], matricule=w["matricule"], mois_bulletin=f"{MOIS_FR[M.month]} {M.year}",
                                 salaire_base=round(w["net"] * 1.2, 2), salaire_brut=round(w["net"] * 1.62, 2),
                                 net_a_payer=w["net"]),
            })
            _pdf(P.tr_in / f, [
                ("ORDRE DE VIREMENT A L'ETRANGER", [f"Montant : {mt:,.2f} DZD".replace(",", " "), en_lettres(mt)[:80],
                                                    f"70 - Reference du paiement : {case70}", f"Donneur d'ordre : {w['rib']}",
                                                    f"Beneficiaire : {w['benef']}", f"Compte : {iban}", f"SWIFT : {w['swift']}",
                                                    f"Banque : {w['banque']}"]),
                ("ANNEXE I", [f"Je soussigne {w['nom_prenom']}", f"Compte : {w['rib']}"]) if w["i"] != 6 else
                ("PAGE BLANCHE", ["(Annexe I absente du dossier)"]),
                ("ANNEXE II", [f"Mois de {MOIS_FR[M.month]} {M.year}", f"Travailleur : {w['nom_prenom']}",
                               f"Part transferable : {a2_part:,.2f}".replace(",", " "), f"Jours : {jours}"]),
                ("BULLETIN DE PAIE", [f"Salarie : {w['nom_prenom']}", f"Matricule : {w['matricule']}",
                                      f"Net a payer : {w['net']:,.2f}".replace(",", " ")]),
            ])
    P.tr_out.mkdir(parents=True, exist_ok=True)
    _write_audit(P.tr_out / f"audit_transferts_{today:%Y%m%d}_0915.xlsx", rows, schema)
    # historique : un lot d'OV d'il y a 4 mois (SHARMA, rejeté : montant au-dessus de l'autorisation) et l'ATF /
    # les autorisations de ce mois-là (SHARMA absent de l'ATF le mois suivant)
    w = next(x for x in W if x["nom"] == "SHARMA")
    m4, m3 = _add_months(M, -4), _add_months(M, -3)
    old = {"META": dict(fichier=f"07310-{w['racine']}-{m4:%Y%m}.pdf", pages_trouvees="ANNEXE_I, ANNEXE_II, BULLETIN, OV"),
           "OV": dict(monnaie="DZD", montant_chiffres=round(w["part"] + 40000, 2), montant_lettres=en_lettres(round(w["part"] + 40000, 2)),
                      periode=f"{MOIS_FR[m4.month]} {m4.year}", compte_debiteur=w["rib"], beneficiaire_nom=w["benef"],
                      beneficiaire_compte=w["iban"], swift=w["swift"][:8], banque_nom=w["banque"]),
           "ANNEXE_I": dict(nom_prenom=w["nom_prenom"]), "ANNEXE_II": dict(mois_transfert=f"{MOIS_FR[m4.month]} {m4.year}",
                                                                            nom_prenom=w["nom_prenom"], compte_bancaire=w["rib"],
                                                                            part_transferable=w["part"]),
           "BULLETIN": dict(nom_prenom=w["nom_prenom"], mois_bulletin=f"{MOIS_FR[m4.month]} {m4.year}", net_a_payer=w["net"])}
    (P.data_root / "historique_ov").mkdir(parents=True, exist_ok=True)
    _write_audit(P.data_root / "historique_ov" / f"audit_transferts_{m4:%Y%m}15_1000.xlsx", [old], schema)
    from . import transfers as _tr
    _tr.save_decisions(P.decisions_csv, pd.DataFrame([{"FICHIER": old["META"]["fichier"], "AUDIT_SOURCE": "historique",
                                                       "DECISION": "REJETE",
                                                       "COMMENTAIRE": "Montant OV supérieur à l'autorisation du mois"}]),
                       "analyste.demo")
    from . import referentiels
    for mm, sans in ((m4, None), (m3, "SHARMA")):
        pd.DataFrame([{"SIRA": k, "MATRICULE SAP": int(x["matricule"]), "Nom et Prenom": x["nom_prenom"], "Adresse": "-",
                       "NET IMPOSABLE": round(x["net"] * 1.47, 2), "IRG": round(x["net"] * 0.47, 2), "SALAIRE NET": x["net"],
                       "MONTANT TRANSFERABLE": x["part"], "BANQUE": "AGENCE MOSTAGANEM", "periode": datetime(mm.year, mm.month, 1)}
                      for k, x in enumerate(W, start=1) if x["nom"] != sans]).to_excel(
            referentiels.folder(P, "atf") / f"ATF_BNP_{mm:%m_%Y}.xlsx", index=False)


def _write_audit(path: Path, rows: list[dict], schema: dict) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Dossiers"
    cols = [(g, k, lab) for g, cs in schema.items() for k, lab in cs]
    c0 = 1
    for g, cs in schema.items():
        ws.cell(row=1, column=c0, value=g)
        c0 += len(cs)
    for i, (_, _, lab) in enumerate(cols, 1):
        ws.cell(row=2, column=i, value=lab)
    for n, d in enumerate(rows, 3):
        for i, (g, k, _) in enumerate(cols, 1):
            ws.cell(row=n, column=i, value=d.get(g, {}).get(k))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _modele_vigilance(path: Path) -> None:
    """Modèle « CHECK LIST TRANSFERT LIBRE » (mêmes libellés que le modèle BNP) pour la démonstration."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Feuil1"
    ws["A2"] = "CHECK LIST TRANSFERT LIBRE"
    for r, lab in ((5, "Client:"), (6, "Réf Dossier:"), (7, "Circuit:"), (8, "Type de dossier:")):
        ws.cell(r, 1, lab)
    for c, lab in zip("ABCDE", ["Contrôle", "Type", "Valeur", "HIT / No HIT / KO", "COMMENTAIRES"]):
        ws[f"{c}10"] = lab
    ws["A11"] = "VIGILANCE"
    ws.merge_cells("A11:E11")
    rub = ["Banque du bénéficiaire", "Pays de la banque du bénéficiaire", "Banque intermédiaire",
           "Pays de la banque intermédiaire", "Bénéficiaire", "Pays du bénéficiaire", "Pays d’origine et provenance",
           "Nom du fabricant", "Pays du fabricant", "Compagnie maritime ou aérienne", "Navire", "Expéditeur sur le BL",
           "Agent signataire P/C compagnie maritime", "Transitaire", "Transporteur LVI", "Port ou aéroport d’embarquement",
           "Port ou aéroport de débarquement", "Lieu de prise en charge sur le BL", "Lieu établissement document transport",
           "Lieu de travail", "Divers individus sur les documents", "Divers entités sur les documents", "Divers lieux"]
    for i, lab in enumerate(rub):
        ws.cell(12 + i, 1, lab)
        ws.cell(12 + i, 2, "I" if "individus" in lab else "C")
    ws["A37"] = "Valideur :"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _decisions_vigilance(P: Paths, home: Path, user: str, today: date) -> None:
    """Contrôles réels du lot, validation des dossiers conformes, vigilance (RAS + un HIT), TL saisis par
    l'agence, blocages et soldes : il reste à la démo à valider, analyser le HIT et générer les fichiers."""
    import io
    from openpyxl import load_workbook
    from . import referentiels, transfers
    from . import vigilance as vg
    import dataclasses
    from .config import load_settings
    from .planning import load_planning
    s = dataclasses.replace(load_settings(data_root=str(P.data_root)), employeur_raison_sociale=RAISON)
    audit_path = transfers.latest_audit(P.tr_out)[0]
    A = transfers.read_audit_xlsx(audit_path)
    PL = load_planning(P.planning_csv)
    R = referentiels.load_all(P)
    sy, ch = transfers.control_each(A, PL, transfers.load_registre(P.registre_csv), R["atf"], R["aut"], s, today, R["rh"])
    sy, ch = transfers.mark_lot_duplicates(sy, ch)
    conformes = list(sy[sy["STATUT_CONTROLE"] == "CONFORME"]["FICHIER"])
    # deux dossiers conformes laissés à valider en direct ; le bénéficiaire tiers (à revoir) validé avec motif
    tiers = [f for f in sy[sy["STATUT_CONTROLE"] == "A_REVOIR"]["FICHIER"] if f.startswith("07310-990109")]
    valides = conformes[:-2] + tiers
    ed = sy[sy["FICHIER"].isin(valides)][["FICHIER", "AUDIT_SOURCE", "STATUT_CONTROLE", "PERIODE_ID"]].assign(
        DECISION="VALIDE", COMMENTAIRE="Contrôles conformes — validé (démo)")
    ed.loc[ed["FICHIER"].isin(tiers), "COMMENTAIRE"] = "Bénéficiaire = épouse du salarié (lien justifié par l'employeur)"
    transfers.save_decisions(P.decisions_csv, ed, "analyste.demo")
    # vigilance : lot préparé avec le modèle, retours RAS sauf un HIT (bénéficiaire tiers)
    _modele_vigilance(P.vigilance_template)
    ds = vg.dossiers_from_ov(sy, A, valides, None, s.vigilance_individus)
    units = vg.build_units(ds)
    m = vg.prepare(units, P.vigilance_template, P.vigilance_dir / "envois" / "demo", s.employeur_raison_sociale)
    jp = vg.journal_path(P)
    # le mois précédent, le HIT « bénéficiaire » de SAHIN a déjà été analysé non avéré : TL-One doit le reconnaître
    import json as _json
    tiers_u = next((u for u in units if u["beneficiaire"]), None)
    if tiers_u:
        jp.parent.mkdir(parents=True, exist_ok=True)
        with jp.open("a", encoding="utf-8") as fh:
            fh.write(_json.dumps({
                "ts": (datetime.combine(today, datetime.min.time()) - timedelta(days=31)).isoformat(timespec="seconds"),
                "user": "analyste.vigilance", "action": "DECISION", "uid": "MOIS_PRECEDENT",
                "ref_client": tiers_u["ref_client"], "fichiers": [], "decision": "LEVEE",
                "motif": "Bénéficiaire : non avéré", "verdicts": [{
                    "rubrique": "Bénéficiaire", "rang": 1, "valeur": tiers_u["beneficiaire"], "resultat": "HIT",
                    "statut": "HIT", "verdict": "NON_AVERE",
                    "commentaire": "Homonymie : la bénéficiaire est l'épouse du salarié (acte de mariage transmis par "
                                   "l'employeur), aucune correspondance avec la personne listée"}]},
                ensure_ascii=False) + "\n")
    vg.append(jp, "analyste.demo", "ENVOI", lot=m["lot"], unites=units)
    ups = []
    hit_uid = None
    for u in units:
        wb = load_workbook(P.vigilance_dir / "envois" / "demo" / u["fichier"])
        ws = wb.worksheets[0]
        r = vg.locate(ws)
        for row in range(r["first"], r["last"] + 1):
            if vg.text(ws.cell(row, 3).value):
                n = len(str(ws.cell(row, 3).value).split("\n"))
                ws.cell(row, 4).value = "\n".join(["OK ==> RAS"] * n)
            elif vg.text(ws.cell(row, 2).value):
                ws.cell(row, 4).value = "Rubrique non renseignée"
        if u["beneficiaire"] and hit_uid is None:          # bénéficiaire ≠ salarié : HIT déjà connu (mois précédent)
            ws.cell(r["beneficiaire"], 4).value = "HIT"
            hit_uid = u["uid"]
        elif "HDFC" in (u.get("swift") or "").upper():      # nouvelle alerte sur la banque : à analyser en direct
            ws.cell(r["banque"], 4).value = "HIT"
        b = io.BytesIO()
        wb.save(b)
        ups.append((u["fichier"], b.getvalue()))
    vg.import_returns(ups, vg.read_journal(jp), jp, "analyste.demo", P.vigilance_dir / "retours")
    vig = {f: v["statut"] for f, v in vg.statuses(vg.read_journal(jp)).items()}
    dec = transfers.load_decisions(P.decisions_csv)
    transfers.register(P.registre_csv, sy, dec, "analyste.demo", vigilance=vig)
    # TL saisis par l'agence (CarthagoTrade) pour les dossiers validés passés en vigilance
    W = {w["client"]: w for w in _workers(demo_month(today))}
    hist = sorted(P.historique_dir.glob("Transfert libre -*.xlsx"))[-1]
    trade = pd.read_excel(hist, dtype=object)
    seq = 26090000
    nouveaux, blocages, soldes, sans_blocage = [], [], [], []
    ok = sy[sy["FICHIER"].isin([f for f, v in vig.items() if v in ("RAS", "HIT_LEVE")])]
    for k, r in enumerate(ok.itertuples()):
        w = W[r.SIEGE_RACINE]
        seq += 53
        per = r.PERIODE_TL
        objet = f"{per[5:7]}{per[:4]}{per[7:]}"
        dom = w["dom_new"] if per.endswith("P2") else w["dom"]
        nouveaux.append(_trade_row(w, f"TL{seq}", float(r.MONTANT_OV), objet, "En cours d'exécution",
                                   today - timedelta(days=4), None, dom))
        if k % 3 == 2:
            sans_blocage.append(w["client"])
        else:               # 1 TL sur 3 sans blocage : fichier de blocages pour l'agence
            blocages.append({"SIEGE": w["client"][:5], "RACINE": w["client"][5:], "ORDINAL": "001", "DEVISE": "DZD",
                             "REFERENCE DENOTAGE DU MOUVEMENT": objet, "MT BLOCAGE DEV CPT": -round(float(r.MONTANT_OV) + 5000, 2),
                             "DATE DEBUT BLOCAGE": f"{today - timedelta(days=3):%Y%m%d}",
                             "DATE LIMITE BLOCAGE": f"{today + timedelta(days=87):%Y%m%d}"})
    for w in W.values():
        # le premier TL sans blocage tombe sur un compte au solde insuffisant ; les autres iront au fichier agence
        solde = 40_000.0 if sans_blocage and w["client"] == sans_blocage[0] else 2_400_000.0
        soldes.append({"compte_local": f"{w['rib'][4:]}DZD", "Solde_Compta": solde + 15000, "Solde_Dispo": solde})
    trade = pd.concat([trade, pd.DataFrame(nouveaux, columns=TRADE_EXPORT_COLUMNS)], ignore_index=True)
    trade.to_excel(hist, index=False)
    stamp = hist.name.split("-", 1)[1].replace(".xlsx", "").strip()
    st2 = datetime.strptime(stamp, "%Y-%m-%dT%H-%M-%S")
    pd.DataFrame(blocages).to_excel(P.historique_dir / f"Blocage-{st2:%Y-%m-%dT%H%M%S}.xlsx", index=False)
    pd.DataFrame(soldes).to_excel(P.historique_dir / f"comptes_solde-{st2:%Y-%m-%dT%H%M%S}.xlsx", index=False)
