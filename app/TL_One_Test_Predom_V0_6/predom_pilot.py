"""TL-One Prédom: test de sélection d'accords sans certification métier.

- Un client peut avoir plusieurs références DOM.
- Une seule référence *candidate* par client, sur statuts éligibles et date décision.
- Ne remplace pas la confrontation au CTS ni la validation humaine.
- Aucun service externe, aucun envoi de données.
"""
from __future__ import annotations
import argparse, csv, json, re, unicodedata
from collections import defaultdict, Counter
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent

def _empty(x):
    return x is None or (isinstance(x, str) and not x.strip())

def _txt(x):
    if x is None: return ""
    return x.strip() if isinstance(x, str) else str(x)

def _status(x):
    return " ".join(_txt(x).casefold().split())

def _date(x):
    if isinstance(x, datetime): return x
    if isinstance(x, date): return datetime.combine(x,datetime.min.time())
    if _empty(x): return None
    if isinstance(x, (int, float)):
        # Une date Excel non formatée comme date ne doit pas devenir une décision valide.
        return None
    text=_txt(x)
    try:
        return datetime.fromisoformat(text.replace('Z','+00:00')).replace(tzinfo=None)
    except ValueError:
        pass
    for fmt in ('%d/%m/%Y %H:%M:%S','%d/%m/%Y %H:%M','%d/%m/%Y',
                '%d-%m-%Y %H:%M:%S','%d-%m-%Y', '%Y-%m-%d'):
        try: return datetime.strptime(text,fmt)
        except ValueError: continue
    return None

def read_table(path: Path, header_row: int = 1, sheet: str | None = None):
    if header_row < 1: raise ValueError('header_row doit être >= 1')
    if not path.is_file(): raise FileNotFoundError(f'Fichier introuvable : {path}')
    suffix=path.suffix.lower()
    if suffix in ('.xlsx','.xlsm'):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError('Installer openpyxl : python -m pip install openpyxl') from exc
        wb=load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet is not None and sheet not in wb.sheetnames:
                raise ValueError(f'Feuille inconnue {sheet!r} ; feuilles disponibles : {wb.sheetnames}')
            ws=wb[sheet] if sheet is not None else wb.active
            iterator=ws.iter_rows(values_only=True)
            for _ in range(header_row-1): next(iterator,None)
            headers=next(iterator,None)
            if headers is None: raise ValueError('Aucun en-tête à la ligne indiquée')
            headers=list(headers)
            # Ignorer uniquement les colonnes terminales réellement vides.
            while headers and _empty(headers[-1]): headers.pop()
            rows=[]
            for pos, values in enumerate(iterator, start=header_row+1):
                vals=list(values)[:len(headers)]
                vals += [None]*(len(headers)-len(vals))
                if all(_empty(v) for v in vals): continue
                rows.append((pos,vals))
        finally: wb.close()
    elif suffix=='.csv':
        with path.open(encoding='utf-8-sig',newline='') as f:
            sample=f.read(4096)
            f.seek(0)
            try: dialect=csv.Sniffer().sniff(sample,delimiters=';,\t')
            except csv.Error: dialect=csv.excel
            reader=csv.reader(f,dialect)
            for _ in range(header_row-1): next(reader,None)
            headers=next(reader,None)
            if headers is None: raise ValueError('Aucun en-tête à la ligne indiquée')
            while headers and _empty(headers[-1]): headers.pop()
            rows=[]
            for pos, values in enumerate(reader,start=header_row+1):
                if len(values)>len(headers): raise ValueError(f'Ligne {pos}: plus de valeurs que d’en-têtes')
                vals=values + [None]*(len(headers)-len(values))
                if all(_empty(v) for v in vals): continue
                rows.append((pos,vals))
    else:
        raise ValueError('Format accepté : .xlsx, .xlsm ou .csv')
    if any(not isinstance(h,str) or not h.strip() for h in headers):
        raise ValueError('En-tête vide/non textuel dans le fichier')
    if len(set(headers))!=len(headers):
        raise ValueError('En-têtes dupliqués: pas de rapprochement possible')
    return headers,rows

def analyze(headers, rows, cfg: dict):
    fields=cfg['colonnes']
    required=('client_id','dom_reference','status','decision_date')
    missing=[f'{key}: {fields.get(key)!r}' for key in required if not isinstance(fields.get(key),str) or fields[key] not in headers]
    if missing:
        raise ValueError('Colonnes exactes absentes: '+', '.join(missing)+'\nEn-têtes trouvés: '+repr(headers))
    allowed={_status(s) for s in cfg['statuts_eligibles']}
    by_client=defaultdict(list)
    records=[]
    for pos,vals in rows:
        raw=dict(zip(headers,vals))
        client=raw[fields['client_id']]; ref=raw[fields['dom_reference']]
        statut=raw[fields['status']]; date_val=raw[fields['decision_date']]
        dt=_date(date_val)
        issues=[]
        if _empty(client): issues.append('CLIENT_ABSENT')
        elif isinstance(client,(int,float)): issues.append('CLIENT_EXCEL_NUMERIQUE_ZEROS_A_VERIFIER')
        if _empty(ref): issues.append('REFERENCE_DOM_ABSENTE')
        elif isinstance(ref,(int,float)): issues.append('DOM_EXCEL_NUMERIQUE_ZEROS_A_VERIFIER')
        if _empty(date_val): issues.append('DATE_DECISION_ABSENTE')
        elif dt is None: issues.append('DATE_DECISION_ILLISIBLE')
        eligible=_status(statut) in allowed
        if not eligible: issues.append('STATUT_NON_ELIGIBLE')
        usable=eligible and dt is not None and not _empty(client) and not _empty(ref) and not isinstance(client,(int,float)) and not isinstance(ref,(int,float))
        r={'ligne_excel':pos,'client':_txt(client),'reference_dom':_txt(ref),'statut':_txt(statut),
           'request_decision_date_brute':_txt(date_val), 'request_decision_date_iso': dt.isoformat(sep=' ') if dt else '',
           'eligible':eligible,'candidate_comparable':usable,'anomalies':issues,
           'selection':'NON_RETENUE'}
        records.append(r)
        if r['client']: by_client[r['client']].append(r)
    summary=[]
    for client, items in sorted(by_client.items()):
        candidates=[r for r in items if r['candidate_comparable']]
        if not candidates:
            state='AUCUNE_CANDIDATE_ELIGIBLE'; winner=''; chosen=[]
        else:
            date_max=max(r['request_decision_date_iso'] for r in candidates)
            latest=[r for r in candidates if r['request_decision_date_iso']==date_max]
            references={r['reference_dom'] for r in latest}
            if len(references)>1:
                state='AMBIGUITE_MEME_DATE'; winner=''; chosen=[]
                for r in latest: r['selection']='AMBIGUE'
            else:
                winner=next(iter(references)); state='PROPOSEE_A_CONFRONTER_CTS'
                chosen=[r for r in latest if r['reference_dom']==winner]
                for r in chosen: r['selection']='PROPOSEE'
        summary.append({'client':client,'nb_lignes_historique':len(items),'nb_eligibles':sum(x['eligible'] for x in items),
                        'reference_proposee':winner,'statut_selection':state,
                        'ca10':'ABSENT_NON_VERIFIE',
                        'validation_dossier':'NON_EFFECTUEE',
                        'date_derniere_decision':max((r['request_decision_date_iso'] for r in chosen),default='')})
    return {'nb_lignes':len(records),'nb_clients':len(by_client), 'bilan_statuts':dict(Counter(r['statut'] for r in records)),
            'dossiers':summary,'historique_predom':records,
            'note':'Choix candidat sur statut et date; confrontations CTS/CTR/DOM/TTR et signature CA10 NON réalisées.'}

def dump_csv(path,rows):
    if not rows: return
    keys=list(rows[0].keys())
    with path.open('w', encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=keys,delimiter=';')
        writer.writeheader()
        for row in rows:
            writer.writerow({k:('|'.join(row[k]) if isinstance(row[k],list) else row[k]) for k in keys})

def run(path, config_path=None, output=None):
    cfg=json.loads(Path(config_path or ROOT/'config_predom.json').read_text(encoding='utf-8'))
    headers,rows=read_table(Path(path),header_row=cfg.get('header_row',1),sheet=cfg.get('sheet_name'))
    result=analyze(headers,rows,cfg)
    dest=Path(output or ROOT/'resultats')
    dest.mkdir(parents=True,exist_ok=True)
    (dest/'rapport_predom.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    dump_csv(dest/'selection_predom.csv',result['dossiers'])
    dump_csv(dest/'historique_predom.csv',result['historique_predom'])
    print(f"Lignes: {result['nb_lignes']} | Clients: {result['nb_clients']}")
    for status,count in Counter(s['statut_selection'] for s in result['dossiers']).items():print(f'  {status}: {count}')
    print('CA10 : ABSENT_NON_VERIFIE (pour tous les dossiers de ce pilote)')
    print('Dossier pilote : aucune validation officielle, aucun cachet, aucun Planning TL.')
    print('Rapports:', dest.resolve())
    return result

if __name__=='__main__':
    p=argparse.ArgumentParser(description='TL-One: pilote Prédom, sans certification')
    p.add_argument('fichier',help='Export Prédom .xlsx / .xlsm / .csv')
    p.add_argument('--config',default=str(ROOT/'config_predom.json'))
    p.add_argument('--output',default=str(ROOT/'resultats'))
    args=p.parse_args()
    run(args.fichier,args.config,args.output)
